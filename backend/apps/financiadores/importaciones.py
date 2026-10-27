"""Planillas privadas: vista previa y aplicación recuperable, sin bajas de padrón."""
from datetime import date, datetime
from copy import copy
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from time import monotonic
from uuid import UUID
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile
from zlib import error as ZlibError

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from django.conf import settings
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError as ApiValidationError

from apps.registros.models import normalizar_documento
from apps.simulacion.contexto import autor_real
from .models import Afiliado, Financiador, Importacion, Plan
from .permisos import requerir_financiador
from .services import catalogo_financiador, registrar_afiliado, registrar_consumo_externo, validar_consumo_externo
from .importacion_cobertura import es_conflicto_cobertura, revisar_cobertura, validar_opcion
from . import padron


ENCABEZADOS = {
    "padron": ("N.º de afiliado", "Documento", "Nombre y apellido", "Plan", "Vigente desde"),
    "consumos": ("N.º de afiliado", "Documento", "Prestación", "Fecha de prestación", "Cantidad", "Referencia externa"),
}
VERSION = 1
SALT = "financiadores.plantilla.v1"
ESTADOS_FILA = ("valida", "rechazada", "revision", "aplicada", "error_tecnico")
CAMPOS_V2 = tuple("titular_documento" if c == "titular" else c for c in padron.CAMPOS_GENERALES
                  if c not in ("solicitud_baja", "baja_efectiva", "fin_prestacional_informado", "motivo_baja", "causa_baja", "cobertura_habilitada"))


def encabezados(tipo, version=1):
    return ENCABEZADOS[tipo] + CAMPOS_V2 if tipo == "padron" and version == 2 else ENCABEZADOS[tipo]


def limites_importacion():
    """Valores conservadores ajustables sin cambiar el formato de las planillas."""
    return {
        "bytes": getattr(settings, "FINANCIADORES_XLSX_MAX_BYTES", 5 * 1024 * 1024),
        "expandido": getattr(settings, "FINANCIADORES_XLSX_MAX_EXPANDIDO", 25 * 1024 * 1024),
        "entradas": getattr(settings, "FINANCIADORES_XLSX_MAX_ENTRADAS", 128),
        "filas": getattr(settings, "FINANCIADORES_XLSX_MAX_FILAS", 2000),
        "segundos": getattr(settings, "FINANCIADORES_XLSX_MAX_SEGUNDOS", 30),
    }


def _tipo_valido(tipo):
    if tipo not in ENCABEZADOS:
        raise ValidationError("El tipo de importación debe ser padrón o consumos.")


def _catalogo(financiador):
    return catalogo_financiador(financiador).order_by("codigo")


def _texto_celda(celda, valor):
    # Forzar texto literal también para entradas que comienzan con =, +, - o @.
    celda.value = "" if valor is None else str(valor)
    celda.data_type = "s"
    celda.number_format = "@"


def generar_plantilla(*, financiador, usuario, tipo, version=1):
    requerir_financiador(usuario, financiador.pk, escritura=True)
    _tipo_valido(tipo)
    if version not in (1, 2) or (version == 2 and tipo != "padron"):
        raise ValidationError("La versión 2 sólo está disponible para padrón.")
    limite = limites_importacion()
    libro = Workbook()
    instrucciones = libro.active
    instrucciones.title = "LEEME"
    for numero, texto in enumerate((
        f"{financiador.nombre}: importación de {tipo}",
        "Complete solamente Carga. No cambie los encabezados ni la hoja _contexto.",
        "N.º de afiliado y documento son texto obligatorio; conserve los ceros iniciales.",
        "Las fechas admiten celdas de fecha o texto AAAA-MM-DD. No se admiten fórmulas.",
        "Revise el resumen antes de confirmar. Se importan válidas y se entregan errores por fila.",
        "Padrón: omitir personas nunca da bajas. Plan vacío significa sin plan.",
        "Consumos: cantidad entera positiva; referencia externa opcional para reconocer reenvíos.",
        f"Límite: {limite['filas']} filas, {limite['bytes'] // (1024 * 1024)} MiB por archivo .xlsx.",
    ), 1):
        _texto_celda(instrucciones.cell(numero, 1), texto)
    instrucciones.column_dimensions["A"].width = 110
    carga = libro.create_sheet("Carga")
    carga.append(encabezados(tipo, version))
    carga.freeze_panes = "A2"
    carga.auto_filter.ref = f"A1:{get_column_letter(len(encabezados(tipo, version)))}1"
    fin = limite["filas"] + 1
    for fila in carga.iter_rows(min_row=2, max_row=min(fin, 201), max_col=len(encabezados(tipo, version))):
        for celda in fila:
            celda.number_format = "@"
        fila[4 if tipo == "padron" else 3].number_format = "yyyy-mm-dd"
        if tipo == "consumos":
            fila[4].number_format = "0"
    for columna in ("A", "B", "C", "D", "E", "F"):
        carga.column_dimensions[columna].width = 27 if columna != "C" else 45
    nombre_catalogo = "Planes" if tipo == "padron" else "Prestaciones"
    catalogo = libro.create_sheet(nombre_catalogo)
    catalogo.append(["Código", "Nombre", "Valor para Carga"])
    opciones = (
        Plan.objects.filter(financiador=financiador, activo=True).order_by("codigo")
        if tipo == "padron" else _catalogo(financiador)
    )
    for opcion in opciones:
        fila = catalogo.max_row + 1
        for col, valor in enumerate((opcion.codigo, opcion.nombre, f"{opcion.codigo} — {opcion.nombre}"), 1):
            _texto_celda(catalogo.cell(fila, col), valor)
    if catalogo.max_row > 1:
        validacion = DataValidation(
            type="list", formula1=f"'{nombre_catalogo}'!$C$2:$C${catalogo.max_row}",
            allow_blank=tipo == "padron",
        )
        validacion.errorTitle = "Seleccione un código del catálogo"
        validacion.error = "Use el código o la opción disponible en la hoja de referencia."
        validacion.showErrorMessage = True
        carga.add_data_validation(validacion)
        validacion.add(f"{'D' if tipo == 'padron' else 'C'}2:{'D' if tipo == 'padron' else 'C'}{fin}")
    for col in ("A", "B", "C"):
        catalogo.column_dimensions[col].width = 45
    contexto = libro.create_sheet("_contexto")
    contexto["A1"] = signing.dumps({"financiador": financiador.pk, "tipo": tipo, "version": version}, salt=SALT)
    if version == 2:
        _texto_celda(instrucciones.cell(10, 1), "Columnas nuevas vacías no borran. Titular por documento: existente o fila válida anterior; un nivel, sin inferencias.")
        _texto_celda(instrucciones.cell(11, 1), "No cargar salud, discapacidad, CBU, CUIT de cuenta ni cobranza. Use recursos protegidos; no se almacenan en este lote.")
        _texto_celda(instrucciones.cell(12, 1), "Se revalida al confirmar. Ediciones posteriores al preview requieren una nueva vista previa; no se fuerza identidad.")
        campos = libro.create_sheet("Campos")
        campos.append(["Campo", "Tipo", "Opciones", "Regla"])
        for campo in CAMPOS_V2:
            if campo == "titular_documento":
                campos.append([campo, "texto", "", "Documento del titular del mismo financiador, existente o fila válida anterior."])
                continue
            modelo_campo = Afiliado._meta.get_field(campo)
            campos.append([campo, modelo_campo.get_internal_type(), ", ".join(str(opcion[0]) for opcion in modelo_campo.choices or ()), "Vacío = no informar; no borra datos existentes."])
        for col in ("A", "B", "C", "D"):
            campos.column_dimensions[col].width = 40 if col != "D" else 95
    contexto.sheet_state = "hidden"
    for hoja in libro:
        for fila in hoja:
            for celda in fila:
                celda.font = Font(name="Arial", size=11)
        if hoja.title not in ("LEEME", "_contexto"):
            for celda in hoja[1]:
                celda.fill = PatternFill("solid", fgColor="164E63")
                celda.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    resultado = BytesIO()
    libro.save(resultado)
    return resultado.getvalue()


def _leer_archivo(archivo, limite):
    nombre = getattr(archivo, "name", "plantilla.xlsx")
    if not str(nombre).lower().endswith(".xlsx"):
        raise ValidationError("Sólo se admiten archivos .xlsx.")
    if not isinstance(archivo, bytes) and hasattr(archivo, "seek"):
        archivo.seek(0)
    contenido = archivo if isinstance(archivo, bytes) else archivo.read(limite["bytes"] + 1)
    if len(contenido) > limite["bytes"]:
        raise ValidationError("El archivo supera el tamaño máximo permitido.")
    return contenido


def _verificar_zip(contenido, limite, inicio):
    try:
        with ZipFile(BytesIO(contenido)) as paquete:
            entradas = paquete.infolist()
            if len(entradas) > limite["entradas"] or sum(e.file_size for e in entradas) > limite["expandido"]:
                raise ValidationError("El contenido del archivo supera los límites permitidos.")
            nombres = [e.filename for e in entradas]
            if len(nombres) != len(set(nombres)):
                raise ValidationError("El archivo contiene entradas duplicadas.")
            for entrada in entradas:
                nombre = entrada.filename.lower()
                if entrada.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                    raise ValidationError("El archivo utiliza una compresión no admitida para .xlsx.")
                if entrada.flag_bits & 1 or any(marca in nombre for marca in ("vbaproject", "externallinks/", "embeddings/")):
                    raise ValidationError("No se admiten archivos cifrados, macros, objetos ni vínculos externos.")
                if nombre.endswith((".xml", ".rels")):
                    raiz = fromstring(paquete.read(entrada), forbid_dtd=True)
                    if nombre.endswith(".rels") and any(n.attrib.get("TargetMode", "").lower() == "external" for n in raiz):
                        raise ValidationError("No se admiten vínculos externos en la planilla.")
                    if nombre == "[content_types].xml":
                        tipos = [n.attrib.get("ContentType", "") for n in raiz]
                        if any(any(marca in tipo.lower() for marca in ("macro", "vba", "oleobject", "activex")) for tipo in tipos):
                            raise ValidationError("No se admiten macros ni objetos activos en la planilla.")
                        if "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml" not in tipos:
                            raise ValidationError("El archivo no es una planilla .xlsx admitida.")
                if monotonic() - inicio > limite["segundos"]:
                    raise ValidationError("La lectura excedió el tiempo permitido; divida la planilla.")
    except (BadZipFile, DefusedXmlException, ParseError, ZlibError, OSError, ValueError, EOFError) as exc:
        raise ValidationError("El archivo no es una planilla .xlsx válida y segura.") from exc


def _error(texto, *columnas):
    """Error de fila con los encabezados de Carga donde se origina (puede no haber)."""
    return ValidationError(texto, params={"columnas": columnas})


def _columnas(exc):
    """Encabezados señalados por el error; los errores de modelo usan el nombre del campo."""
    if hasattr(exc, "error_dict"):
        return list(exc.error_dict)
    return list(dict.fromkeys(c for e in getattr(exc, "error_list", ()) for c in (e.params or {}).get("columnas", ())))


def _rechazar(fila, exc):
    fila.update(estado="revision" if _es_duplicado(exc) or es_conflicto_cobertura(exc) else "rechazada",
                errores=_mensajes(exc), columnas_error=_columnas(exc) if isinstance(exc, ValidationError) else [])


def _texto(valor, campo, *, obligatorio=True, maximo=120, columna=None):
    columna = columna or campo
    if valor is None or valor == "":
        if obligatorio:
            raise _error(f"{campo} es obligatorio.", columna)
        return ""
    if not isinstance(valor, str):
        raise _error(f"{campo} debe estar guardado como texto para conservar sus ceros iniciales.", columna)
    valor = valor.strip()
    if (obligatorio and not valor) or len(valor) > maximo:
        raise _error(f"{campo} debe contener entre {1 if obligatorio else 0} y {maximo} caracteres.", columna)
    return valor


def _fecha(valor, columna=None):
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, str):
        try:
            return date.fromisoformat(valor.strip()).isoformat()
        except ValueError:
            pass
    raise _error("La fecha debe ser una fecha de Excel o texto AAAA-MM-DD.", *([columna] if columna else []))


def _cantidad(valor):
    try:
        if isinstance(valor, bool):
            raise ValueError
        cantidad = Decimal(str(valor))
        if not cantidad.is_finite() or cantidad != cantidad.to_integral_value() or not 0 < cantidad <= 100000:
            raise ValueError
        return int(cantidad)
    except (InvalidOperation, ValueError, TypeError):
        raise _error("La cantidad debe ser un entero entre 1 y 100.000.", "Cantidad")


def _codigo(valor, campo, *, obligatorio=True):
    codigo = _texto(valor, campo, obligatorio=obligatorio, maximo=300).split(" — ", 1)[0].strip()
    if len(codigo) > 60:
        raise _error(f"El código de {campo.lower()} no puede superar 60 caracteres.", campo)
    return codigo


def _normalizar(valores, tipo, version=1):
    documento = normalizar_documento(_texto(valores[1], "Documento", maximo=80))
    if not documento:
        raise _error("El documento debe contener letras o números.", "Documento")
    datos = {"numero": _texto(valores[0], "N.º de afiliado", maximo=80), "documento": documento}
    if tipo == "padron":
        datos.update(nombre=_texto(valores[2], "Nombre y apellido", maximo=160), plan=_codigo(valores[3], "Plan", obligatorio=False), desde=_fecha(valores[4], "Vigente desde"))
    else:
        datos.update(prestacion=_codigo(valores[2], "Prestación"), fecha=_fecha(valores[3], "Fecha de prestación"), cantidad=_cantidad(valores[4]), referencia=_texto(valores[5], "Referencia externa", obligatorio=False, maximo=120))
    if tipo == "padron" and version == 2:
        nuevos = {}
        for campo, valor in zip(CAMPOS_V2, valores[5:]):
            if valor is None or valor == "":
                continue
            if campo == "titular_documento":
                datos[campo] = normalizar_documento(_texto(valor, "Documento del titular", maximo=80, columna=campo))
                continue
            if isinstance(valor, (date, datetime)):
                valor = _fecha(valor)
            if campo in ("cuil", "cuit_empleador"):
                valor = _texto(valor, campo, maximo=11)
            nuevos[campo] = valor
        try:
            validados = padron.campos_entrada(Afiliado, [c for c in CAMPOS_V2 if c != "titular_documento"], nuevos, controles=(), requerido=False)
        except ApiValidationError:
            raise ValidationError("Los campos nuevos contienen un valor no admitido. Revisá tipos, fechas y opciones de la plantilla.")
        datos.update({k: v.isoformat() if isinstance(v, date) else v for k, v in validados.items()})
    return datos


def _mensajes(error):
    return list(error.messages) if isinstance(error, ValidationError) else ["No se pudo aplicar la fila; reintente la importación."]


def _es_duplicado(error):
    return any(getattr(e, "code", None) == "posible_duplicado" for e in getattr(error, "error_list", [error]))


def _referencias(financiador, tipo, datos=()):
    if tipo == "padron":
        return {
            "planes": {p.codigo: p for p in Plan.objects.filter(financiador=financiador, activo=True)},
            "afiliados": {normalizar_documento(a.documento): a for a in Afiliado.objects.filter(financiador=financiador, documento__in={doc for d in datos for doc in (d["documento"], d.get("titular_documento", ""))})},
            "cuil_lote": {},
        }
    return {
        "afiliados": {
            (a.numero, normalizar_documento(a.documento)): a
            for a in Afiliado.objects.filter(financiador=financiador, documento__in={d["documento"] for d in datos})
        },
        "prestaciones": {p.codigo: p for p in _catalogo(financiador)},
    }


def _validar_datos(datos, tipo, financiador, referencias, motivo=""):
    if tipo == "padron":
        plan = referencias["planes"].get(datos["plan"]) if datos["plan"] else None
        if datos["plan"] and plan is None:
            raise _error("El plan no pertenece al financiador o no está activo.", "Plan")
        desde = date.fromisoformat(datos["desde"])
        anterior = referencias["afiliados"].get(datos["documento"])
        if anterior and anterior.finalizado_en:
            raise ValidationError("La afiliación está finalizada. Reactivala expresamente desde el padrón; la importación no reactiva afiliaciones.")
        if desde > timezone.localdate():
            raise _error("La vigencia del padrón debe ser una fecha actual o anterior.", "Vigente desde")
        if anterior and desde < anterior.desde:
            raise _error("La actualización no puede anteceder la última vigencia del padrón.", "Vigente desde")
        if datos.get("_version") == 2:
            if anterior and "_revision" in datos and anterior.revision != datos["_revision"]:
                raise ValidationError("El afiliado cambió desde la vista previa; generá una nueva importación para revisar el conflicto.")
            obj = copy(anterior) if anterior else Afiliado(financiador=financiador, documento=datos["documento"], numero=datos["numero"], nombre=datos["nombre"], desde=desde, plan=plan)
            for campo in CAMPOS_V2:
                if campo in datos and campo != "titular_documento":
                    valor = datos[campo]
                    if Afiliado._meta.get_field(campo).get_internal_type() == "DateField":
                        valor = date.fromisoformat(valor)
                    setattr(obj, campo, valor)
            if datos.get("titular_documento"):
                titular = referencias["afiliados"].get(datos["titular_documento"])
                if not titular or titular.documento == obj.documento or titular.titular_id or (anterior and anterior.pk and Afiliado.objects.filter(titular_id=anterior.pk).exists()):
                    raise _error("El titular debe existir en este financiador o en una fila válida anterior, sin ciclos ni familiares de familiares.", "titular_documento")
                obj.titular_id = titular.pk or -1
            padron.validar_generales(obj)
            if obj.cuil:
                if Afiliado.objects.filter(financiador=financiador, cuil=obj.cuil).exclude(documento=obj.documento).exists() or referencias["cuil_lote"].get(obj.cuil, obj.documento) != obj.documento:
                    raise _error("El CUIL corresponde a otra identidad; requiere revisión, sin fusión automática.", "cuil")
        return {**datos, "plan": plan, "desde": desde}
    afiliado = referencias["afiliados"].get((datos["numero"], datos["documento"]))
    if afiliado is None:
        raise _error("El número y documento no corresponden a un afiliado del padrón; actualice primero el padrón.", "N.º de afiliado", "Documento")
    prestacion = referencias["prestaciones"].get(datos["prestacion"])
    if prestacion is None:
        raise _error("La prestación no pertenece al catálogo configurado del financiador.", "Prestación")
    entrada = dict(financiador=financiador, afiliado=afiliado, prestacion=prestacion, fecha=date.fromisoformat(datos["fecha"]), cantidad=datos["cantidad"], referencia=datos["referencia"], motivo_duplicado=motivo)
    validar_consumo_externo(**entrada)
    return entrada


def _resumen(filas, anterior=None):
    opcion = (anterior or {}).get("cobertura_habilitada")
    return {**({"cobertura_habilitada": opcion} if opcion is not None else {}),
            "total": len(filas), **{estado: sum(f["estado"] == estado for f in filas) for estado in ESTADOS_FILA}}


def previsualizar_importacion(*, financiador, usuario, tipo, archivo, clave, cobertura_habilitada=None):
    requerir_financiador(usuario, financiador.pk, escritura=True)
    _tipo_valido(tipo)
    validar_opcion(cobertura_habilitada)
    if tipo != "padron" and cobertura_habilitada is not None:
        raise ValidationError("La opción de cobertura corresponde sólo a importaciones de padrón.")
    try:
        clave = UUID(str(clave))
    except (ValueError, TypeError, AttributeError):
        raise ValidationError("La clave de importación debe ser un UUID válido.")
    inicio = monotonic()
    limite = limites_importacion()
    contenido = _leer_archivo(archivo, limite)
    opcion = b"" if cobertura_habilitada is None else str(cobertura_habilitada).encode()
    huella = sha256(opcion + contenido).hexdigest()
    anterior = Importacion.objects.filter(financiador=financiador, clave=clave).first()
    if anterior:
        if anterior.huella != huella or anterior.tipo != tipo:
            raise ValidationError("La clave ya pertenece a otra importación.")
        return anterior
    _verificar_zip(contenido, limite, inicio)
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationError("No se pudo leer la estructura de la planilla .xlsx.") from exc
    try:
        try:
            token = libro["_contexto"]["A1"].value
            contexto = signing.loads(token, salt=SALT)
        except (KeyError, TypeError, signing.BadSignature):
            raise ValidationError("Descargue y utilice una plantilla emitida por este sistema.")
        version = contexto.get("version") if isinstance(contexto, dict) else None
        if version not in (1, 2) or (version == 2 and tipo != "padron") or contexto != {"financiador": financiador.pk, "tipo": tipo, "version": version}:
            raise ValidationError("La plantilla corresponde a otro financiador, tipo o versión.")
        columnas = encabezados(tipo, version)
        if "Carga" not in libro:
            raise ValidationError("Falta la hoja Carga.")
        carga = libro["Carga"]
        # No confiar en dimensiones declaradas: el XML puede contener celdas fuera de ellas.
        carga.reset_dimensions()
        iterador = carga.iter_rows()
        cabecera = next(iterador, ())
        if tuple(c.value for c in cabecera) != columnas:
            raise ValidationError("Los encabezados de Carga no coinciden con la plantilla.")
        filas = []
        for numero, celdas in enumerate(iterador, 2):
            if numero > limite["filas"] + 1:
                raise ValidationError("La planilla supera la cantidad máxima de filas; divídala.")
            if monotonic() - inicio > limite["segundos"]:
                raise ValidationError("La lectura excedió el tiempo permitido; divida la planilla.")
            if all(c.value is None for c in celdas):
                continue
            valores = [c.value for c in celdas]
            fila = {"fila": numero, "datos": {}, "estado": "valida", "errores": [], "original": [v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else str(v) if v is not None else "" for v in valores[:len(columnas)]]}
            # No conservar parentesco de salud rechazado ni siquiera en los crudos.
            parentesco_invalido = False
            if version == 2:
                posicion = columnas.index("parentesco")
                if len(valores) > posicion and valores[posicion] not in (None, "", "conyuge", "hijo", "hija", "conviviente", "otro"):
                    parentesco_invalido = True
                    fila["original"][posicion] = "[omitido: parentesco no admitido]"
            try:
                # Fuera de la plantilla no hay encabezado: se señala la posición (base 1).
                if len(celdas) > len(columnas) and any(c.value is not None for c in celdas[len(columnas):]):
                    raise _error("La fila contiene columnas adicionales a la plantilla.",
                                 *(i for i, c in enumerate(celdas[len(columnas):], len(columnas) + 1) if c.value is not None))
                if any(c.data_type in ("f", "e") for c in celdas):
                    raise _error("No se admiten fórmulas ni errores de Excel en las filas de carga.",
                                 *(columnas[i] if i < len(columnas) else i + 1 for i, c in enumerate(celdas) if c.data_type in ("f", "e")))
                if parentesco_invalido:
                    raise _error("El parentesco sensible se carga únicamente por el recurso de salud; su valor fue omitido del lote.", "parentesco")
                valores += [None] * (len(columnas) - len(valores))
                fila["datos"] = _normalizar(valores, tipo, version)
                if version == 2:
                    fila["datos"]["_version"] = 2
            except ValidationError as exc:
                _rechazar(fila, exc)
            filas.append(fila)
        referencias = _referencias(financiador, tipo, [f["datos"] for f in filas if f["estado"] == "valida"])
        vistos = set()
        for fila in filas:
            if monotonic() - inicio > limite["segundos"]:
                raise ValidationError("La lectura excedió el tiempo permitido; divida la planilla.")
            if fila["estado"] != "valida":
                continue
            try:
                if version == 2:
                    documento = fila["datos"]["documento"]
                    if documento in vistos:
                        raise _error("Documento repetido dentro del lote; dejá una única fila por identidad.", "Documento")
                    anterior = referencias["afiliados"].get(documento)
                    fila["datos"]["_revision"] = anterior.revision if anterior else 0
                _validar_datos(fila["datos"], tipo, financiador, referencias)
                if tipo == "padron":
                    revisar_cobertura(fila, referencias["afiliados"].get(fila["datos"]["documento"]), cobertura_habilitada)
                if version == 2:
                    vistos.add(documento)
                    previo = referencias["afiliados"].get(documento)
                    futuro = copy(previo) if previo else Afiliado(financiador=financiador, documento=documento, pk=-fila["fila"])
                    if fila["datos"].get("titular_documento"):
                        futuro.titular_id = referencias["afiliados"][fila["datos"]["titular_documento"]].pk
                    if "cuil" in fila["datos"]:
                        futuro.cuil = fila["datos"]["cuil"]
                    referencias["afiliados"][documento] = futuro
                    if futuro.cuil:
                        referencias["cuil_lote"][futuro.cuil] = documento
            except ValidationError as exc:
                _rechazar(fila, exc)
    finally:
        libro.close()
    if not filas:
        raise ValidationError("La hoja Carga no contiene filas para importar.")
    lote, creado = Importacion.objects.get_or_create(financiador=financiador, clave=clave, defaults={"tipo": tipo, "version": version, "huella": huella, "filas": filas, "estado": "preview", "resumen": _resumen(filas, {"cobertura_habilitada": cobertura_habilitada}), "creado_por": usuario})
    if not creado and (lote.huella != huella or lote.tipo != tipo):
        raise ValidationError("La clave ya pertenece a otra importación.")
    return lote


def aplicar_importacion(*, importacion, usuario, revisiones_duplicados=None, revisiones_cobertura=None):
    if importacion.version == 3:
        from .ostel_importacion import aplicar
        return aplicar(importacion=importacion, usuario=usuario)
    """Cada fila y su efecto se confirman juntos; interrumpir permite retomar el lote."""
    requerir_financiador(usuario, importacion.financiador_id, escritura=True)
    revisiones = revisiones_duplicados or {}
    cobertura = revisiones_cobertura or {}
    if not isinstance(cobertura, dict) or any(not isinstance(v, str) or not v.strip() or len(v) > 255 for v in cobertura.values()):
        raise ValidationError("Cada conflicto de cobertura necesita un motivo de hasta 255 caracteres.")
    if not isinstance(revisiones, dict) or any(not isinstance(v, str) or not v.strip() or len(v) > 255 for v in revisiones.values()):
        raise ValidationError("Cada revisión de duplicado necesita un motivo de hasta 255 caracteres.")
    inicio, procesadas = monotonic(), 0
    importacion.refresh_from_db()
    posiciones = [
        indice for indice, fila in enumerate(importacion.filas)
        if fila["estado"] in ("valida", "error_tecnico")
        or (fila["estado"] == "revision" and (str(fila["fila"]) in revisiones or fila["fila"] in revisiones))
        or (fila["estado"] == "revision" and fila.get("conflicto_cobertura") and str(fila["fila"]) in cobertura)
    ]
    for indice in posiciones:
        try:
            with transaction.atomic():
                # El JOIN sólo carga contexto; bloquearlo también impediría
                # confirmar auditorías de ediciones concurrentes del padrón.
                lote = Importacion.objects.select_for_update(of=("self",)).select_related("financiador").get(pk=importacion.pk)
                requerir_financiador(usuario, lote.financiador_id, escritura=True)
                fila = lote.filas[indice]
                motivo = revisiones.get(str(fila["fila"]), revisiones.get(fila["fila"], "")).strip()
                motivo_cobertura = cobertura.get(str(fila["fila"]), "").strip()
                if fila["estado"] in ("aplicada", "rechazada") or (fila["estado"] == "revision" and not motivo and not motivo_cobertura):
                    continue
                if procesadas >= 100 or monotonic() - inicio >= 10:
                    break
                procesadas += 1
                try:
                    with transaction.atomic():
                        if lote.tipo == "padron":
                            Financiador.objects.select_for_update(no_key=True).get(pk=lote.financiador_id)
                            documentos = [fila["datos"]["documento"], fila["datos"].get("titular_documento", "")]
                            list(Afiliado.objects.select_for_update().filter(financiador=lote.financiador, documento__in=documentos).order_by("pk"))
                        referencias = _referencias(lote.financiador, lote.tipo, [fila["datos"]])
                        entrada = _validar_datos(fila["datos"], lote.tipo, lote.financiador, referencias, motivo)
                        if lote.tipo == "padron":
                            anterior = referencias["afiliados"].get(fila["datos"]["documento"])
                            revisar_cobertura(fila, anterior, lote.resumen.get("cobertura_habilitada"), motivo_cobertura)
                            base = {k: entrada[k] for k in ("numero", "documento", "nombre", "plan", "desde")}
                            # Sólo las altas usan la opción. Una creación concurrente falla
                            # explícitamente en vez de pisar la habilitación recién registrada.
                            base.update(solo_crear=anterior is None,
                                        cobertura_habilitada=lote.resumen.get("cobertura_habilitada") if anterior is None else None,
                                        motivo=fila.get("motivo_cobertura", "Importación de padrón"))
                            resultado = registrar_afiliado(financiador=lote.financiador, usuario=usuario, origen="importacion", lote=lote, **base)
                            if lote.version == 2:
                                nuevos = {c: fila["datos"][c] for c in CAMPOS_V2 if c in fila["datos"] and c != "titular_documento"}
                                nuevos = padron.campos_entrada(Afiliado, nuevos.keys(), nuevos, controles=(), requerido=False)
                                if fila["datos"].get("titular_documento"):
                                    nuevos["titular"] = referencias["afiliados"][fila["datos"]["titular_documento"]].pk
                                nuevos.update(revision=resultado.revision, motivo="Importación de padrón")
                                resultado = padron.editar_ficha(afiliado=resultado, usuario=usuario, entrada=nuevos, origen="importacion", lote=lote)
                        else:
                            resultado = registrar_consumo_externo(usuario=usuario, **entrada)
                    fila.update(estado="aplicada", errores=[], columnas_error=[], resultado_id=resultado.pk, motivo_duplicado=motivo, aplicado_por=autor_real(usuario).pk, aplicado_en=timezone.now().isoformat())
                except ValidationError as exc:
                    _rechazar(fila, exc)
                lote.resumen = _resumen(lote.filas, lote.resumen)
                lote.estado = "preview" if any(f["estado"] in ("valida", "revision", "error_tecnico") for f in lote.filas) else "aplicada"
                lote.save(update_fields=["filas", "resumen", "estado"])
        except PermissionDenied:
            raise
        except Exception:
            # No publicar datos del archivo ni excepciones del motor en mensajes o logs.
            with transaction.atomic():
                lote = Importacion.objects.select_for_update().get(pk=importacion.pk)
                if lote.filas[indice]["estado"] != "aplicada":
                    lote.filas[indice].update(estado="error_tecnico", errores=["No se pudo aplicar la fila; puede reintentar este lote."])
                    lote.resumen = _resumen(lote.filas, lote.resumen)
                    lote.estado = "preview"
                    lote.save(update_fields=["filas", "resumen", "estado"])
            break
    with transaction.atomic():
        lote = Importacion.objects.select_for_update().get(pk=importacion.pk)
        requerir_financiador(usuario, lote.financiador_id, escritura=True)
        lote.resumen = _resumen(lote.filas, lote.resumen)
        lote.estado = "preview" if any(f["estado"] in ("valida", "revision", "error_tecnico") for f in lote.filas) else "aplicada"
        lote.save(update_fields=["resumen", "estado"])
    importacion.refresh_from_db()
    return importacion


def descargar_rechazadas(*, importacion, usuario):
    """El caller registra la descarga en la auditoría de la API; no hay archivo público."""
    requerir_financiador(usuario, importacion.financiador_id)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Filas para corregir"
    columnas = ("NUMERO", "NROAFI", "NRODOC", "CUIL", "APENOM") if importacion.version == 3 else encabezados(importacion.tipo, importacion.version)
    hoja.append(["Fila original", *columnas, "Estado", "Detalle"])
    for fila in importacion.filas:
        if fila["estado"] not in ("rechazada", "revision", "error_tecnico"):
            continue
        valores = [fila["original"].get(c) for c in columnas] if importacion.version == 3 else fila.get("original", [])[:len(columnas)]
        valores += [""] * (len(columnas) - len(valores))
        valores = [fila["fila"], *valores, fila["estado"], "; ".join(fila["errores"])]
        numero = hoja.max_row + 1
        for col, valor in enumerate(valores, 1):
            _texto_celda(hoja.cell(numero, col), valor)
    hoja.freeze_panes = "A2"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
