"""Importación por filas: ámbito, revisión previa, seguridad y recuperación."""
from datetime import date
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier
from types import SimpleNamespace
from unittest import skipUnless
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection
from django.test import TestCase, TransactionTestCase, override_settings
from openpyxl import load_workbook

from apps.accounts.models import Usuario
from apps.registros.models import Ciudadano
from apps.simulacion import contexto
from . import importaciones
from .models import (
    Afiliado, ConsumoExterno, Financiador, Importacion, MembresiaFinanciador,
    Plan, PrestacionComun, ReglaCobertura,
)
from .services import registrar_afiliado, registrar_consumo_externo


class ImportacionesSetup:
    def setUp(self):
        self.usuario = Usuario.objects.create_user("importador@example.test", "test")
        self.financiador = Financiador.objects.create(nombre="Obra social de prueba", tipo="obra_social")
        self.membresia = MembresiaFinanciador.objects.create(financiador=self.financiador, usuario=self.usuario, rol="operador")
        self.plan = Plan.objects.create(financiador=self.financiador, codigo="BASICO", nombre="Básico")
        self.prestacion = PrestacionComun.objects.create(codigo="RX", nombre="Radiografía", categoria="imagenes")
        self.otra_prestacion = PrestacionComun.objects.create(codigo="ECO", nombre="Ecografía", categoria="ecografia")
        ReglaCobertura.objects.create(financiador=self.financiador, plan=self.plan, prestacion=self.prestacion, porcentaje=80, cupo=6, vigente_desde=date(2020, 1, 1), creado_por=self.usuario)
        self.afiliado = registrar_afiliado(financiador=self.financiador, usuario=self.usuario, numero="000123", documento="00123456", nombre="Persona de prueba", plan=self.plan, desde=date(2020, 1, 1))

    def archivo(self, tipo="consumos", filas=None):
        contenido = importaciones.generar_plantilla(financiador=self.financiador, usuario=self.usuario, tipo=tipo)
        if filas is None:
            return contenido
        libro = load_workbook(BytesIO(contenido))
        for numero, valores in enumerate(filas, 2):
            for columna, valor in enumerate(valores, 1):
                libro["Carga"].cell(numero, columna, valor)
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        return salida.getvalue()

    def consumo(self, **cambios):
        datos = dict(numero="000123", documento="00123456", prestacion="RX", fecha=date(2020, 6, 1), cantidad=1, referencia="")
        datos.update(cambios)
        return list(datos.values())

    def preview(self, archivo, tipo="consumos", clave=None):
        return importaciones.previsualizar_importacion(financiador=self.financiador, usuario=self.usuario, tipo=tipo, archivo=archivo, clave=clave or uuid4())

    def aplicar(self, lote, revisiones=None):
        return importaciones.aplicar_importacion(importacion=lote, usuario=self.usuario, revisiones_duplicados=revisiones)


class CoberturaImportacionesTests(ImportacionesSetup, TestCase):
    def archivo_cobertura(self, version, filas):
        contenido = importaciones.generar_plantilla(financiador=self.financiador, usuario=self.usuario, tipo="padron", version=version)
        libro = load_workbook(BytesIO(contenido))
        for n, fila in enumerate(filas, 2):
            for c, valor in enumerate(fila, 1):
                libro["Carga"].cell(n, c, valor)
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        return salida.getvalue()

    def previa(self, contenido, habilitada, clave=None):
        return importaciones.previsualizar_importacion(financiador=self.financiador, usuario=self.usuario,
            tipo="padron", archivo=contenido, clave=clave or uuid4(), cobertura_habilitada=habilitada)

    def test_ambas_versiones_parcial_y_revision_conservan_cobertura(self):
        for version in (1, 2):
            for solicitada in (False, True):
                with self.subTest(version=version, solicitada=solicitada):
                    self.afiliado.cobertura_habilitada = not solicitada
                    self.afiliado.save(update_fields=["cobertura_habilitada"])
                    doc = f"80000{version}{int(solicitada)}"
                    contenido = self.archivo_cobertura(version, [
                        ["000123", "00123456", "Persona de prueba", "BASICO", date(2020, 1, 1)],
                        ["NUEVO", doc, "Persona nueva ficticia", "", date(2020, 1, 1)]])
                    lote = self.previa(contenido, solicitada)
                    self.assertEqual((lote.resumen["valida"], lote.resumen["revision"]), (1, 1))
                    self.aplicar(lote)
                    self.assertEqual(lote.resumen["aplicada"], 1)
                    self.assertEqual(Afiliado.objects.get(documento=doc).cobertura_habilitada, solicitada)
                    fila = lote.filas[0]
                    lote = importaciones.aplicar_importacion(importacion=lote, usuario=self.usuario,
                        revisiones_cobertura={str(fila["fila"]): "Confirmo conservar la cobertura actual"})
                    self.assertEqual(lote.resumen["aplicada"], 2)
                    self.assertEqual(lote.resumen["cobertura_habilitada"], solicitada)
                    self.assertTrue(lote.filas[0]["motivo_cobertura"])
                    self.afiliado.refresh_from_db()
                    self.assertEqual(self.afiliado.cobertura_habilitada, not solicitada)
                    self.aplicar(lote)
                    self.assertEqual(Afiliado.objects.filter(documento=doc).count(), 1)

    def test_revision_cobertura_no_fuerza_identidad_o_foto_vieja(self):
        contenido = self.archivo_cobertura(2, [["000123", "00123456", "Persona de prueba", "BASICO", date(2020, 1, 1)]])
        lote = self.previa(contenido, False)
        self.afiliado.revision += 1
        self.afiliado.save(update_fields=["revision"])
        lote = importaciones.aplicar_importacion(importacion=lote, usuario=self.usuario,
            revisiones_cobertura={"2": "Intento con foto vieja"})
        self.assertEqual(lote.resumen["aplicada"], 0)
        self.afiliado.refresh_from_db()
        self.assertTrue(self.afiliado.cobertura_habilitada)

    def test_opcion_inmutable_y_no_aplicable_a_consumos(self):
        contenido = self.archivo_cobertura(2, [["NUEVO", "80000111", "Ficticio", "", date(2020, 1, 1)]])
        clave = uuid4()
        lote = self.previa(contenido, True, clave)
        self.assertEqual(self.previa(contenido, True, clave).pk, lote.pk)
        with self.assertRaises(ValidationError):
            self.previa(contenido, False, clave)
        with self.assertRaises(ValidationError):
            importaciones.previsualizar_importacion(financiador=self.financiador, usuario=self.usuario,
                tipo="consumos", archivo=b"", clave=uuid4(), cobertura_habilitada=True)

    def test_api_multipart_conserva_false_y_confirma_conflicto(self):
        from rest_framework.test import APIClient
        client = APIClient()
        client.force_authenticate(self.usuario)
        base = f"/api/financiadores/{self.financiador.pk}/"
        contenido = self.archivo_cobertura(2, [["000123", "00123456", "Persona de prueba", "BASICO", date(2020, 1, 1)]])
        respuesta = client.post(base + "importaciones/", {"tipo": "padron", "clave": str(uuid4()),
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "cobertura_habilitada": False}, format="multipart")
        self.assertEqual(respuesta.status_code, 201, respuesta.data)
        self.assertFalse(respuesta.data["resumen"]["cobertura_habilitada"])
        self.assertEqual(respuesta.data["resumen"]["revision"], 1)
        aplicada = client.post(base + "confirmar-importacion/", {"importacion": respuesta.data["id"],
            "revisiones_cobertura": {"2": "Conservar cobertura actual comprobada"}}, format="json")
        self.assertEqual(aplicada.status_code, 200, aplicada.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 1)
        self.afiliado.refresh_from_db()
        self.assertTrue(self.afiliado.cobertura_habilitada)


class ImportacionesTests(ImportacionesSetup, TestCase):
    def test_preview_no_aplica_y_confirmacion_importa_validas_con_ceros(self):
        lote = self.preview(self.archivo(filas=[self.consumo(), self.consumo(documento="DESCONOCIDO")]))
        self.assertEqual(lote.resumen["valida"], 1)
        self.assertEqual(lote.resumen["rechazada"], 1)
        self.assertEqual(lote.filas[0]["datos"]["numero"], "000123")
        self.assertFalse(ConsumoExterno.objects.exists())
        self.aplicar(lote)
        self.assertEqual(lote.estado, "aplicada")
        self.assertEqual(lote.resumen["aplicada"], 1)
        self.assertEqual(ConsumoExterno.objects.get().afiliado, self.afiliado)
        self.assertFalse(Ciudadano.objects.exists())
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 1)

    def test_fila_aplicada_atribuye_al_superusuario_que_simula(self):
        lote = self.preview(self.archivo(filas=[self.consumo()]))
        root = Usuario.objects.create_superuser("root-importacion@example.test", "x")
        contexto.activar(SimpleNamespace(cuenta=self.usuario), root)
        try:
            self.aplicar(lote)
        finally:
            contexto.limpiar()
        lote.refresh_from_db()
        self.assertEqual(lote.filas[0]["aplicado_por"], root.pk)

    def test_mismo_lote_reintenta_y_clave_no_admite_otro_contenido(self):
        archivo = self.archivo(filas=[self.consumo()])
        clave = uuid4()
        lote = self.preview(archivo, clave=clave)
        self.assertEqual(self.preview(archivo, clave=clave).pk, lote.pk)
        with self.assertRaisesMessage(ValidationError, "otra importación"):
            self.preview(self.archivo(filas=[self.consumo(cantidad=2)]), clave=clave)
        self.assertEqual(Importacion.objects.count(), 1)

    def test_plantilla_personalizada_no_permite_codigo_ajeno(self):
        libro = load_workbook(BytesIO(self.archivo()), read_only=True)
        self.assertEqual(list(libro["Prestaciones"].values)[1][0], "RX")
        self.assertEqual(libro["Prestaciones"].max_row, 2)
        libro.close()
        lote = self.preview(self.archivo(filas=[self.consumo(prestacion="ECO")]))
        self.assertEqual(lote.resumen["rechazada"], 1)
        self.aplicar(lote)
        self.assertFalse(ConsumoExterno.objects.exists())

    def test_plantilla_de_otro_financiador_se_rechaza_aun_con_permiso(self):
        otro = Financiador.objects.create(nombre="Mutual de prueba", tipo="mutual")
        MembresiaFinanciador.objects.create(financiador=otro, usuario=self.usuario, rol="operador")
        with self.assertRaisesMessage(ValidationError, "otro financiador"):
            importaciones.previsualizar_importacion(financiador=otro, usuario=self.usuario, tipo="consumos", archivo=self.archivo(filas=[self.consumo()]), clave=uuid4())

    def test_auditor_no_importa_y_revocacion_impide_confirmar(self):
        lote = self.preview(self.archivo(filas=[self.consumo()]))
        self.membresia.rol = "auditor"
        self.membresia.save()
        with self.assertRaises(PermissionDenied):
            self.aplicar(lote)
        with self.assertRaises(PermissionDenied):
            self.preview(self.archivo(filas=[self.consumo()]))
        self.assertFalse(ConsumoExterno.objects.exists())

    def test_padron_actualiza_sin_bajas_ni_ciudadanos(self):
        omitido = registrar_afiliado(financiador=self.financiador, usuario=self.usuario, numero="FAMILIA", documento="000002", nombre="Otra persona", plan=self.plan, desde=date(2020, 1, 1))
        filas = [["000999", "00123456", "Persona actualizada", "BASICO", date(2020, 2, 1)], ["FAMILIA", "000003", "Persona nueva", "BASICO", date(2020, 1, 1)]]
        lote = self.preview(self.archivo("padron", filas), "padron")
        self.assertEqual(lote.resumen["valida"], 2)
        self.aplicar(lote)
        self.afiliado.refresh_from_db()
        self.assertEqual(self.afiliado.numero, "000999")
        self.assertTrue(Afiliado.objects.filter(pk=omitido.pk).exists())
        self.assertEqual(Afiliado.objects.count(), 3)
        self.assertFalse(Ciudadano.objects.exists())

    def test_padron_sin_plan_y_no_acepta_fecha_anterior_a_ultima_vigencia(self):
        filas = [["000123", "00123456", "Persona de prueba", "", date(2020, 2, 1)], ["000123", "00123456", "Persona de prueba", "BASICO", date(2019, 1, 1)]]
        lote = self.preview(self.archivo("padron", filas), "padron")
        self.assertEqual(lote.resumen["valida"], 1)
        self.assertEqual(lote.resumen["rechazada"], 1)
        self.aplicar(lote)
        self.afiliado.refresh_from_db()
        self.assertIsNone(self.afiliado.plan_id)

    def test_revalida_plan_desactivado_despues_del_resumen(self):
        fila = ["NUEVO", "000003", "Persona nueva", "BASICO", date(2020, 1, 1)]
        lote = self.preview(self.archivo("padron", [fila]), "padron")
        self.plan.activo = False
        self.plan.save()
        self.aplicar(lote)
        self.assertEqual(lote.resumen["rechazada"], 1)
        self.assertFalse(Afiliado.objects.filter(documento="000003").exists())

    def test_no_adivina_ceros_perdidos_ni_admite_cantidades_fraccionarias(self):
        lote = self.preview(self.archivo(filas=[self.consumo(numero=123), self.consumo(cantidad=1.5)]))
        self.assertEqual(lote.resumen["rechazada"], 2)
        self.assertIn("texto", lote.filas[0]["errores"][0])
        self.assertIn("entero", lote.filas[1]["errores"][0])

    def test_documento_con_separadores_conserva_identidad_y_ceros(self):
        lote = self.preview(self.archivo(filas=[self.consumo(documento="00.123.456")]))
        self.assertEqual(lote.filas[0]["datos"]["documento"], "00123456")
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.get().afiliado_id, self.afiliado.pk)
        fila = ["000123", "00.123.456", "Misma persona", "BASICO", date(2020, 2, 1)]
        padron = self.preview(self.archivo("padron", [fila]), "padron")
        self.aplicar(padron)
        self.assertEqual(Afiliado.objects.count(), 1)

    def test_posible_duplicado_exige_revision_con_motivo(self):
        registrar_consumo_externo(financiador=self.financiador, usuario=self.usuario, afiliado=self.afiliado, prestacion=self.prestacion, fecha=date(2020, 6, 1), cantidad=1)
        lote = self.preview(self.archivo(filas=[self.consumo()]))
        self.assertEqual(lote.resumen["revision"], 1)
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 1)
        self.aplicar(lote, {"2": "Son dos prácticas distintas, verificadas con el prestador."})
        self.assertEqual(ConsumoExterno.objects.count(), 2)
        self.assertEqual(lote.estado, "aplicada")
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 2)

    def test_duplicados_dentro_del_lote_se_revalidan_al_aplicar(self):
        lote = self.preview(self.archivo(filas=[self.consumo(), self.consumo()]))
        self.aplicar(lote)
        self.assertEqual(lote.resumen["aplicada"], 1)
        self.assertEqual(lote.resumen["revision"], 1)
        self.assertEqual(ConsumoExterno.objects.count(), 1)

    def test_referencia_repetida_no_duplica_y_payload_distinto_se_rechaza(self):
        filas = [self.consumo(referencia="ORIGEN-1"), self.consumo(referencia="ORIGEN-1"), self.consumo(referencia="ORIGEN-1", cantidad=2)]
        lote = self.preview(self.archivo(filas=filas))
        self.aplicar(lote)
        self.assertEqual(lote.resumen["aplicada"], 2)
        self.assertEqual(lote.resumen["rechazada"], 1)
        self.assertEqual(ConsumoExterno.objects.count(), 1)
        self.assertEqual(lote.filas[0]["resultado_id"], lote.filas[1]["resultado_id"])

    def test_error_tecnico_se_retoma_sin_repetir_filas_aplicadas(self):
        lote = self.preview(self.archivo(filas=[self.consumo(referencia="A"), self.consumo(referencia="B")]))
        original = importaciones.registrar_consumo_externo
        invocaciones = []

        def falla_segunda(**kwargs):
            invocaciones.append(kwargs["referencia"])
            if kwargs["referencia"] == "B":
                raise RuntimeError("Error técnico simulado")
            return original(**kwargs)

        with patch.object(importaciones, "registrar_consumo_externo", side_effect=falla_segunda):
            self.aplicar(lote)
        self.assertEqual(lote.resumen["aplicada"], 1)
        self.assertEqual(lote.resumen["error_tecnico"], 1)
        self.assertEqual(ConsumoExterno.objects.count(), 1)
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 2)
        self.assertEqual(lote.estado, "aplicada")

    def test_aplicacion_acotada_retoma_las_filas_restantes(self):
        filas = [self.consumo(referencia=f"ORIGEN-{i}") for i in range(101)]
        lote = self.preview(self.archivo(filas=filas))
        self.aplicar(lote)
        self.assertGreater(lote.resumen["aplicada"], 0)
        self.assertLessEqual(lote.resumen["aplicada"], 100)
        self.assertGreater(lote.resumen["valida"], 0)
        self.assertEqual(lote.estado, "preview")
        while lote.resumen["valida"]:
            anterior = lote.resumen["aplicada"]
            self.aplicar(lote)
            self.assertGreater(lote.resumen["aplicada"], anterior)
        self.assertEqual(ConsumoExterno.objects.count(), 101)
        self.assertEqual(lote.estado, "aplicada")

    def test_fallo_despues_de_crear_consumo_revierte_efecto_y_permite_reintento(self):
        lote = self.preview(self.archivo(filas=[self.consumo()]))
        original = importaciones.registrar_consumo_externo

        def falla_despues(**kwargs):
            original(**kwargs)
            raise RuntimeError("Corte antes de confirmar la fila")

        with patch.object(importaciones, "registrar_consumo_externo", side_effect=falla_despues):
            self.aplicar(lote)
        self.assertFalse(ConsumoExterno.objects.exists())
        self.assertEqual(lote.resumen["error_tecnico"], 1)
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 1)

    def test_permiso_revalidado_por_fila_preserva_progreso(self):
        lote = self.preview(self.archivo(filas=[self.consumo(referencia="A"), self.consumo(referencia="B")]))
        original = importaciones.requerir_financiador
        invocaciones = []

        def revocacion(*args, **kwargs):
            invocaciones.append(1)
            if len(invocaciones) == 3:
                raise PermissionDenied("Permiso revocado")
            return original(*args, **kwargs)

        with patch.object(importaciones, "requerir_financiador", side_effect=revocacion):
            with self.assertRaises(PermissionDenied):
                self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 1)
        self.aplicar(lote)
        self.assertEqual(ConsumoExterno.objects.count(), 2)

    def test_formulas_se_rechazan_y_exportan_como_texto_literal(self):
        lote = self.preview(self.archivo(filas=[self.consumo(documento="=1+1")]))
        self.assertEqual(lote.resumen["rechazada"], 1)
        contenido = importaciones.descargar_rechazadas(importacion=lote, usuario=self.usuario)
        libro = load_workbook(BytesIO(contenido), data_only=False)
        self.assertEqual(libro.active["C2"].value, "=1+1")
        self.assertEqual(libro.active["C2"].data_type, "s")
        self.assertEqual(libro.active["E2"].value, "2020-06-01")
        libro.close()

    def test_formato_encabezados_y_limites_se_rechazan(self):
        with self.assertRaisesMessage(ValidationError, ".xlsx"):
            self.preview(SimpleUploadedFile("datos.xls", b"not-xlsx"))
        with self.assertRaises(ValidationError):
            self.preview(b"not-xlsx")
        archivo = self.archivo(filas=[self.consumo()])
        with override_settings(FINANCIADORES_XLSX_MAX_BYTES=100):
            with self.assertRaisesMessage(ValidationError, "tamaño"):
                self.preview(archivo)
        with override_settings(FINANCIADORES_XLSX_MAX_EXPANDIDO=100):
            with self.assertRaisesMessage(ValidationError, "límites"):
                self.preview(archivo)
        libro = load_workbook(BytesIO(archivo))
        libro["Carga"]["A1"] = "Número cambiado"
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        with self.assertRaisesMessage(ValidationError, "encabezados"):
            self.preview(salida.getvalue())

    def test_dtd_y_vinculos_externos_se_rechazan_antes_de_openpyxl(self):
        archivo = self.archivo(filas=[self.consumo()])
        for malicioso in (
            b'<!DOCTYPE data [<!ENTITY x "boom">]><data>&x;</data>',
            b'<Relationships><Relationship TargetMode="External" Target="https://example.test"/></Relationships>',
        ):
            salida = BytesIO()
            with ZipFile(BytesIO(archivo)) as origen, ZipFile(salida, "w", ZIP_DEFLATED) as destino:
                for entrada in origen.infolist():
                    destino.writestr(entrada.filename, origen.read(entrada))
                destino.writestr("malicioso.rels", malicioso)
            with self.assertRaises(ValidationError):
                self.preview(salida.getvalue())

    def test_macros_y_compresion_no_admitida_se_rechazan(self):
        archivo = self.archivo(filas=[self.consumo()])
        for compresion, macro in ((ZIP_DEFLATED, True), (ZIP_BZIP2, False)):
            salida = BytesIO()
            with ZipFile(BytesIO(archivo)) as origen, ZipFile(salida, "w", compresion) as destino:
                for entrada in origen.infolist():
                    destino.writestr(entrada.filename, origen.read(entrada))
                if macro:
                    destino.writestr("xl/vbaProject.bin", b"macro")
            with self.assertRaises(ValidationError):
                self.preview(salida.getvalue())

    @override_settings(FINANCIADORES_XLSX_MAX_FILAS=2)
    def test_limite_filas_y_columnas_no_confia_en_dimensiones_xml(self):
        archivo = self.archivo(filas=[self.consumo(), self.consumo(), self.consumo()])
        with self.assertRaisesMessage(ValidationError, "máxima de filas"):
            self.preview(archivo)
        lote = self.preview(self.archivo(filas=[self.consumo() + ["columna inesperada"]]))
        self.assertEqual(lote.resumen["rechazada"], 1)


@skipUnless(connection.vendor == "postgresql", "El bloqueo de filas exige PostgreSQL.")
class ImportacionesConcurrenciaTests(ImportacionesSetup, TransactionTestCase):
    def test_dos_confirmaciones_del_mismo_lote_generan_un_consumo(self):
        lote = self.preview(self.archivo(filas=[self.consumo()]))
        barrera = Barrier(2)

        def confirmar():
            close_old_connections()
            try:
                copia = Importacion.objects.get(pk=lote.pk)
                usuario = Usuario.objects.get(pk=self.usuario.pk)
                barrera.wait(timeout=10)
                resultado = importaciones.aplicar_importacion(importacion=copia, usuario=usuario)
                return resultado.resumen["aplicada"]
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as procesos:
            tareas = [procesos.submit(confirmar), procesos.submit(confirmar)]
            self.assertEqual([tarea.result(timeout=20) for tarea in tareas], [1, 1])
        self.assertEqual(ConsumoExterno.objects.count(), 1)


class SubsanarImportacionTests(ImportacionesSetup, TestCase):
    """Revisar fila en la plantilla del sistema: corrige, excluye y revalida el lote."""

    def lote(self, filas, version=1):
        contenido = importaciones.generar_plantilla(financiador=self.financiador, usuario=self.usuario, tipo="padron", version=version)
        libro = load_workbook(BytesIO(contenido))
        for n, fila in enumerate(filas, 2):
            for c, valor in enumerate(fila, 1):
                libro["Carga"].cell(n, c, valor)
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        return importaciones.previsualizar_importacion(financiador=self.financiador, usuario=self.usuario,
                                                       tipo="padron", archivo=salida.getvalue(), clave=uuid4())

    def subsanar(self, lote, fila, valores=None, accion="corregir", revision=None):
        actual = next(f for f in lote.filas if f["fila"] == fila)
        return importaciones.subsanar_importacion(importacion=lote, usuario=self.usuario, fila=fila,
            revision=revision or actual.get("revision", 1), valores=valores, motivo="Verificado con la fuente", accion=accion)

    def fila(self, lote, numero):
        return next(f for f in lote.filas if f["fila"] == numero)

    def test_corrige_una_fila_rechazada_y_queda_en_el_historial(self):
        lote = self.lote([["000200", "00200000", "Persona ficticia", "NOEXISTE", date(2020, 1, 1)]])
        self.assertEqual(self.fila(lote, 2)["columnas_error"], ["Plan"])
        valores = ["000200", "00200000", "Persona ficticia", "BASICO", "2020-01-01"]
        lote = self.subsanar(lote, 2, valores)
        fila = self.fila(lote, 2)
        self.assertEqual((fila["estado"], fila["revision"], fila["original"]), ("valida", 2, valores))
        self.assertEqual(fila["correcciones"][0]["anterior"][3], "NOEXISTE")
        self.assertEqual(lote.resumen["valida"], 1)
        self.aplicar(lote)
        self.assertEqual(Afiliado.objects.get(documento="00200000").plan, self.plan)

    def test_correccion_invalida_sigue_rechazada_con_su_columna(self):
        lote = self.lote([["000200", "00200000", "Persona ficticia", "", date(2020, 1, 1)]])
        lote = self.subsanar(lote, 2, ["000200", "00200000", "Persona ficticia", "", "no es fecha"])
        fila = self.fila(lote, 2)
        self.assertEqual((fila["estado"], fila["columnas_error"], fila["datos"]), ("rechazada", ["Vigente desde"], {}))
        with self.assertRaises(ValidationError):
            self.subsanar(lote, 2, ["000200", "00200000", "Persona ficticia", "", "2020-01-01"], revision=1)

    def test_excluir_no_se_aplica_y_cierra_el_lote(self):
        lote = self.lote([["000200", "00200000", "Persona ficticia", "NOEXISTE", date(2020, 1, 1)],
                          ["000201", "00200001", "Persona ficticia", "", date(2020, 1, 1)]])
        lote = self.subsanar(lote, 2, accion="omitir")
        self.assertEqual((lote.resumen["omitida"], lote.resumen["valida"]), (1, 1))
        lote = self.aplicar(lote)
        self.assertEqual(lote.estado, "aplicada")
        self.assertFalse(Afiliado.objects.filter(documento="00200000").exists())
        with self.assertRaises(ValidationError):
            self.subsanar(lote, 2, ["000200", "00200000", "Persona ficticia", "", "2020-01-01"])

    def test_revalida_entre_filas_sin_releer_las_demas(self):
        repetida = ["000300", "00300000", "Persona ficticia", "", date(2020, 1, 1)]
        lote = self.lote([repetida, ["000301", "00300000", "Otra ficticia", "", date(2020, 1, 1)],
                          [300301, "00300002", "Número sin ceros", "", date(2020, 1, 1)]], version=2)
        self.assertEqual(self.fila(lote, 3)["errores"], ["Documento repetido dentro del lote; dejá una única fila por identidad."])
        self.assertEqual(self.fila(lote, 4)["estado"], "rechazada")
        lote = self.subsanar(lote, 2, accion="omitir")
        # Al excluir la primera, la repetida deja de serlo; la del número sigue rechazada.
        self.assertEqual(self.fila(lote, 3)["estado"], "valida")
        self.assertEqual(self.fila(lote, 4)["estado"], "rechazada")

    def test_valores_incompletos_o_lote_ostel_no_se_corrigen(self):
        lote = self.lote([["000200", "00200000", "Persona ficticia", "NOEXISTE", date(2020, 1, 1)]])
        with self.assertRaises(ValidationError):
            self.subsanar(lote, 2, ["000200"])
        Importacion.objects.filter(pk=lote.pk).update(version=3)
        with self.assertRaises(ValidationError):
            self.subsanar(lote, 2, accion="omitir")
