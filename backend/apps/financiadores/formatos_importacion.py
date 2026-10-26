"""Registro de generadores de plantillas; no homologa ni aplica movimientos."""
from io import BytesIO

from django.core.exceptions import ValidationError


def _sistema(*, financiador, usuario, tipo, version):
    from .importaciones import generar_plantilla
    contenido = generar_plantilla(financiador=financiador, usuario=usuario, tipo=tipo, version=version)
    return contenido, "plantilla-padron" if tipo == "padron" and version == 2 else "plantilla"


def _ostel(**contexto):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from .ostel_prevalidacion import COLUMNAS

    libro = Workbook()
    try:
        hoja = libro.active
        hoja.title = "Hoja1"
        # 40 columnas: CUILTIT es obligatorio desde la entrega del 26/10/2026.
        hoja.append(COLUMNAS)
        hoja.freeze_panes = "A2"
        for numero, celda in enumerate(hoja[1], 1):
            celda.font = Font(bold=True)
            columna = hoja.column_dimensions[get_column_letter(numero)]
            columna.width = 18
            columna.number_format = "@"
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue(), "plantilla-ostel"
    finally:
        libro.close()


# Para otro formato: registrar sus tipos de carga y su generador.
# El importador del sistema conserva su contrato y sus controles.
FORMATOS_PLANTILLA = {
    "sistema": {"tipos": ("padron", "consumos"), "generar": _sistema},
    "ostel": {"tipos": ("padron",), "generar": _ostel},
}


def generar_plantilla_formato(formato, *, financiador, usuario, tipo, version=1):
    definicion = FORMATOS_PLANTILLA.get(formato)
    if definicion is None or tipo not in definicion["tipos"]:
        raise ValidationError("El formato no está disponible para este tipo de carga.")
    return definicion["generar"](financiador=financiador, usuario=usuario, tipo=tipo, version=version)
