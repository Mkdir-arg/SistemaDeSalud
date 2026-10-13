"""Escenario ficticio issue 130, exclusivo de PostgreSQL demo/padron130.

Nunca borra ni actualiza un escenario anterior. Group marca una carga completa;
las instantáneas iniciales identifican sus afiliados aunque luego se editen.
"""
import json
import os
from datetime import timedelta
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from django.utils import timezone

from apps.accounts.models import Usuario
from apps.registros.models import Ciudadano
from apps.financiadores import models as m
from apps.financiadores.padron import editar_protegido, validar_generales
from apps.financiadores.services import auditar


ESCENARIO = "padron130-v1"
ORIGEN = "demo_padron130"
# Lock propio de este loader; no necesita filas previamente existentes.
LOCK_ESCENARIO = 130013001
CASOS_A = (
    "titular", "conyuge", "menor", "estudiante", "integrante_5",
    "monotributo", "pasivo", "desempleo", "continuidad", "legado",
    "finalizado", "futuro",
)
CUENTAS = {
    # clave de manifiesto: (organización, rol, salud, cobranza)
    "admin": ("A", "admin", True, True),
    "operador": ("A", "operador", False, False),
    "salud": ("A", "operador", True, False),
    "cobranza": ("A", "operador", False, True),
    "auditor": ("A", "auditor", True, True),
    "orgB": ("B", "admin", False, False),
}


def exigir_destino():
    """Antes de consultas y archivos; desarrollo tampoco es un destino válido."""
    if settings.ENTORNO != "demo":
        raise CommandError("seed_padron_demo exige ENTORNO=demo. No se modificó ningún dato.")
    if connection.vendor != "postgresql" or connection.settings_dict["NAME"] != "padron130":
        raise CommandError("seed_padron_demo exige PostgreSQL con NAME exacto padron130.")


def clave_externa():
    # No usar clave_demo(): ese helper permite una contraseña pública por defecto.
    clave = os.environ.get("DEMO_PASSWORD")
    if not clave or not clave.strip() or clave.strip().casefold() == "demo1234":
        raise CommandError("Provisioná DEMO_PASSWORD externa; se prohíbe la clave demo por defecto.")
    return clave


def identificador_fiscal(documento, prefijos=(20, 27, 23)):
    """CUIL/CUIT sintético con verificador local; sin consultar organismos."""
    for prefijo in prefijos:
        cuerpo = f"{prefijo}{documento:08d}"
        dv = (11 - sum(int(n) * p for n, p in zip(cuerpo, (5, 4, 3, 2, 7, 6, 5, 4, 3, 2))) % 11) % 11
        if dv != 10:
            return cuerpo + str(dv)
    raise CommandError("No se pudo generar el identificador ficticio.")


def cbu_ficticio(indice):
    """Banco ficticio 000; sólo verifica dígitos, no existencia de una cuenta."""
    def bloque(cuerpo, pesos):
        dv = (10 - sum(int(n) * p for n, p in zip(cuerpo, pesos)) % 10) % 10
        return cuerpo + str(dv)

    return bloque("0000130", (7, 1, 3, 9, 7, 1, 3)) + bloque(
        f"{130000 + indice:013d}", (3, 9, 7, 1, 3, 9, 7, 1, 3, 9, 7, 1, 3),
    )


def ruta_manifiesto(valor, tipo="manifiesto"):
    if not valor:
        return None
    ruta = Path(valor).expanduser().resolve()
    codigo = Path(__file__).resolve()
    raiz = next((p for p in codigo.parents if (p / ".git").exists()), Path(settings.BASE_DIR).resolve())
    if ruta == raiz or raiz in ruta.parents:
        raise CommandError(f"El {tipo} debe guardarse fuera del repositorio/código de la aplicación.")
    if ruta.exists() or not ruta.parent.is_dir():
        raise CommandError(f"Elegí un archivo de {tipo} nuevo en un directorio existente.")
    return ruta


class Command(BaseCommand):
    help = "Carga única ficticia de padrón issue130; exige demo/padron130 vacía y clave externa."

    def add_arguments(self, parser):
        parser.add_argument("--manifiesto", help="JSON sin secretos, archivo nuevo fuera del repositorio.")
        parser.add_argument("--muestra-xlsx", help="Muestra sobre plantilla v2 firmada en runtime, archivo nuevo fuera del repositorio.")

    def handle(self, *args, **options):
        exigir_destino()
        ruta = ruta_manifiesto(options.get("manifiesto"))
        ruta_muestra = ruta_manifiesto(options.get("muestra_xlsx"), "XLSX")
        if ruta and ruta == ruta_muestra:
            raise CommandError("El manifiesto y la muestra XLSX requieren rutas distintas.")
        try:
            with transaction.atomic():
                if connection.vendor == "postgresql":
                    with connection.cursor() as cursor:
                        cursor.execute("SELECT pg_advisory_xact_lock(%s)", [LOCK_ESCENARIO])
                marcador = Group.objects.filter(name=ESCENARIO).first()
                existente = marcador is not None
                if not existente:
                    self._exigir_vacia()
                    self._crear(clave_externa())
                    marcador = Group.objects.create(name=ESCENARIO)
                    marcador.user_set.add(*Usuario.objects.filter(email__in=self._emails()))
                manifiesto = self._manifiesto(marcador)
                muestra = self._muestra(manifiesto) if ruta_muestra else None
        except ValidationError:
            # No interpolar campos/valores protegidos ni la excepción original.
            raise CommandError("La carga ficticia no cumple el contrato del padrón; se revirtió completa.") from None
        if ruta or ruta_muestra:
            try:
                if ruta:
                    with ruta.open("x", encoding="utf-8") as archivo:
                        json.dump(manifiesto, archivo, ensure_ascii=False, indent=2)
                        archivo.write("\n")
                if ruta_muestra:
                    with ruta_muestra.open("xb") as archivo:
                        archivo.write(muestra)
            except OSError as exc:
                raise CommandError(
                    "La carga está completa en la base, pero no se pudieron guardar todos los archivos. "
                    "Repetí con otra ruta nueva: no se resembrará ni se resetearán claves."
                ) from exc
        # La salida tiene sólo IDs, logins y nombres ficticios; nunca salud/bancos/claves.
        self.stdout.write(json.dumps({"existente": existente, **manifiesto}, ensure_ascii=False, indent=2))

    @staticmethod
    def _muestra(manifiesto):
        # Se reutiliza la función del servidor y su SECRET_KEY actual, sin token estático.
        from openpyxl import load_workbook
        from apps.financiadores.importaciones import generar_plantilla

        org = m.Financiador.objects.get(pk=manifiesto["financiadores"]["A"]["id"])
        admin = Usuario.objects.get(pk=manifiesto["usuarios"]["admin"]["id"])
        libro = load_workbook(BytesIO(generar_plantilla(financiador=org, usuario=admin, tipo="padron", version=2)))
        hoja = libro["Carga"]
        columnas = [c.value for c in hoja[1]]
        titular = m.Afiliado.objects.get(pk=manifiesto["afiliados"]["titular"]["id"])
        finalizado = m.Afiliado.objects.get(pk=manifiesto["afiliados"]["finalizado"]["id"])
        filas = [
            {"N.º de afiliado": "P130-IMPORT-01", "Documento": "99013101", "Nombre y apellido": "Titular Importación Ficticia",
             "tipo_documento": "DNI", "cuil": identificador_fiscal(99013101), "regimen": "adherente"},
            {"N.º de afiliado": "P130-IMPORT-02", "Documento": "99013102", "Nombre y apellido": "Familiar Importación Ficticia",
             "tipo_documento": "DNI", "cuil": identificador_fiscal(99013102), "titular_documento": "99013101", "parentesco": "hijo"},
            {"N.º de afiliado": "P130-IMPORT-03", "Documento": "99013103", "Nombre y apellido": "Verificador Inválido Ficticio",
             "tipo_documento": "DNI", "cuil": "00000000001"},
            {"N.º de afiliado": "P130-IMPORT-04", "Documento": "99013104", "Nombre y apellido": "Conflicto Identidad Ficticio",
             "tipo_documento": "OTRO", "cuil": titular.cuil},
            {"N.º de afiliado": "P130-IMPORT-05", "Documento": "99013105", "Nombre y apellido": "Familiar Sin Titular Ficticio",
             "titular_documento": "99013888", "parentesco": "hijo"},
            {"N.º de afiliado": finalizado.numero, "Documento": finalizado.documento, "Nombre y apellido": finalizado.nombre},
        ]
        hoy = timezone.localdate().isoformat()
        for numero, datos in enumerate(filas, 2):
            datos.update({"Plan": "P130-A-BASE", "Vigente desde": hoy})
            for columna, nombre in enumerate(columnas, 1):
                celda = hoja.cell(numero, columna, datos.get(nombre, ""))
                celda.data_type, celda.number_format = "s", "@"
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    @staticmethod
    def _emails():
        return [f"{cuenta.lower()}@padron130.example.test" for cuenta in CUENTAS]

    @staticmethod
    def _exigir_vacia():
        if any(modelo.objects.exists() for modelo in (m.Financiador, m.Afiliado, Usuario, Ciudadano)):
            raise CommandError("Primera carga: financiadores y personas deben estar vacíos. No se borra ni mezcla información.")
        if m.HistorialAfiliacion.objects.filter(origen=ORIGEN).exists():
            raise CommandError("El escenario está incompleto; no se repara ni se resiembra automáticamente.")

    def _crear(self, password):
        hoy = timezone.localdate()
        organizaciones = {
            letra: m.Financiador.objects.create(nombre=f"Padrón 130 — Organización {letra} ficticia", tipo="obra_social")
            for letra in ("A", "B")
        }
        usuarios = {}
        for cuenta, (letra, rol, salud, cobranza) in CUENTAS.items():
            usuario = Usuario.objects.create_user(
                f"{cuenta.lower()}@padron130.example.test", password,
                nombre=f"Demo {cuenta}", apellido="Padrón 130",
            )
            usuarios[cuenta] = usuario
            m.MembresiaFinanciador.objects.create(
                financiador=organizaciones[letra], usuario=usuario, rol=rol,
                acceso_padron_salud=salud, acceso_padron_cobranza=cobranza,
            )
        # Concesiones explícitas para este escenario, atribuidas al admin A ficticio.
        for cuenta, (letra, _, salud, cobranza) in CUENTAS.items():
            for bloque, concedido in (("salud", salud), ("cobranza", cobranza)):
                if concedido:
                    auditar(usuarios["admin"], "designar_padron_" + bloque, usuarios[cuenta].pk,
                            financiador=organizaciones[letra], motivo="Concesión explícita del escenario ficticio issue130.")
        planes = {
            (letra, tipo): m.Plan.objects.create(financiador=org, codigo=f"P130-{letra}-{tipo.upper()}", nombre=f"Demo {tipo}", tipo=tipo)
            for letra, org in organizaciones.items()
            for tipo in (("base", "superador", "privado") if letra == "A" else ("base",))
        }
        afiliados = {}
        for indice, caso in enumerate(CASOS_A, 1):
            afiliados[caso] = self._afiliado(organizaciones["A"], planes, usuarios["admin"], hoy, indice, caso, afiliados)
        # Mismo documento sintético en otra organización: identidades separadas.
        self._afiliado(organizaciones["B"], planes, usuarios["orgB"], hoy, 1, "orgB", {})
        integrante = afiliados["integrante_5"]
        editar_protegido(afiliado=integrante, usuario=usuarios["admin"], bloque="salud", entrada={
            "revision": 1, "motivo": "Carga ficticia protegida issue130.",
            "cud": True, "cud_numero": "FICTICIO-130", "cud_presentado_en": hoy - timedelta(days=60),
            "cud_vencimiento": hoy + timedelta(days=45), "cud_referencia": "Constancia ficticia sin adjunto",
            "dbt": True, "cronicas": "Condición crónica ficticia para ensayo; sin historia clínica.",
            "parentesco_sensible": "discapacidad",
        })
        for indice, caso in enumerate(("titular", "monotributo", "continuidad"), 1):
            editar_protegido(afiliado=afiliados[caso], usuario=usuarios["admin"], bloque="cobranza", entrada={
                "revision": 1, "motivo": "Carga ficticia protegida issue130.",
                "situacion": "deuda" if caso == "monotributo" else "al_dia",
                "fuente": "Declaración administrativa ficticia; sin conciliación bancaria",
                "informado_en": hoy, "modalidad": "debito_cbu", "cbu": cbu_ficticio(indice),
                "banco": "Banco ficticio 000 (sin integración)",
                "cuit_titular": identificador_fiscal(99001300, (30, 33, 34)),
                "titular_cuenta": "Cuenta ficticia de tercero Demo 130", "dia_vencimiento": 15,
            })

    @staticmethod
    def _afiliado(org, planes, admin, hoy, indice, caso, previos):
        letra = "B" if caso == "orgB" else "A"
        tipo_plan = "superador" if caso in ("monotributo", "pasivo") else "privado" if caso in ("continuidad", "futuro") else "base"
        desde = hoy + timedelta(days=30) if caso == "futuro" else hoy - timedelta(days=180)
        datos = dict(financiador=org, numero=f"P130-{letra}-{indice:02d}", documento=str(99013000 + indice),
                     nombre=f"Persona {indice:02d} FamiliaDemo{indice:02d} {letra}", plan=planes[(letra, tipo_plan)], desde=desde)
        if caso != "legado":
            edad = {"titular": 56, "conyuge": 52, "menor": 14, "estudiante": 22, "integrante_5": 28, "pasivo": 68}.get(caso, 35)
            regimen = caso if caso in ("monotributo", "pasivo", "desempleo") else "adherente" if caso in ("continuidad", "futuro") else "dependencia"
            datos.update(tipo_documento="DNI", cuil=identificador_fiscal(99013000 + indice),
                         nombres=f"Persona {indice:02d}", apellidos=f"FamiliaDemo{indice:02d} {letra}", sexo="X",
                         fecha_nacimiento=hoy.replace(year=hoy.year - edad, day=min(hoy.day, 28)),
                         nacionalidad="Argentina", calle="Calle Ficticia", altura=str(130 + indice),
                         codigo_postal="0000", localidad="Localidad Demo", partido="Partido Demo", provincia="Provincia Demo",
                         email=f"afiliado-{letra.lower()}-{indice:02d}@padron130.example.test", regimen=regimen,
                         alta_administrativa=hoy - timedelta(days=180), delegacion="Delegación ficticia 130",
                         credencial_tipo="Externa informada (ficticia)", credencial_fecha=hoy - timedelta(days=20))
            if regimen == "dependencia":
                datos.update(empleador="Empleador Ficticio 130", actividad_laboral="Actividad demo",
                             cuit_empleador=identificador_fiscal(99001300, (30, 33, 34)), domicilio_laboral="Domicilio laboral ficticio")
            elif regimen == "monotributo":
                datos["actividad_laboral"] = "Actividad independiente ficticia"
            elif regimen == "desempleo":
                datos.update(subsidio_desde=hoy - timedelta(days=90), subsidio_hasta=hoy + timedelta(days=60),
                             solicitud_baja=hoy, baja_efectiva=hoy + timedelta(days=30),
                             fin_prestacional_informado=hoy + timedelta(days=60), causa_baja="cese_laboral",
                             motivo_baja="Baja futura ficticia informada; sin corte automático")
            if caso in ("conyuge", "menor", "estudiante", "integrante_5"):
                datos.update(titular=previos["titular"], parentesco="conyuge" if caso == "conyuge" else "hijo")
            if caso == "continuidad":
                datos.update(solicitud_baja=hoy - timedelta(days=20), baja_efectiva=hoy - timedelta(days=10),
                             fin_prestacional_informado=hoy + timedelta(days=50), causa_baja="renuncia",
                             motivo_baja="Baja administrativa ficticia; derechos aún vigentes")
        if caso == "finalizado":
            datos.update(finalizado_en=timezone.now() - timedelta(days=10), finalizado_por=admin,
                         motivo_finalizacion="Finalización prestacional ficticia explícita")
        afiliado = m.Afiliado(**datos)
        afiliado.full_clean()
        afiliado.save()
        if caso != "legado":
            tramite = caso == "estudiante"
            documento = m.DocumentoLegajo(
                afiliado=afiliado, tipo="estudio" if tramite else "dni",
                estado="presentado" if tramite else "verificado",
                presentado_en=hoy - timedelta(days=30), verificado_en=None if tramite else hoy - timedelta(days=20),
                vence_en=hoy - timedelta(days=1) if tramite else hoy + timedelta(days=45),
                referencia="Presentación ficticia registrada; sin archivo adjunto",
            )
            documento.full_clean()
            documento.save()
            afiliado.estado_tramite = "en_revision" if tramite else "completo"
        validar_generales(afiliado)
        afiliado.save(update_fields=["estado_tramite"])
        m.HistorialAfiliacion.objects.create(
            afiliado=afiliado, numero=afiliado.numero, documento=afiliado.documento, plan=afiliado.plan,
            desde=desde, registrado_por=admin, origen=ORIGEN, motivo="Carga inicial ficticia issue130.",
        )
        auditar(admin, "seed_padron130", afiliado.pk, financiador=org)
        return afiliado

    @staticmethod
    def _manifiesto(marcador):
        historias = list(m.HistorialAfiliacion.objects.filter(origen=ORIGEN).select_related("afiliado__financiador", "plan").order_by("id"))
        usuarios = {u.email: u for u in marcador.user_set.all()}
        esperados = {f"P130-A-{i:02d}": caso for i, caso in enumerate(CASOS_A, 1)}
        esperados["P130-B-01"] = "orgB"
        if len(historias) != 13 or {h.numero for h in historias} != set(esperados) or set(usuarios) != set(Command._emails()):
            raise CommandError("El marcador del escenario está incompleto; no se resiembra ni repara automáticamente.")
        org_ids = {h.afiliado.financiador_id for h in historias}
        if (len(org_ids) != 2 or m.Financiador.objects.exclude(pk__in=org_ids).exists()
                or Usuario.objects.exclude(pk__in=[u.pk for u in usuarios.values()]).exists()
                or Ciudadano.objects.exists()):
            raise CommandError("Se detectaron financiadores/personas ajenos al escenario; no se modifica ningún dato.")
        organizaciones, afiliados, planes = {}, {}, {}
        for h in historias:
            letra = h.numero.split("-")[1]
            org = h.afiliado.financiador
            if letra in organizaciones and organizaciones[letra]["id"] != org.pk:
                raise CommandError("El escenario cambió de organización; no se resiembra automáticamente.")
            organizaciones[letra] = {"id": org.pk, "nombre": org.nombre}
            afiliados[esperados[h.numero]] = {"id": h.afiliado_id, "numero": h.afiliado.numero, "nombre": h.afiliado.nombre, "financiador_id": org.pk}
            if h.plan:
                planes[h.plan.codigo] = {"id": h.plan_id, "financiador_id": h.plan.financiador_id}
        cuentas = {}
        for cuenta, email in zip(CUENTAS, Command._emails()):
            usuario = usuarios[email]
            letra = CUENTAS[cuenta][0]
            membresia = m.MembresiaFinanciador.objects.filter(usuario=usuario, financiador_id=organizaciones[letra]["id"]).first()
            if not membresia:
                raise CommandError("Falta una membresía del escenario; no se resiembra automáticamente.")
            cuentas[cuenta] = {"id": usuario.pk, "email": email, "nombre": usuario.nombre, "rol": membresia.rol}
        return {
            "escenario": ESCENARIO, "version": 1,
            "fecha_carga": timezone.localtime(historias[0].registrado).date().isoformat(),
            "financiadores": organizaciones, "planes": planes,
            "usuarios": cuentas,
            "afiliados": afiliados,
        }
