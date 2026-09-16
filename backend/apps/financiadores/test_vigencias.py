"""Vigencias explícitas, acceso histórico acotado y trazabilidad hospitalaria."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from django.test import override_settings
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import Membresia
from apps.auditoria.models import AccesoClinico
from apps.casos.models import Caso
from apps.finanzas.cobros import capturar_cobros_atencion
from apps.finanzas.dinero import (
    decidir_ajuste, previsualizar_reintegro, reducir_obligacion, registrar_movimiento,
    reintegrar_movimiento,
)
from apps.finanzas.models import ObligacionFinanciera
from apps.finanzas.models_cobros import SnapshotCobroAtencion
from apps.instituciones.models import Institucion
from apps.registros.models import Ciudadano

from . import importaciones
from .cobertura import liberar
from .cobros import resolver_saldo
from .models import (
    Afiliado, AfiliacionCaso, ArancelConvenio, ConsumoExterno, Convenio,
    DistribucionCobro, Financiador, MembresiaFinanciador, Plan,
    ReservaCobertura,
)
from .test_cobertura import CoberturaSetup


class VigenciasApiSetup(CoberturaSetup):
    def setUp(self):
        super().setUp()
        self.membresia.rol = "admin"
        self.membresia.save(update_fields=["rol"])
        self.base = f"/api/financiadores/{self.financiador.pk}/"
        self.client.force_authenticate(self.operador)

    def post(self, accion, datos):
        return self.client.post(self.base + accion + "/", datos, format="json")

    def finalizar(self):
        response = self.post("finalizar-afiliacion", {
            "afiliado": self.afiliado.pk, "motivo": "Baja informada por la persona afiliada",
        })
        self.assertEqual(response.status_code, 200, response.data)
        return response

    def cerrar(self):
        response = self.post("cerrar-convenio", {
            "convenio": self.convenio.pk, "motivo": "Finalización del acuerdo institucional",
        })
        self.assertEqual(response.status_code, 200, response.data)
        return response

    def nuevo_caso(self):
        return Caso.objects.create(
            institucion=self.institucion, version=self.caso.version,
            ciudadano=self.paciente, area_actual=self.area,
        )

    def seleccionar_en_nuevo_caso(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post("/api/coberturas/afiliacion/", {
            "caso": self.nuevo_caso().pk, "afiliado": self.afiliado.pk,
            "motivo": "Verificación al ingresar",
        }, format="json")
        self.client.force_authenticate(self.operador)
        return response

    def cobrar(self, obligacion, **datos):
        return registrar_movimiento(
            obligacion=obligacion, importe=obligacion.importe_original,
            fecha=self.hoy, clave=uuid4(), usuario=self.admin, **datos,
        )

    def actividad(self, **filtros):
        response = self.client.get(self.base + "actividad/", filtros)
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def realizada(self, acepta=False):
        if acepta:
            self.reservar(acepta=True)
        hecho = self.atencion()
        return ReservaCobertura.objects.get(hecho=hecho)


class VigenciasAfiliacionTests(VigenciasApiSetup, APITestCase):
    def test_finalizar_exige_motivo_y_conserva_la_seleccion_del_caso(self):
        response = self.post("finalizar-afiliacion", {"afiliado": self.afiliado.pk, "motivo": " "})
        self.assertEqual(response.status_code, 400)
        response = self.finalizar()
        self.assertFalse(response.data["vigente"])
        self.afiliado.refresh_from_db()
        self.assertIsNotNone(self.afiliado.finalizado_en)
        self.assertEqual(self.afiliado.finalizado_por_id, self.operador.pk)
        self.assertEqual(self.afiliado.motivo_finalizacion, "Baja informada por la persona afiliada")
        self.seleccion.refresh_from_db()
        self.assertEqual((self.seleccion.afiliado_id, self.seleccion.plan_id), (self.afiliado.pk, self.plan.pk))
        self.assertEqual(AfiliacionCaso.objects.filter(caso=self.caso).count(), 1)
        self.assertEqual(self.evaluar()["importe_financiador"], "80.00")
        self.assertEqual(self.seleccionar_en_nuevo_caso().status_code, 400)

    def test_padron_manual_no_reactiva_una_afiliacion_finalizada(self):
        self.finalizar()
        response = self.post("padron", {
            "numero": self.afiliado.numero, "documento": self.afiliado.documento,
            "nombre": self.afiliado.nombre, "plan": self.plan.pk, "desde": str(self.hoy),
        })
        self.assertEqual(response.status_code, 400)
        self.afiliado.refresh_from_db()
        self.assertIsNotNone(self.afiliado.finalizado_en)
        self.assertEqual(Afiliado.objects.count(), 1)

    def test_excel_rechaza_reactivacion_implicita(self):
        self.finalizar()
        plantilla = importaciones.generar_plantilla(
            financiador=self.financiador, usuario=self.operador, tipo="padron",
        )
        libro = load_workbook(BytesIO(plantilla))
        for columna, valor in enumerate([
            self.afiliado.numero, self.afiliado.documento, self.afiliado.nombre,
            self.plan.codigo, self.hoy,
        ], 1):
            libro["Carga"].cell(2, columna, valor)
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        lote = importaciones.previsualizar_importacion(
            financiador=self.financiador, usuario=self.operador, tipo="padron",
            archivo=salida.getvalue(), clave=uuid4(),
        )
        importaciones.aplicar_importacion(importacion=lote, usuario=self.operador)
        self.afiliado.refresh_from_db()
        self.assertIsNotNone(self.afiliado.finalizado_en)
        self.assertEqual(lote.resumen.get("aplicada", 0), 0)
        self.assertEqual(lote.resumen["rechazada"], 1)

    def test_reactivar_conserva_identidad_consumos_y_cupo(self):
        externo = self.externo(cantidad=5)
        self.finalizar()
        response = self.post("reactivar-afiliacion", {
            "afiliado": self.afiliado.pk, "plan": self.plan.pk,
            "motivo": "Reingreso confirmado por padrón",
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["vigente"])
        self.assertEqual(response.data["id"], self.afiliado.pk)
        self.assertEqual(Afiliado.objects.count(), 1)
        self.assertTrue(ConsumoExterno.objects.filter(pk=externo.pk, afiliado=self.afiliado).exists())
        self.assertEqual(self.evaluar(cantidad=2)["cubiertas"], 1)
        self.assertEqual(self.seleccionar_en_nuevo_caso().status_code, 201)

    def test_reactivar_admite_ausencia_explicita_de_plan(self):
        self.finalizar()
        response = self.post("reactivar-afiliacion", {
            "afiliado": self.afiliado.pk, "plan": None, "motivo": "Reingreso sin plan específico",
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["vigente"])
        self.assertIsNone(response.data["plan"])

    def test_operador_finaliza_y_reactiva_pero_auditor_no(self):
        self.membresia.rol = "auditor"
        self.membresia.save(update_fields=["rol"])
        for accion in ("finalizar-afiliacion", "reactivar-afiliacion"):
            with self.subTest(accion=accion):
                response = self.post(accion, {"afiliado": self.afiliado.pk, "plan": self.plan.pk, "motivo": "Cambio solicitado"})
                self.assertEqual(response.status_code, 403)
        self.membresia.rol = "operador"
        self.membresia.save(update_fields=["rol"])
        self.finalizar()
        self.assertEqual(self.post("reactivar-afiliacion", {
            "afiliado": self.afiliado.pk, "plan": self.plan.pk, "motivo": "Reingreso",
        }).status_code, 200)

    def test_no_modifica_afiliacion_de_otro_financiador(self):
        otra = Financiador.objects.create(nombre="Otra mutual", tipo="mutual")
        ajeno = Afiliado.objects.create(financiador=otra, numero="1", documento="99999", nombre="Otra persona", desde=self.hoy)
        for accion in ("finalizar-afiliacion", "reactivar-afiliacion"):
            with self.subTest(accion=accion):
                response = self.post(accion, {"afiliado": ajeno.pk, "plan": None, "motivo": "Intento ajeno"})
                self.assertIn(response.status_code, (400, 404))
        ajeno.refresh_from_db()
        self.assertIsNone(ajeno.finalizado_en)

    def test_plan_inactivo_no_admite_nuevas_afiliaciones_y_preserva_el_caso(self):
        response = self.post("editar-plan", {
            "plan": self.plan.pk, "nombre": "Plan conservado", "activo": False,
            "motivo": "Cierre a nuevas altas", "codigo": "NO-DEBE-CAMBIAR",
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.plan.refresh_from_db()
        self.assertEqual((self.plan.nombre, self.plan.codigo, self.plan.activo), ("Plan conservado", "BASE", False))
        self.seleccion.refresh_from_db()
        self.assertEqual(self.seleccion.plan_id, self.plan.pk)
        self.assertEqual(self.evaluar()["importe_financiador"], "80.00")
        self.assertEqual(self.seleccionar_en_nuevo_caso().status_code, 400)
        response = self.post("padron", {"numero": "2", "documento": "22222", "nombre": "Nueva persona", "plan": self.plan.pk, "desde": str(self.hoy)})
        self.assertEqual(response.status_code, 400)
        self.finalizar()
        response = self.post("reactivar-afiliacion", {"afiliado": self.afiliado.pk, "plan": self.plan.pk, "motivo": "Reingreso"})
        self.assertEqual(response.status_code, 400)

    def test_editar_plan_requiere_administrador_ambito_y_motivo(self):
        datos = {"plan": self.plan.pk, "nombre": "Renombrado", "activo": True, "motivo": "Cambio comercial"}
        for rol in ("operador", "auditor"):
            self.membresia.rol = rol
            self.membresia.save(update_fields=["rol"])
            self.assertEqual(self.post("editar-plan", datos).status_code, 403)
        self.membresia.rol = "admin"
        self.membresia.save(update_fields=["rol"])
        self.assertEqual(self.post("editar-plan", {**datos, "motivo": ""}).status_code, 400)
        otra = Financiador.objects.create(nombre="Otra mutual", tipo="mutual")
        plan = Plan.objects.create(financiador=otra, codigo="OTRO", nombre="Ajeno")
        self.assertIn(self.post("editar-plan", {**datos, "plan": plan.pk}).status_code, (400, 404))
        plan.refresh_from_db()
        self.assertEqual(plan.nombre, "Ajeno")


class VigenciasConvenioTests(VigenciasApiSetup, APITestCase):
    def test_cierre_preserva_arancel_reserva_y_no_admite_reapertura_por_aceptacion(self):
        arancel = ArancelConvenio.objects.create(
            convenio=self.convenio, prestacion=self.prestacion, importe=Decimal("150"),
            vigente_desde=self.hoy, creado_por=self.admin,
        )
        reserva = self.reservar(acepta=True)
        evaluacion = reserva.evaluacion.copy()
        self.cerrar()
        self.convenio.refresh_from_db()
        self.assertEqual(self.convenio.estado, "finalizado")
        self.assertIsNotNone(self.convenio.cerrado_en)
        self.assertEqual(self.convenio.cerrado_por_id, self.operador.pk)
        self.assertTrue(self.convenio.motivo_cierre)
        reserva.refresh_from_db()
        self.assertEqual(reserva.evaluacion, evaluacion)
        self.assertTrue(ArancelConvenio.objects.filter(pk=arancel.pk, convenio=self.convenio).exists())
        self.assertEqual(self.post("aceptar-convenio", {"convenio": self.convenio.pk}).status_code, 400)
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.post("/api/coberturas/aceptar-convenio/", {"convenio": self.convenio.pk}).status_code, 400)

    def test_nuevo_convenio_preserva_el_anterior_y_rechaza_dos_vigentes(self):
        self.assertEqual(self.post("convenios", {"institucion": self.institucion.pk}).status_code, 400)
        self.cerrar()
        nuevo = self.post("convenios", {"institucion": self.institucion.pk})
        self.assertEqual(nuevo.status_code, 201, nuevo.data)
        self.assertNotEqual(nuevo.data["id"], self.convenio.pk)
        self.assertEqual(self.post("convenios", {"institucion": self.institucion.pk}).status_code, 400)
        self.assertEqual(Convenio.objects.filter(financiador=self.financiador, institucion=self.institucion).count(), 2)

    def test_rechazo_solo_para_propuesta_de_la_contraparte(self):
        self.cerrar()
        propuesta = self.post("convenios", {"institucion": self.institucion.pk}).data
        datos = {"convenio": propuesta["id"], "motivo": "Condiciones sin acuerdo"}
        self.assertEqual(self.post("rechazar-convenio", datos).status_code, 400)
        self.assertEqual(self.post("cerrar-convenio", datos).status_code, 400)
        self.conceder("configurar_cobros")
        self.client.force_authenticate(self.usuario)
        response = self.client.post("/api/coberturas/rechazar-convenio/", datos, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["estado"], "rechazado")
        self.assertEqual(self.client.post("/api/coberturas/aceptar-convenio/", {"convenio": propuesta["id"]}).status_code, 400)
        self.client.force_authenticate(self.operador)
        self.assertEqual(self.post("convenios", {"institucion": self.institucion.pk}).status_code, 201)

    def test_financiador_rechaza_propuesta_hospitalaria_y_no_un_convenio_activo(self):
        self.assertEqual(self.post("rechazar-convenio", {"convenio": self.convenio.pk, "motivo": "Rechazo tardío"}).status_code, 400)
        self.cerrar()
        propuesta = Convenio.objects.create(financiador=self.financiador, institucion=self.institucion, propuesto_por="hospital", creado_por=self.admin)
        response = self.post("rechazar-convenio", {"convenio": propuesta.pk, "motivo": "No se acordaron las condiciones"})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["estado"], "rechazado")

    def test_cierre_hospitalario_requiere_permiso_en_el_hospital_del_convenio(self):
        datos = {"convenio": self.convenio.pk, "motivo": "Fin del acuerdo"}
        self.client.force_authenticate(self.usuario)
        self.assertEqual(self.client.post("/api/coberturas/cerrar-convenio/", datos).status_code, 403)
        self.conceder("configurar_cobros")
        otra = Institucion.objects.create(nombre="Hospital ajeno")
        ajeno = Convenio.objects.create(financiador=self.financiador, institucion=otra, estado="activo", propuesto_por="plataforma", creado_por=self.admin)
        self.assertEqual(self.client.post("/api/coberturas/cerrar-convenio/", {**datos, "convenio": ajeno.pk}).status_code, 403)
        response = self.client.post("/api/coberturas/cerrar-convenio/", datos)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["estado"], "finalizado")

    def test_cierre_requiere_motivo_y_administrador_del_financiador(self):
        datos = {"convenio": self.convenio.pk, "motivo": ""}
        self.assertEqual(self.post("cerrar-convenio", datos).status_code, 400)
        for rol in ("operador", "auditor"):
            self.membresia.rol = rol
            self.membresia.save(update_fields=["rol"])
            self.assertEqual(self.post("cerrar-convenio", {**datos, "motivo": "Cierre"}).status_code, 403)
        self.convenio.refresh_from_db()
        self.assertEqual(self.convenio.estado, "activo")

    def test_hospital_no_acepta_ni_rechaza_su_propia_propuesta(self):
        self.cerrar()
        propuesta = Convenio.objects.create(financiador=self.financiador, institucion=self.institucion, propuesto_por="hospital", creado_por=self.admin)
        self.conceder("configurar_cobros")
        self.client.force_authenticate(self.usuario)
        for accion in ("aceptar-convenio", "rechazar-convenio"):
            with self.subTest(accion=accion):
                response = self.client.post(f"/api/coberturas/{accion}/", {"convenio": propuesta.pk, "motivo": "Decisión unilateral"})
                self.assertEqual(response.status_code, 400)
        propuesta.refresh_from_db()
        self.assertEqual(propuesta.estado, "propuesto")

    def test_financiador_no_puede_cerrar_convenio_ajeno(self):
        otra = Financiador.objects.create(nombre="Otra mutual", tipo="mutual")
        ajeno = Convenio.objects.create(financiador=otra, institucion=self.institucion, estado="activo", propuesto_por="plataforma", creado_por=self.admin)
        response = self.post("cerrar-convenio", {"convenio": ajeno.pk, "motivo": "Cierre ajeno"})
        self.assertIn(response.status_code, (400, 404))
        ajeno.refresh_from_db()
        self.assertEqual(ajeno.estado, "activo")


class ActividadHistoricaTests(VigenciasApiSetup, APITestCase):
    def test_actividad_vigente_expone_solo_datos_administrativos(self):
        self.reservar()
        actividad = self.actividad()
        self.assertEqual(actividad["count"], 1)
        self.assertEqual(actividad["results"][0]["acceso"], "vigente")
        self.assertEqual(set(actividad["results"][0]), {
            "id", "fecha", "hospital", "prestacion", "numero", "documento",
            "cantidad", "cubiertas", "estado", "discrepancia", "importe_financiador",
            "estado_cobro", "acceso", "nombre", "plan", "codigo",
            "importe_asignado", "importe_acuerdos",
        })

    def test_reserva_cubierta_pendiente_sigue_visible_hasta_liberarla(self):
        reserva = self.reservar()
        self.cerrar()
        actividad = self.actividad()
        self.assertEqual(actividad["count"], 1)
        self.assertEqual(actividad["results"][0]["acceso"], "pendiente_historico")
        liberar(reserva=reserva, usuario=self.admin, motivo="No se realizó la consulta", no_realizada=True)
        self.assertEqual(self.actividad()["count"], 0)

    def test_reserva_no_cubierta_no_habilita_acceso_tras_baja(self):
        self.externo(cantidad=6)
        reserva = self.reservar()
        self.assertEqual(reserva.cubiertas, 0)
        self.finalizar()
        self.assertEqual(self.actividad()["count"], 0)

    def test_distribucion_resuelta_no_equivale_a_cobro_del_financiador(self):
        reserva = self.realizada(acepta=True)
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.assertEqual(distribucion.estado, "resuelta")
        self.cerrar()
        self.assertEqual(self.actividad()["count"], 1)
        self.cobrar(distribucion.obligacion_financiador)
        self.assertEqual(self.actividad()["count"], 0)

    def test_evaluacion_propia_pendiente_no_se_oculta_al_cerrar_convenio(self):
        self.politica(importe=None)
        reserva = self.realizada()
        self.assertEqual(DistribucionCobro.objects.get(reserva=reserva).estado, "arancel_pendiente")
        self.cerrar()
        actividad = self.actividad()
        self.assertEqual(actividad["count"], 1)
        self.assertEqual(actividad["results"][0]["acceso"], "pendiente_historico")

    def test_copago_paciente_pendiente_no_extiende_acceso_del_financiador_pagado(self):
        reserva = self.realizada()
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.assertEqual(distribucion.estado, "pendiente")
        self.cobrar(distribucion.obligacion_financiador)
        self.finalizar()
        self.assertEqual(self.actividad()["count"], 0)
        self.assertEqual(DistribucionCobro.objects.get(pk=distribucion.pk).estado, "pendiente")

    def test_resolucion_posterior_a_cargo_del_financiador_conserva_acceso_hasta_cobrarse(self):
        reserva = self.realizada()
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.cobrar(distribucion.obligacion_financiador)
        resolucion = resolver_saldo(
            reserva=reserva, usuario=self.admin, decision="financiador", importe=Decimal("20"),
            motivo="La obra social asume la diferencia", evidencia="Acuerdo firmado", clave=uuid4(),
        )
        self.finalizar()
        self.assertEqual(self.actividad()["count"], 1)
        self.cobrar(resolucion.obligacion)
        self.assertEqual(self.actividad()["count"], 0)

    def test_reintegro_pendiente_conserva_acceso_aunque_el_cobro_este_completo(self):
        reserva = self.realizada(acepta=True)
        obligacion = DistribucionCobro.objects.get(reserva=reserva).obligacion_financiador
        movimiento = self.cobrar(obligacion)
        datos = dict(original=movimiento, importe=Decimal("10"), fecha=self.hoy,
                     motivo="Reintegro a la obra social", efecto="reducir", usuario=self.admin, aprobado=False)
        preview = previsualizar_reintegro(**datos)
        reintegrar_movimiento(**datos, clave=uuid4(), version_esperada=preview["version_esperada"])
        self.cerrar()
        self.assertEqual(self.actividad()["count"], 1)

    def test_reduccion_pendiente_y_saldo_a_devolver_conservan_acceso(self):
        reserva = self.realizada(acepta=True)
        obligacion = DistribucionCobro.objects.get(reserva=reserva).obligacion_financiador
        self.cobrar(obligacion)
        ajuste = reducir_obligacion(obligacion=obligacion, importe=Decimal("10"), motivo="Revisión acordada", clave=uuid4(), usuario=self.admin, aprobado=False)
        self.cerrar()
        self.assertEqual(self.actividad()["count"], 1)
        decidir_ajuste(obligacion=obligacion, ajuste=ajuste, usuario=self.admin, aprobar=True)
        self.assertEqual(self.actividad()["count"], 1)

    def test_resumen_no_revela_discrepancias_fuera_del_alcance_de_actividad(self):
        reserva = self.reservar()
        liberar(reserva=reserva, usuario=self.admin, motivo="Cancelada antes de realizarse", no_realizada=True)
        ReservaCobertura.objects.filter(pk=reserva.pk).update(discrepancia=True)
        self.assertEqual(self.client.get(self.base + "resumen/").data["discrepancias"], 1)
        self.cerrar()
        self.assertEqual(self.actividad()["count"], 0)
        self.assertEqual(self.client.get(self.base + "resumen/").data["discrepancias"], 0)

    def test_otro_financiador_no_obtiene_actividad_ni_auditoria_del_primero(self):
        self.reservar()
        otra = Financiador.objects.create(nombre="Mutual ajena", tipo="mutual")
        MembresiaFinanciador.objects.create(financiador=otra, usuario=self.operador, rol="auditor")
        response = self.client.get(f"/api/financiadores/{otra.pk}/actividad/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 0)
        self.assertFalse(AccesoClinico.objects.filter(tipo="financiador").exists())

    def test_auditoria_registra_paciente_hospital_y_usuario_solo_para_la_pagina_visible(self):
        primera = self.reservar()
        otro_caso, prestacion = self.otro_hospital()
        segunda = self.reservar(caso=otro_caso, prestacion=prestacion)
        actividad = self.actividad(page_size=1)
        self.assertEqual(actividad["count"], 2)
        self.assertEqual(actividad["results"][0]["id"], segunda.pk)
        acceso = AccesoClinico.objects.get(tipo="financiador", recurso="financiadores-actividad")
        self.assertEqual((acceso.usuario_id, acceso.ciudadano_id, acceso.institucion_id),
                         (self.operador.pk, otro_caso.ciudadano_id, otro_caso.institucion_id))
        self.assertFalse(AccesoClinico.objects.filter(tipo="financiador", ciudadano=self.paciente).exists())
        segunda_pagina = self.actividad(page_size=1, page=2)
        self.assertEqual(segunda_pagina["results"][0]["id"], primera.pk)
        self.assertEqual(AccesoClinico.objects.filter(tipo="financiador").count(), 2)
        Membresia.objects.create(usuario=self.usuario, institucion=self.institucion, rol="admin")
        self.client.force_authenticate(self.usuario)
        response = self.client.get("/api/accesos-clinicos/", {"tipo": "financiador"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["count"], 1)
        self.assertEqual(response.data["results"][0]["ciudadano"], self.paciente.pk)
        self.assertEqual(self.client.get("/api/accesos-clinicos/", {"tipo": "financiador", "institucion": otro_caso.institucion_id}).data["count"], 0)

    def test_auditoria_utiliza_paciente_original_del_hecho(self):
        self.realizada()
        nuevo = Ciudadano.objects.create(institucion=self.institucion, nombre="Paciente corregido", documento="99999")
        self.caso.ciudadano = nuevo
        self.caso.save(update_fields=["ciudadano"])
        self.actividad()
        acceso = AccesoClinico.objects.get(tipo="financiador", recurso="financiadores-actividad")
        self.assertEqual(acceso.ciudadano_id, self.paciente.pk)

    @override_settings(DEBUG=False)
    def test_fallo_de_auditoria_impide_entregar_actividad_al_financiador(self):
        self.reservar()
        self.client.raise_request_exception = False
        with patch("apps.auditoria.mixins.AccesoClinico.objects.bulk_create", side_effect=RuntimeError("Auditoría no disponible")):
            response = self.client.get(self.base + "actividad/")
        self.assertEqual(response.status_code, 500)
        self.assertNotIn(self.afiliado.documento.encode(), response.content)
        self.assertNotIn(self.afiliado.nombre.encode(), response.content)
        self.assertFalse(AccesoClinico.objects.filter(tipo="financiador").exists())


class CierreYCapturaHistoricaTests(VigenciasApiSetup, APITestCase):
    def comprobar_revaluacion_al_realizar(self):
        reserva = self.reservar(acepta=True)
        evaluacion_anterior = reserva.evaluacion.copy()
        aceptacion_anterior = reserva.aceptacion.copy()
        self.assertEqual(evaluacion_anterior["convenio"], self.convenio.pk)
        self.assertEqual(aceptacion_anterior["importe"], "20.00")
        self.cerrar()
        reserva.refresh_from_db()
        self.assertEqual(reserva.evaluacion, evaluacion_anterior)
        self.assertEqual(reserva.aceptacion, aceptacion_anterior)
        hecho = self.atencion()
        reserva.refresh_from_db()
        self.assertEqual(reserva.estado, "realizada")
        self.assertEqual(reserva.hecho_id, hecho.pk)
        self.assertEqual(reserva.evaluacion["evaluacion_confirmada"], evaluacion_anterior)
        self.assertIsNone(reserva.evaluacion["convenio"])
        self.assertEqual(reserva.evaluacion["estado"], "no_cubierta")
        self.assertEqual(reserva.evaluacion["importe_paciente"], "100.00")
        self.assertEqual(reserva.cubiertas, 0)
        self.assertEqual(reserva.aceptacion, {})
        self.assertTrue(reserva.discrepancia)
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.assertEqual(distribucion.estado, "pendiente")
        self.assertEqual(distribucion.importe_paciente, Decimal("100"))
        self.assertIsNone(distribucion.obligacion_paciente_id)
        self.assertIsNone(distribucion.obligacion_financiador_id)
        self.assertFalse(ObligacionFinanciera.objects.filter(hecho=hecho).exists())
        self.assertTrue(SnapshotCobroAtencion.objects.get(hecho=hecho).capturado)

    def test_cierre_conserva_reserva_pero_revalida_y_quita_aceptacion_al_realizar(self):
        self.comprobar_revaluacion_al_realizar()

    def test_cambio_de_convenio_revoca_aceptacion_incluso_sin_regla_particular(self):
        self.regla.delete()
        self.convenio.porcentaje_default = Decimal("80")
        self.convenio.save(update_fields=["porcentaje_default"])
        self.assertIsNone(self.evaluar()["regla"])
        self.comprobar_revaluacion_al_realizar()

    def test_nueva_evaluacion_sin_convenio_no_genera_deuda_paciente_sin_aceptacion(self):
        self.cerrar()
        evaluacion = self.evaluar()
        self.assertEqual(evaluacion["estado"], "no_cubierta")
        self.assertEqual(evaluacion["importe_financiador"], "0.00")
        self.assertEqual(evaluacion["importe_paciente"], "100.00")
        self.assertIsNone(evaluacion["convenio"])
        reserva = self.reservar()
        hecho = self.atencion()
        reserva.refresh_from_db()
        self.assertEqual(reserva.hecho_id, hecho.pk)
        self.assertEqual(reserva.aceptacion, {})
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.assertEqual(distribucion.estado, "pendiente")
        self.assertEqual(distribucion.importe_paciente, Decimal("100"))
        self.assertFalse(ObligacionFinanciera.objects.filter(hecho=hecho).exists())

    def test_cierre_posterior_no_cambia_recuperacion_del_hecho_ni_arancel_historico(self):
        arancel = ArancelConvenio.objects.create(
            convenio=self.convenio, prestacion=self.prestacion, importe=Decimal("150"),
            vigente_desde=self.hoy, creado_por=self.admin,
        )
        reserva = self.reservar(acepta=True)
        evaluacion_anterior = reserva.evaluacion.copy()
        aceptacion_anterior = reserva.aceptacion.copy()
        with patch("apps.financiadores.cobros.capturar_cobertura", side_effect=RuntimeError("Fallo transitorio de captura")):
            hecho = self.atencion()
        self.assertFalse(SnapshotCobroAtencion.objects.get(hecho=hecho).capturado)
        self.cerrar()
        nuevo = self.post("convenios", {"institucion": self.institucion.pk})
        self.assertEqual(nuevo.status_code, 201, nuevo.data)
        self.client.force_authenticate(self.admin)
        aceptado = self.client.post("/api/coberturas/aceptar-convenio/", {"convenio": nuevo.data["id"]})
        self.assertEqual(aceptado.status_code, 200, aceptado.data)
        ArancelConvenio.objects.create(
            convenio_id=nuevo.data["id"], prestacion=self.prestacion, importe=Decimal("999"),
            vigente_desde=self.hoy, creado_por=self.admin,
        )
        capturar_cobros_atencion(hecho.pk)
        reserva.refresh_from_db()
        self.assertEqual(reserva.estado, "realizada")
        self.assertEqual(reserva.evaluacion, evaluacion_anterior)
        self.assertEqual(reserva.aceptacion, aceptacion_anterior)
        self.assertEqual(reserva.evaluacion["convenio"], self.convenio.pk)
        self.assertEqual(reserva.evaluacion["excepcion"], arancel.pk)
        self.assertFalse(reserva.discrepancia)
        distribucion = DistribucionCobro.objects.get(reserva=reserva)
        self.assertEqual(distribucion.obligacion_financiador.importe_original, Decimal("120"))
        self.assertEqual(distribucion.obligacion_paciente.importe_original, Decimal("30"))
        self.assertTrue(SnapshotCobroAtencion.objects.get(hecho=hecho).capturado)
