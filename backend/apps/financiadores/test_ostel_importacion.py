"""Circuito administrativo con datos inventados y DB de pruebas aislada."""
from uuid import uuid4
from io import BytesIO

from django.test import TestCase
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import load_workbook

from apps.accounts.models import Usuario
from .models import Afiliado, Financiador, MembresiaFinanciador, Importacion
from .ostel_synthetic import Lote, generar_xlsx, _persona


class ImportacionAdministrativaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.org = Financiador.objects.create(nombre="Administrativa ficticia", tipo="obra_social")
        cls.user = Usuario.objects.create_user("admin-ostel@example.invalid", nombre="Operador ficticio")
        cls.miembro = MembresiaFinanciador.objects.create(financiador=cls.org, usuario=cls.user, rol="operador")

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f"/api/financiadores/{self.org.pk}/"

    def preview(self, filas=None, fuente="base", clave=None, **opciones):
        contenido = generar_xlsx(Lote("entrada", tuple(filas or [_persona(1, NUMERO="000001")])) )
        return self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido),
            "clave": str(clave or uuid4()), "fuente": fuente, **opciones,
        }, format="multipart")

    def aplicar(self, lote):
        return self.client.post(self.base + "confirmar-importacion/", {"importacion": lote["id"]}, format="json")

    def corregir(self, lote, indice=0, **extra):
        fila = lote["filas"][indice]
        return self.client.post(self.base + "subsanar-ostel/", {
            "importacion": lote["id"], "fila": fila["fila"], "revision": fila["revision"],
            "datos": fila["datos"], "motivo": "Comprobación de documento ficticio", **extra,
        }, format="json")

    def test_parcial_reintento_y_cobertura_no_habilitada(self):
        r = self.preview([_persona(1, NUMERO="000001"), _persona(2, NUMERO="000002", CUIL="123")])
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data["resumen"]["valida"], r.data["resumen"]["revision"]), (1, 1))
        aplicada = self.aplicar(r.data)
        self.assertEqual(aplicada.status_code, 200, aplicada.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 1)
        obj = Afiliado.objects.get()
        self.assertFalse(obj.cobertura_habilitada)
        self.assertEqual(obj.numero_interno, "000001")
        self.assertEqual(obj.numero, "FIC185-0001-AAA")
        self.assertIsNone(obj.plan_id)
        self.aplicar(aplicada.data)
        self.assertEqual(Afiliado.objects.count(), 1)

    def test_opcion_cobertura_altas_y_conflictos_conservan_ambos_estados(self):
        for n, habilitada in ((21, True), (22, False)):
            with self.subTest(habilitada=habilitada):
                filas = [_persona(n, NUMERO=str(n))]
                lote = self.preview(filas, cobertura_habilitada=habilitada).data
                self.assertEqual(lote["resumen"]["cobertura_habilitada"], habilitada)
                self.aplicar(lote)
                obj = Afiliado.objects.get(numero_interno=str(n))
                self.assertEqual(obj.cobertura_habilitada, habilitada)
                conflicto = self.preview(filas, cobertura_habilitada=not habilitada).data
                self.assertEqual(conflicto["resumen"]["revision"], 1)
                self.assertEqual(conflicto["filas"][0]["conflicto_cobertura"]["actual"], habilitada)
                self.assertEqual(self.aplicar(conflicto).data["resumen"]["aplicada"], 0)
                revisada = self.corregir(conflicto).data
                self.assertEqual(revisada["resumen"]["valida"], 1)
                aplicada = self.aplicar(revisada).data
                self.assertEqual(aplicada["resumen"]["aplicada"], 1)
                self.assertEqual(aplicada["resumen"]["cobertura_habilitada"], not habilitada)
                self.assertTrue(aplicada["filas"][0]["motivo_cobertura"])
                obj.refresh_from_db()
                self.assertEqual(obj.cobertura_habilitada, habilitada)

    def test_opcion_cobertura_no_fuerza_baja_ni_identidad_y_clave_la_fija(self):
        from datetime import date
        r = self.preview([_persona(31, NUMERO="31", FECBAJ=date(2026, 1, 1)),
                          _persona(32, NUMERO="32", CUIL="123")], cobertura_habilitada=True).data
        self.assertEqual(r["resumen"]["revision"], 2)
        self.assertEqual(self.aplicar(r).data["resumen"]["aplicada"], 0)
        contenido = generar_xlsx(Lote("entrada", (_persona(33, NUMERO="33"),)))
        clave = str(uuid4())
        def enviar(valor):
            return self.client.post(self.base + "previsualizar-ostel/", {
                "archivo": SimpleUploadedFile("prueba.xlsx", contenido), "clave": clave,
                "cobertura_habilitada": valor}, format="multipart")
        self.assertEqual(enviar(True).status_code, 200)
        self.assertEqual(enviar(True).status_code, 200)
        self.assertEqual(enviar(False).status_code, 400)

    def test_conflicto_cobertura_revisado_no_acepta_revision_desactualizada(self):
        self.aplicar(self.preview(cobertura_habilitada=False).data)
        r = self.preview(cobertura_habilitada=True).data
        r = self.corregir(r).data
        obj = Afiliado.objects.get()
        obj.revision += 1
        obj.save(update_fields=["revision"])
        r = self.aplicar(r).data
        self.assertEqual(r["resumen"]["aplicada"], 0)
        self.assertEqual(r["resumen"]["revision"], 1)
        obj.refresh_from_db()
        self.assertFalse(obj.cobertura_habilitada)

    def test_conflictos_en_ambas_direcciones_y_omision(self):
        for filas in (
            [_persona(1, NUMERO="000001"), _persona(2, NUMERO="000001")],
            [_persona(1, NUMERO="000001"), _persona(1, NUMERO="000002")],
        ):
            with self.subTest(filas=filas):
                r = self.preview(filas)
                self.assertEqual(r.status_code, 200, r.data)
                self.assertEqual(r.data["resumen"]["revision"], 2)
                res = self.corregir(r.data, indice=1, accion="omitir")
                self.assertEqual(res.status_code, 200, res.data)
                self.assertEqual((res.data["resumen"]["valida"], res.data["resumen"]["omitida"]), (1, 1))

    def test_subsanar_conserva_original_y_rechaza_revision_vieja(self):
        r = self.preview([_persona(1, NUMERO="000001", CUIL="123")]).data
        from .ostel_synthetic import _cuil
        d = {**r["filas"][0]["datos"], "cuil": _cuil("99018501")}
        res = self.corregir(r, datos=d)
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["resumen"]["valida"], 1)
        self.assertEqual(res.data["filas"][0]["original"]["CUIL"], "123")
        self.assertEqual(len(res.data["filas"][0]["correcciones"]), 1)
        self.assertEqual(self.corregir(r, datos=d).status_code, 400)
        self.aplicar(res.data)
        self.assertEqual(self.corregir(res.data, datos=d).status_code, 400)

    def test_codpar_conserva_tipo_y_literal_sin_crear_familias(self):
        codigos = ("0007", " ZZ ", 7, None)
        filas = tuple(_persona(n, NUMERO=f"00000{n}", NROAFI="COMPARTIDO",
                               CODPAR=codigo, NROREF="000001")
                      for n, codigo in enumerate(codigos, 1))
        contenido = generar_xlsx(Lote("codigos-sin-homologar", filas))
        clave = str(uuid4())

        def enviar():
            respuesta = self.client.post(self.base + "previsualizar-ostel/", {
                "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": clave,
                "cobertura_habilitada": False,
            }, format="multipart")
            self.assertEqual(respuesta.status_code, 200, respuesta.data)
            return respuesta.data

        lote = enviar()
        originales = [fila["original"] for fila in lote["filas"]]
        self.assertEqual([original["CODPAR"] for original in originales], list(codigos))
        self.assertEqual(lote["resumen"]["valida"], 4)
        datos = {**lote["filas"][0]["datos"], "nombre": "PERSONA FICTICIA CORREGIDA"}
        corregida = self.corregir(lote, datos=datos)
        self.assertEqual(corregida.status_code, 200, corregida.data)
        self.assertEqual([fila["original"] for fila in corregida.data["filas"]], originales)
        aplicada = self.aplicar(corregida.data)
        self.assertEqual(aplicada.status_code, 200, aplicada.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 4)
        for resultado in (self.aplicar(aplicada.data).data, enviar()):
            self.assertEqual(resultado["id"], lote["id"])
            self.assertEqual([fila["original"] for fila in resultado["filas"]], originales)
        persistida = Importacion.objects.get(pk=lote["id"])
        self.assertEqual([fila["original"] for fila in persistida.filas], originales)
        self.assertEqual(len(persistida.filas[0]["correcciones"]), 1)
        self.assertEqual(Importacion.objects.count(), 1)
        self.assertEqual(Afiliado.objects.count(), 4)
        for afiliado in Afiliado.objects.all():
            self.assertIsNone(afiliado.titular_id)
            self.assertEqual(afiliado.parentesco, "")
            self.assertFalse(afiliado.cobertura_habilitada)

    def test_codpar_no_cambia_cobertura_ni_familia_existentes(self):
        from .ostel_synthetic import _cuil

        titular = Afiliado.objects.create(financiador=self.org, numero="TITULAR",
            documento="99018502", nombre="PERSONA FICTICIA TITULAR", desde=timezone.localdate())
        afiliado = Afiliado.objects.create(financiador=self.org, numero="COMPARTIDO",
            numero_interno="000001", documento="99018501", cuil=_cuil("99018501"),
            nombre="PERSONA FICTICIA OSTEL 0001", desde=timezone.localdate(),
            titular=titular, parentesco="hijo", cobertura_habilitada=True)
        lote = self.preview([_persona(1, NUMERO="000001", NROAFI="COMPARTIDO",
                                     CODPAR="SIN-DICCIONARIO", NROREF="NO-ES-TITULAR")],
                            cobertura_habilitada=False)
        self.assertEqual(lote.status_code, 200, lote.data)
        self.assertEqual(lote.data["resumen"]["revision"], 1)
        corregida = self.corregir(lote.data)
        self.assertEqual(corregida.status_code, 200, corregida.data)
        aplicada = self.aplicar(corregida.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 1)
        afiliado.refresh_from_db()
        self.assertEqual(afiliado.titular_id, titular.pk)
        self.assertEqual(afiliado.parentesco, "hijo")
        self.assertTrue(afiliado.cobertura_habilitada)
        self.assertEqual(Importacion.objects.get(pk=lote.data["id"]).filas[0]["original"]["CODPAR"],
                         "SIN-DICCIONARIO")

    def test_cuiltit_y_datos_crudos_se_conservan_sin_crear_familia(self):
        """Entrega del 26/10/2026: CUILTIT es evidencia; las familias se arman en #181."""
        from .ostel_importacion import ORIGINAL
        from .ostel_synthetic import _cuil
        cuiltit = _cuil("99018501")
        filas = (_persona(1, NUMERO="000001", CODPAR="AAA", CODPROV="B",
                          CODLOCA="000123", CODZON="B04"),
                 _persona(2, NUMERO="000002", CODPAR="H01", CUILTIT=cuiltit))
        r = self.preview(filas, cobertura_habilitada=False)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["resumen"]["valida"], 2)
        originales = [fila["original"] for fila in r.data["filas"]]
        from .ostel_prevalidacion import COLUMNAS
        self.assertEqual(ORIGINAL, COLUMNAS)
        self.assertEqual(list(originales[0]), list(COLUMNAS))
        self.assertEqual(originales[0]["EMAIL1"], "ostel-1@example.invalid")
        self.assertEqual(originales[0]["A_CALLE"], "CALLE FICTICIA")
        self.assertEqual([o["CUILTIT"] for o in originales], [cuiltit, cuiltit])
        self.assertEqual([o["CODPAR"] for o in originales], ["AAA", "H01"])
        self.assertEqual((originales[0]["CODPROV"], originales[0]["CODLOCA"], originales[0]["CODZON"]),
                         ("B", "000123", "B04"))
        aplicada = self.aplicar(r.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 2)
        self.assertEqual([f["original"] for f in Importacion.objects.get(pk=r.data["id"]).filas], originales)
        for afiliado in Afiliado.objects.all():
            self.assertIsNone(afiliado.titular_id)
            self.assertEqual(afiliado.parentesco, "")
            self.assertFalse(afiliado.cobertura_habilitada)

    FICHA_COMPLETA = dict(SEXO="F", A_CALLE="CALLE NUEVA FICTICIA", A_NRO="0123", A_PISO="2",
                          A_DPTO="B", A_TEL1="0000-0001", CELU1="0000-0002",
                          EMAIL1="nuevo@example.invalid", CODPOS="0001",
                          DESLOCA="LOCALIDAD FICTICIA", CODPROV="C", CODLOCA="000999")

    def test_alta_nueva_completa_la_ficha_desde_ostel(self):
        from datetime import date
        from .models import HistorialAfiliacion
        r = self.preview([_persona(1, NUMERO="000001", FEC_NAC=date(1990, 5, 6), **self.FICHA_COMPLETA)])
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["filas"][0]["advertencias"], [])
        aplicada = self.aplicar(r.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 1, aplicada.data)
        obj = Afiliado.objects.get()
        self.assertEqual(
            (obj.sexo, obj.fecha_nacimiento, obj.calle, obj.altura, obj.piso, obj.departamento,
             obj.codigo_postal, obj.localidad, obj.provincia, obj.telefono, obj.celular, obj.email),
            ("F", date(1990, 5, 6), "CALLE NUEVA FICTICIA", "0123", "2", "B", "0001",
             "LOCALIDAD FICTICIA", "Ciudad Autónoma de Buenos Aires", "0000-0001", "0000-0002",
             "nuevo@example.invalid"))
        historial = HistorialAfiliacion.objects.filter(afiliado=obj, origen="importacion").last()
        self.assertEqual(historial.cambios["provincia"],
                         {"anterior": "", "nuevo": "Ciudad Autónoma de Buenos Aires"})
        self.assertEqual(historial.cambios["fecha_nacimiento"], {"anterior": None, "nuevo": "1990-05-06"})

    def test_existente_se_pisa_sin_borrar_con_vacios_ni_bloquear_por_invalidos(self):
        from .models import HistorialAfiliacion
        from .ostel_synthetic import _cuil
        existente = Afiliado.objects.create(
            financiador=self.org, numero="FIC185-0001-AAA", numero_interno="000001",
            documento="99018501", cuil=_cuil("99018501"), nombre="PERSONA FICTICIA OSTEL 0001",
            desde=timezone.localdate(), calle="CALLE VIEJA", piso="9", provincia="Mendoza",
            email="viejo@example.invalid", sexo="M")
        revision = existente.revision
        filas = [_persona(1, NUMERO="000001", SEXO="F", A_CALLE="CALLE NUEVA FICTICIA",
                          A_PISO=None, EMAIL1="no-es-un-correo", CODPROV="9"),
                 _persona(2, NUMERO="000002", **self.FICHA_COMPLETA)]
        r = self.preview(filas)
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["resumen"]["valida"], 2, r.data["filas"][0]["errores"])
        avisos = " ".join(r.data["filas"][0]["advertencias"])
        self.assertIn("EMAIL1", avisos)
        self.assertIn("CODPROV", avisos)
        self.assertEqual(r.data["resumen"]["advertencias"], 1)
        aplicada = self.aplicar(r.data)
        self.assertEqual(aplicada.data["resumen"]["aplicada"], 2, aplicada.data)
        # La advertencia sigue visible en el lote después de aplicar.
        self.assertEqual(aplicada.data["resumen"]["advertencias"], 1)
        existente.refresh_from_db()
        self.assertEqual((existente.calle, existente.piso, existente.provincia, existente.email, existente.sexo),
                         ("CALLE NUEVA FICTICIA", "9", "Mendoza", "viejo@example.invalid", "F"))
        self.assertGreater(existente.revision, revision)
        cambios = HistorialAfiliacion.objects.filter(afiliado=existente, origen="importacion").last().cambios
        self.assertEqual(cambios["calle"], {"anterior": "CALLE VIEJA", "nuevo": "CALLE NUEVA FICTICIA"})
        self.assertEqual(cambios["sexo"], {"anterior": "M", "nuevo": "F"})
        self.assertNotIn("email", cambios)
        self.assertNotIn("provincia", cambios)
        self.assertNotIn("piso", cambios)
        self.assertEqual(Afiliado.objects.get(numero_interno="000002").provincia,
                         "Ciudad Autónoma de Buenos Aires")

    def test_provincias_por_letra_son_supuesto_iso(self):
        from .ostel_importacion import PROVINCIAS, _ficha
        self.assertEqual(len(PROVINCIAS), 24)
        self.assertEqual({k: PROVINCIAS[k] for k in "BCSXM"}, {
            "B": "Buenos Aires", "C": "Ciudad Autónoma de Buenos Aires", "S": "Santa Fe",
            "X": "Córdoba", "M": "Mendoza"})
        ficha, avisos, campos = _ficha({"CODPROV": "I", "CODLOCA": "000001"})
        self.assertEqual(ficha, {})
        self.assertEqual(len(avisos), 1)
        self.assertEqual(campos, ["provincia"])

    def test_revisar_fila_corrige_la_ficha_advertida(self):
        r = self.preview([_persona(1, NUMERO="000001", SEXO="F", EMAIL1="no-es-un-correo", CODPROV="9", A_CALLE="CALLE ARCHIVO")]).data
        fila = r["filas"][0]
        self.assertEqual(fila["advertidos"], ["email", "provincia"])
        editable = {"email": "corregido@example.invalid", "provincia": "", "calle": "CALLE CORREGIDA"}
        res = self.corregir(r, ficha=editable)
        self.assertEqual(res.status_code, 200, res.data)
        fila = res.data["filas"][0]
        # El email corregido deja de estar advertido; la provincia vacía sigue sin actualizarse.
        self.assertEqual(fila["advertidos"], ["provincia"])
        self.assertEqual(len(fila["advertencias"]), 1)
        self.assertEqual(fila["ficha"], {"email": "corregido@example.invalid", "calle": "CALLE CORREGIDA"})
        self.assertEqual(fila["correcciones"][0]["ficha_nueva"], fila["ficha"])
        self.assertEqual(fila["correcciones"][0]["ficha_anterior"]["calle"], "CALLE ARCHIVO")
        self.assertEqual(res.data["resumen"]["advertencias"], 1)
        self.assertEqual(self.aplicar(res.data).data["resumen"]["aplicada"], 1)
        obj = Afiliado.objects.get()
        self.assertEqual((obj.email, obj.calle, obj.provincia), ("corregido@example.invalid", "CALLE CORREGIDA", ""))

    def test_revisar_fila_rechaza_ficha_invalida_sin_cambiar_el_lote(self):
        r = self.preview([_persona(1, NUMERO="000001", SEXO="F", CODPROV="B", EMAIL1="no-es-un-correo")]).data
        for ficha in ({"email": "sigue-mal"}, {"provincia": "Narnia"}, {"sexo": "X"},
                      {"fecha_nacimiento": "2999-01-01"}, {"fecha_nacimiento": "06/05/1990"}, {"nombre": "OTRO"}):
            with self.subTest(ficha=ficha):
                self.assertEqual(self.corregir(r, ficha=ficha).status_code, 400)
        persistida = Importacion.objects.get(pk=r["id"]).filas[0]
        self.assertEqual((persistida["revision"], persistida["correcciones"]), (1, []))
        self.assertEqual(persistida["advertidos"], ["email"])

    def test_revisar_fila_de_lote_previo_conserva_sus_advertencias(self):
        r = self.preview([_persona(1, NUMERO="000001", SEXO="F", CODPROV="B", EMAIL1="no-es-un-correo")]).data
        lote = Importacion.objects.get(pk=r["id"])
        del lote.filas[0]["advertidos"]  # lote guardado antes de registrar los campos advertidos
        lote.save(update_fields=["filas"])
        res = self.corregir(r, ficha={"email": "corregido@example.invalid"})
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data["filas"][0]["ficha"]["email"], "corregido@example.invalid")
        self.assertEqual(len(res.data["filas"][0]["advertencias"]), 1)

    def test_sin_cuiltit_o_con_otra_columna_40_no_crea_lote(self):
        from openpyxl import Workbook
        from .ostel_synthetic import COLUMNAS
        for encabezados in (COLUMNAS[:39], COLUMNAS[:39] + ("OTRA_COL",)):
            with self.subTest(ancho=len(encabezados), ultima=encabezados[-1]):
                libro = Workbook()
                libro.active.title = "Hoja1"
                libro.active.append(encabezados)
                libro.active.append(_persona(1)[:len(encabezados)])
                salida = BytesIO()
                libro.save(salida)
                respuesta = self.client.post(self.base + "previsualizar-ostel/", {
                    "archivo": SimpleUploadedFile("ficticio.xlsx", salida.getvalue()), "clave": str(uuid4()),
                }, format="multipart")
                self.assertEqual(respuesta.status_code, 400)
                self.assertFalse(Importacion.objects.exists())

    def test_bajas_no_aplican_ni_se_fuerzan_mediante_correccion(self):
        r = self.preview(fuente="bajas").data
        self.assertEqual(r["resumen"]["valida"], 0)
        self.assertEqual(self.corregir(r).status_code, 400)
        self.aplicar(r)
        self.assertFalse(Afiliado.objects.exists())

    def test_nroafi_compartido_permitido_y_repeticion_archivo_idempotente(self):
        clave = uuid4()
        filas = [_persona(1, NUMERO="000001", NROAFI="FAMILIA"), _persona(2, NUMERO="000002", NROAFI="FAMILIA")]
        r = self.preview(filas, clave=clave)
        self.assertEqual(r.data["resumen"]["valida"], 2)
        # XLSX puede tener metadatos variables: reintento exacto se comprueba aparte.
        self.aplicar(r.data)
        self.assertEqual(Afiliado.objects.filter(numero="FAMILIA").count(), 2)
        segundo = self.preview(filas)
        self.aplicar(segundo.data)
        self.assertEqual(Afiliado.objects.count(), 2)

    def test_actualizacion_preserva_cobertura_plan_fechas_y_detecta_edicion(self):
        r = self.preview().data
        self.aplicar(r)
        obj = Afiliado.objects.get()
        obj.cobertura_habilitada = True
        obj.save(update_fields=["cobertura_habilitada"])
        r = self.preview().data
        obj.revision += 1
        obj.save(update_fields=["revision"])
        res = self.aplicar(r)
        self.assertEqual(res.data["resumen"]["revision"], 1)
        corregida = self.corregir(res.data, datos={**res.data["filas"][0]["datos"], "nombre": "PERSONA FICTICIA CAMBIADA"})
        self.assertEqual(corregida.status_code, 200, corregida.data)
        self.aplicar(corregida.data)
        obj.refresh_from_db()
        self.assertTrue(obj.cobertura_habilitada)
        self.assertEqual(obj.desde, timezone.localdate())
        self.assertEqual(obj.nombre, "PERSONA FICTICIA CAMBIADA")

    def test_diferencias_entre_fuentes_requieren_revision_explicita(self):
        self.aplicar(self.preview().data)
        r = self.preview([_persona(1, NUMERO="000001", APENOM="PERSONA FICTICIA OTRA FUENTE")]).data
        self.assertEqual(r["resumen"]["revision"], 1)
        self.aplicar(r)
        self.assertNotEqual(Afiliado.objects.get().nombre, "PERSONA FICTICIA OTRA FUENTE")
        self.assertEqual(r["filas"][0]["padron_actual"]["nombre"], "PERSONA FICTICIA OSTEL 0001")
        corregida = self.corregir(r)
        self.aplicar(corregida.data)
        self.assertEqual(Afiliado.objects.get().nombre, "PERSONA FICTICIA OTRA FUENTE")

    def test_dimension_declarada_no_expande_la_lectura(self):
        from zipfile import ZipFile, ZIP_DEFLATED
        import re
        original = generar_xlsx(Lote("entrada", (_persona(1, NUMERO="000001"),)))
        salida = BytesIO()
        with ZipFile(BytesIO(original)) as fuente, ZipFile(salida, "w", ZIP_DEFLATED) as destino:
            for nombre in fuente.namelist():
                contenido = fuente.read(nombre)
                if nombre.startswith("xl/worksheets/"):
                    contenido = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD1048576"', contenido)
                destino.writestr(nombre, contenido)
        r = self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", salida.getvalue()),
            "clave": str(uuid4()), "fuente": "base"}, format="multipart")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["resumen"]["total"], 1)

    def test_alta_manual_default_y_excepcion_sin_plan(self):
        for n, extra in ((1, {}), (2, {"cobertura_habilitada": False})):
            r = self.client.post(self.base + "padron/", {
                "numero": str(n), "documento": str(99018500 + n), "nombre": "PERSONA FICTICIA",
                "desde": timezone.localdate().isoformat(), **extra}, format="json")
            self.assertEqual(r.status_code, 201, r.data)
            self.assertEqual(r.data["cobertura_habilitada"], n == 1)

    def test_cambio_de_destino_exige_ver_nueva_identidad_aunque_revision_coincida(self):
        from .ostel_synthetic import _cuil
        for n in (1, 2):
            Afiliado.objects.create(financiador=self.org, documento=str(99018500 + n),
                numero=str(n), nombre=f"PERSONA FICTICIA EXISTENTE {n}", desde=timezone.localdate())
        r = self.preview().data
        destino = {**r["filas"][0]["datos"], "documento": "99018502", "cuil": _cuil("99018502")}
        nuevo = self.corregir(r, datos=destino)
        self.assertEqual(nuevo.status_code, 200, nuevo.data)
        fila = nuevo.data["filas"][0]
        self.assertEqual(fila["padron_actual"]["documento"], "99018502")
        self.assertFalse(fila["cambios_revisados"])
        self.aplicar(nuevo.data)
        self.assertEqual(Afiliado.objects.get(documento="99018502").nombre, "PERSONA FICTICIA EXISTENTE 2")

    def test_reintento_mismos_bytes_y_clave_recupera_lote(self):
        contenido = generar_xlsx(Lote("entrada", (_persona(1, NUMERO="000001"),)))
        clave = str(uuid4())
        def enviar():
            return self.client.post(self.base + "previsualizar-ostel/", {
                "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": clave, "fuente": "base"}, format="multipart")
        a, b = enviar(), enviar()
        self.assertEqual(a.status_code, 200, a.data)
        self.assertEqual(b.status_code, 200, b.data)
        self.assertEqual(a.data["id"], b.data["id"])
        self.assertEqual(Importacion.objects.count(), 1)

    def test_auditor_y_otro_financiador_no_modifican_lotes(self):
        r = self.preview().data
        MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(rol="auditor")
        self.assertEqual(self.preview().status_code, 403)
        self.assertEqual(self.aplicar(r).status_code, 403)
        self.assertEqual(self.corregir(r).status_code, 403)
        MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(rol="operador")
        otro = Financiador.objects.create(nombre="Otro ficticio", tipo="obra_social")
        ajeno = Importacion.objects.create(financiador=otro, tipo="padron", version=3,
            clave=uuid4(), huella="a" * 64, creado_por=self.user)
        self.assertEqual(self.aplicar({"id": ajeno.pk}).status_code, 400)

    def test_exportacion_y_filtro_administrativo(self):
        r = self.preview([_persona(1, NUMERO="000001"), _persona(2, NUMERO="000002", CUIL="123")]).data
        self.aplicar(r)
        listado = self.client.get(self.base + "padron/", {"estado": "administrativas"})
        self.assertEqual(listado.status_code, 200, listado.data)
        self.assertEqual(listado.data["count"], 1)
        self.assertFalse(listado.data["results"][0]["vigente"])
        from .importaciones import descargar_rechazadas
        contenido = descargar_rechazadas(importacion=Importacion.objects.get(pk=r["id"]), usuario=self.user)
        libro = load_workbook(BytesIO(contenido), read_only=True)
        self.assertEqual(libro.active.max_row, 2)
        libro.close()
