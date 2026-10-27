"""Archivo con comentarios de Excel en las celdas con error, con datos ficticios."""
from datetime import date
from io import BytesIO
from uuid import uuid4

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Usuario
from . import importaciones
from .models import EventoCobertura, Financiador, MembresiaFinanciador, Plan
from .ostel_prevalidacion import COLUMNAS
from .ostel_synthetic import Lote, _persona, generar_xlsx


def libro(contenido):
    return load_workbook(BytesIO(contenido))


def comentario(hoja, fila, columna):
    celda = hoja.cell(fila, columna)
    return celda.comment.text if celda.comment else None


class AnotacionesTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.org = Financiador.objects.create(nombre="Anotaciones ficticia", tipo="obra_social")
        cls.user = Usuario.objects.create_user("anotar@example.invalid", nombre="Operador ficticio")
        MembresiaFinanciador.objects.create(financiador=cls.org, usuario=cls.user, rol="operador")
        cls.plan = Plan.objects.create(financiador=cls.org, codigo="BASICO", nombre="Básico")

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.base = f"/api/financiadores/{self.org.pk}/"

    def anotar(self, contenido, **extra):
        return self.client.post(self.base + "anotar-importacion/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), **extra}, format="multipart")

    def ostel(self, filas, columnas=COLUMNAS):
        datos = Workbook()
        hoja = datos.active
        hoja.title = "Hoja1"
        hoja.append(columnas)
        for fila in filas:
            hoja.append(fila[:len(columnas)])
        salida = BytesIO()
        datos.save(salida)
        return salida.getvalue()

    def test_archivo_ostel_rechazado_marca_la_celda_exacta_sin_cambiar_valores(self):
        contenido = self.ostel([_persona(1, NUMERO="000001")], COLUMNAS[:-1])
        rechazo = self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": str(uuid4())}, format="multipart")
        self.assertEqual(rechazo.status_code, 400)
        respuesta = self.anotar(contenido)
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, "data", None))
        self.assertEqual(respuesta["Cache-Control"], "private, no-store")
        hoja = libro(respuesta.content)["Hoja1"]
        self.assertIn("CUILTIT", comentario(hoja, 1, 40))
        self.assertIsNone(comentario(hoja, 1, 1))
        self.assertEqual(hoja.cell(2, 9).value, "X")  # los datos del archivo quedan como estaban
        self.assertTrue(EventoCobertura.objects.filter(accion="anotar_importacion").exists())

    def test_formula_y_hoja_ajena_se_marcan_donde_estan(self):
        datos = libro(self.ostel([_persona(1, NUMERO="000001")]))
        datos["Hoja1"].cell(2, 23, "=1+1")
        datos.create_sheet("Notas")
        salida = BytesIO()
        datos.save(salida)
        anotado = libro(self.anotar(salida.getvalue()).content)
        self.assertIn("Fórmula", comentario(anotado["Hoja1"], 2, 23))
        self.assertIn("Hoja no reconocida", comentario(anotado["Notas"], 1, 1))

    def test_lote_ostel_marca_errores_por_columna_y_advertencias(self):
        filas = (_persona(1, NUMERO="000001", SEXO="F", CODPROV="B", CUIL="20990185027"),
                 _persona(2, NUMERO="000002", SEXO="F", CODPROV="B", EMAIL1="no-es-un-correo"))
        contenido = generar_xlsx(Lote("entrada", filas))
        lote = self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": str(uuid4())}, format="multipart").data
        self.assertEqual(lote["filas"][0]["columnas_error"], ["NRODOC", "CUIL"])
        hoja = libro(self.anotar(contenido, importacion=lote["id"]).content)["Hoja1"]
        for columna in ("NRODOC", "CUIL"):
            self.assertIn("no concuerdan", comentario(hoja, 2, COLUMNAS.index(columna) + 1))
        self.assertIn("Advertencia: EMAIL1", comentario(hoja, 3, COLUMNAS.index("EMAIL1") + 1))
        self.assertIsNone(comentario(hoja, 3, 1))

    def test_otro_archivo_o_lote_ajeno_no_se_anotan(self):
        contenido = generar_xlsx(Lote("entrada", (_persona(1, NUMERO="000001", CUIL="20990185027"),)))
        lote = self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": str(uuid4())}, format="multipart").data
        otro = generar_xlsx(Lote("entrada", (_persona(2, NUMERO="000002", CUIL="20990185027"),)))
        self.assertEqual(self.anotar(otro, importacion=lote["id"]).status_code, 400)
        self.assertEqual(self.anotar(b"no es un xlsx").status_code, 400)
        ajena = Financiador.objects.create(nombre="Ajena ficticia", tipo="mutual")
        ajeno = Usuario.objects.create_user("ajeno@example.invalid", nombre="Ajeno")
        MembresiaFinanciador.objects.create(financiador=ajena, usuario=ajeno, rol="operador")
        self.client.force_authenticate(ajeno)
        respuesta = self.client.post(self.base + "anotar-importacion/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "importacion": lote["id"]}, format="multipart")
        self.assertIn(respuesta.status_code, (403, 404))

    def test_archivo_sin_errores_no_devuelve_copia(self):
        contenido = self.ostel([_persona(1, NUMERO="000001", SEXO="F", CODPROV="B")])
        self.assertEqual(self.anotar(contenido).status_code, 400)

    def test_plantilla_del_sistema_marca_columnas_y_se_puede_volver_a_subir(self):
        plantilla = libro(importaciones.generar_plantilla(financiador=self.org, usuario=self.user, tipo="padron"))
        for n, fila in enumerate((["000123", "00123456", "Persona ficticia", "NOEXISTE", date(2020, 1, 1)],
                                  ["000124", "00123457", "Persona ficticia", "", "no es fecha"],
                                  ["000125", "00123458", "Persona ficticia", "BASICO", date(2020, 1, 1)]), 2):
            for c, valor in enumerate(fila, 1):
                plantilla["Carga"].cell(n, c, valor)
        salida = BytesIO()
        plantilla.save(salida)
        contenido = salida.getvalue()
        lote = self.client.post(self.base + "importaciones/", {"tipo": "padron", "clave": str(uuid4()),
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido)}, format="multipart").data
        self.assertEqual(lote["resumen"]["rechazada"], 2, lote["filas"])
        respuesta = self.anotar(contenido, importacion=lote["id"])
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, "data", None))
        carga = libro(respuesta.content)["Carga"]
        self.assertIn("plan", comentario(carga, 2, 4).lower())
        self.assertIn("fecha", comentario(carga, 3, 5).lower())
        self.assertIsNone(comentario(carga, 4, 1))
        # Conserva la hoja firmada: la copia corregida vuelve a entrar como plantilla del sistema.
        nueva = self.client.post(self.base + "importaciones/", {"tipo": "padron", "clave": str(uuid4()),
            "archivo": SimpleUploadedFile("corregido.xlsx", respuesta.content)}, format="multipart")
        self.assertEqual(nueva.status_code, 201, nueva.data)

    def test_bajas_y_errores_sin_columna_marcan_su_celda_y_nunca_toda_una_columna(self):
        filas = (_persona(1, NUMERO="000001", SEXO="F", CODPROV="B", FECBAJ=date(2026, 9, 30), EMAIL1="no-es-un-correo"),
                 _persona(2, NUMERO="000002", SEXO="F", CODPROV="B", ESTADO="B"),
                 _persona(3, NUMERO="000003", SEXO="F", CODPROV="B", NROAFI=1.5))
        contenido = generar_xlsx(Lote("entrada", filas))
        lote = self.client.post(self.base + "previsualizar-ostel/", {
            "archivo": SimpleUploadedFile("ficticio.xlsx", contenido), "clave": str(uuid4())}, format="multipart").data
        # Una baja nunca se aplica: su ficha no genera advertencias.
        self.assertEqual(lote["filas"][0]["advertencias"], [])
        self.assertEqual([f["columnas_error"] for f in lote["filas"]], [["FECBAJ"], ["ESTADO"], ["NROAFI"]])
        hoja = libro(self.anotar(contenido, importacion=lote["id"]).content)["Hoja1"]
        self.assertIn("Evidencia de baja", comentario(hoja, 2, COLUMNAS.index("FECBAJ") + 1))
        self.assertIn("Evidencia de baja", comentario(hoja, 3, COLUMNAS.index("ESTADO") + 1))
        self.assertIn("Tipos de identificadores", comentario(hoja, 4, COLUMNAS.index("NROAFI") + 1))
        self.assertEqual([comentario(hoja, n, 1) for n in (2, 3, 4)], [None, None, None])
        self.assertIsNone(comentario(hoja, 2, COLUMNAS.index("EMAIL1") + 1))
