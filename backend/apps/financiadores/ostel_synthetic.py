"""Fixtures OSTEL creadas desde cero; no validan ni aplican reglas del padrón.

Determinismo de valores, tipos y orden (no del ZIP XLSX). Los códigos son
estímulos de prueba, no un diccionario homologado por OSTEL.
"""

from dataclasses import dataclass
from datetime import date
from io import BytesIO


COLUMNAS = tuple(
    "TIPO,NROAFI,CODEMP,CODPAR,APENOM,TIPODOC,NRODOC,CUIL,SEXO,COD_PLAN,"
    "FEC_NAC,FECING,FECBAJ,FECVTO,MOTIVO,NROBENE,NROPENS,CARNET,NIV,INCAPA,"
    "FACTURA,TABVAL,NUMERO,NROREF,A_CALLE,A_NRO,A_PISO,A_DPTO,A_TEL1,CELU1,"
    "EMAIL1,CODPOS,DESLOCA,CODPROV,CODZON,CODLOCA,CODEST,CODNAC,ESTADO"
    .split(",")
)
HOJAS = ("Hoja1", "afimotook")


@dataclass(frozen=True)
class Lote:
    nombre: str
    filas: tuple[tuple, ...]


@dataclass(frozen=True)
class Escenario:
    codigo: str
    lotes: tuple[Lote, ...]


def _cuil(documento):
    """Checksum módulo 11; no acredita existencia, asignación ni identidad."""
    for prefijo in ("20", "27", "23"):
        base = prefijo + documento
        resto = sum(int(n) * p for n, p in zip(base, (5, 4, 3, 2, 7, 6, 5, 4, 3, 2))) % 11
        digito = (11 - resto) % 11
        if digito != 10:
            return base + str(digito)
    raise ValueError("No se pudo construir el identificador ficticio")


def _persona(indice, **cambios):
    documento = str(99018500 + indice)
    valores = dict.fromkeys(COLUMNAS)
    valores.update(
        TIPO="SINTETICO", NROAFI=f"FIC185-{indice:04d}-AAA",
        CODEMP="EMP-FICTICIA", CODPAR="00",
        APENOM=f"PERSONA FICTICIA OSTEL {indice:04d}", TIPODOC="DNI",
        NRODOC=documento, CUIL=_cuil(documento), SEXO="X",
        COD_PLAN="PLAN-FICTICIO", FEC_NAC=date(1990, 1, 1),
        FECING=date(2026, 1, 1), FECVTO=date(2026, 12, 31),
        A_CALLE="CALLE FICTICIA", A_NRO="001", EMAIL1=f"ostel-{indice}@example.invalid",
        CODPOS="0000", DESLOCA="LOCALIDAD FICTICIA", CODPROV="00",
        CODZON="00", CODLOCA="000", CODEST="00", CODNAC="00",
        ESTADO="SINTETICO",
    )
    desconocidas = cambios.keys() - valores.keys()
    if desconocidas:
        raise ValueError(f"Columnas desconocidas: {sorted(desconocidas)}")
    valores.update(cambios)
    return tuple(valores[c] for c in COLUMNAS)


def generar_escenarios():
    """Devuelve un catálogo nuevo, inmutable y sin aleatoriedad ni reloj.

    Cada escenario es independiente. Sus lotes se recorren en el orden dado;
    no se deben concatenar los escenarios como si fueran movimientos reales.
    """
    def caso(codigo, *filas):
        return Escenario(codigo, (Lote("entrada", tuple(filas)),))

    referencia = "FIC185-0001-AAA"
    baja = _persona(1, FECBAJ=date(2026, 6, 30), MOTIVO="BAJA-FICTICIA")
    cuil = _cuil("99018501")
    invalido = cuil[:-1] + str((int(cuil[-1]) + 1) % 10)
    cuil_parcial = _cuil("99018502")
    invalido_parcial = cuil_parcial[:-1] + str((int(cuil_parcial[-1]) + 1) % 10)
    return (
        caso("base", _persona(1)),
        caso("numero_compartido", _persona(1), _persona(2, NROAFI=referencia)),
        caso("referencia_presente", _persona(1),
             _persona(2, NROREF=referencia, CODPAR="01")),
        caso("referencia_ausente", _persona(2, NROREF=referencia, CODPAR="01")),
        caso("referencia_no_aaa", _persona(1, NROAFI="FIC185-0001-ZZZ"),
             _persona(2, NROREF="FIC185-0001-ZZZ", CODPAR="01")),
        caso("cuil_invalido", _persona(1, CUIL=invalido)),
        caso("documento_discordante", _persona(1, CUIL=_cuil("99018502"))),
        caso("filas_vacias", (None,) * len(COLUMNAS), _persona(1),
             (None,) * len(COLUMNAS), _persona(2), (None,) * len(COLUMNAS)),
        caso("tipos_mixtos", _persona(1, NROAFI="000185", NRODOC="099018501"),
             _persona(2, NROAFI=185, NRODOC=99018502, FECING="2026-01-01"),
             _persona(3, NROAFI=185.5, NRODOC=True, FECING="no-es-fecha")),
        Escenario("baja_repetida", (Lote("01_baja", (baja,)), Lote("02_repeticion", (baja,)))),
        Escenario("posible_reingreso", (
            Lote("01_baja", (baja,)),
            Lote("02_posterior", (_persona(1, FECING=date(2026, 8, 1)),)),
        )),
        caso("vencimiento_lejano", _persona(1, FECVTO=date(2099, 12, 31))),
        caso("baja_anterior_ingreso", _persona(1, FECBAJ=date(2025, 12, 31))),
        caso("lote_parcial", _persona(1), _persona(2, CUIL=invalido_parcial),
             _persona(3, FECBAJ=date(2025, 12, 31))),
        caso("referencia_ambigua", _persona(1), _persona(2, NROAFI=referencia),
             _persona(3, NROREF=referencia, CODPAR="01")),
        caso("fila_repetida", _persona(1), _persona(1)),
        caso("codigos_sin_homologar", _persona(
            1, NUMERO="000042", TIPO="TIPO-FICTICIO", TIPODOC="DOC-FICTICIO",
            CODPAR="PAR-FICTICIO", MOTIVO="MOT-FICTICIO", INCAPA="INC-FICTICIA",
        )),
        caso("contactos_ausentes", _persona(
            1, A_CALLE=None, A_NRO=None, EMAIL1=None, CODPOS=None, DESLOCA=None,
            CODPROV=None, CODZON=None, CODLOCA=None,
        )),
        caso("texto_literal", _persona(
            1, NROAFI=" 000185-Ñ ", NUMERO="000042", NROREF="REF-FICTICIA/α",
            APENOM="PERSONA FICTICIA OSTEL ÑANDÚ",
        )),
        # Respuesta OSTEL recibida el 20/10: escenarios de entrada, no reglas
        # de aplicación. No se inventa CUILTIT ni se utiliza NROREF en su lugar.
        caso("identidad_individual", _persona(1, NUMERO="000001"),
             _persona(2, NUMERO="000002", NROAFI=referencia)),
        caso("cuil_distinto_mismo_numero", _persona(1, NUMERO="000001"),
             _persona(2, NUMERO="000001")),
        caso("numero_distinto_mismo_cuil", _persona(1, NUMERO="000001"),
             _persona(1, NUMERO="000002")),
        caso("categorias_sin_cobertura",
             _persona(1, TIPO="A", COD_PLAN="CC", FECBAJ=date(2026, 6, 30)),
             _persona(2, TIPO="P", COD_PLAN="CC", FECBAJ=None)),
        caso("parentescos_confirmados", *(
            _persona(i, NUMERO=f"{i:06d}", CODPAR=codigo)
            for i, codigo in enumerate((
                "AAA", "CO1", "CO2", "CC1", "CC3", "H01", "H15",
                "HJ1", "HJ6", "ME1", "ME8", "IN1", "IN5", "MD1", "MD5",
                "MG1", "MG6", "PA1", "PA2", "PP1", "PP2", "HE1", "HE6",
                "TI1", "TI3", "SO1", "SO4", "AB1", "NI1", "NI8",
                "MF1", "MF5", "OT1", "OT5", "XC1", "XC2", "XQ1", "XQ2",
                "XH1", "XH9", "XJ1", "XJ6", "XI1", "XI6", "XM1", "XM6",
                "XG1", "XG6", "XE1", "XE6", "XP1", "XP2", "XN1", "XO1",
                "XD1", "XD2",
            ), start=1)
        )),
        Escenario("cambio_empresa", (
            Lote("01_anterior", (_persona(1, NUMERO="000001", CODEMP="EMP-A"),)),
            Lote("02_posterior", (_persona(1, NUMERO="000001", CODEMP="EMP-B"),)),
        )),
        Escenario("cambio_categoria", (
            Lote("01_anterior", (_persona(1, NUMERO="000001", TIPO="A"),)),
            Lote("02_posterior", (_persona(1, NUMERO="000001", TIPO="P"),)),
        )),
        Escenario("fuentes_discordantes", (
            Lote("01_fuente_a", (baja,)),
            Lote("02_fuente_b", (_persona(1, COD_PLAN="OTRO-PLAN-FICTICIO"),)),
        )),
    )


def generar_xlsx(lote, *, hoja="Hoja1"):
    """Devuelve bytes para uso temporal, sin escribir archivos ni importar Django."""
    if hoja not in HOJAS:
        raise ValueError(f"Hoja no soportada: {hoja}")
    if any(len(fila) != len(COLUMNAS) for fila in lote.filas):
        raise ValueError("Cada fila debe tener exactamente 39 columnas")

    from openpyxl import Workbook

    libro = Workbook()
    try:
        tabla = libro.active
        tabla.title = hoja
        tabla.append(COLUMNAS)
        for fila in lote.filas:
            tabla.append(fila)
            if all(valor is None for valor in fila):
                # Excel omite celdas vacías sin estilo, perdiendo filas finales.
                # Una celda con formato conserva la posición sin inventar datos.
                tabla.cell(tabla.max_row, len(COLUMNAS)).number_format = "@"
        # El texto es literal, incluso si un futuro estímulo comienza con '='.
        for fila in tabla:
            for celda in fila:
                if isinstance(celda.value, str):
                    celda.data_type = "s"
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()
    finally:
        libro.close()


def exportar_temporales():
    """Prepara los lotes en ambas hojas y un manifiesto técnico de fixtures.

    La carpeta permanece disponible para inspección manual; quien llama debe
    eliminarla cuando termine. No recibe fuentes ni sobrescribe datos existentes.
    """
    from pathlib import Path
    from tempfile import mkdtemp
    from hashlib import sha256
    import json

    destino = Path(mkdtemp(prefix="ostel-185-ficticio-"))
    archivos = []
    for hoja in HOJAS:
        for escenario in generar_escenarios():
            carpeta = destino / hoja / escenario.codigo
            carpeta.mkdir(parents=True)
            for orden, lote in enumerate(escenario.lotes, start=1):
                ruta = carpeta / f"{lote.nombre}.xlsx"
                contenido = generar_xlsx(lote, hoja=hoja)
                with ruta.open("xb") as archivo:
                    archivo.write(contenido)
                archivos.append({
                    "archivo": ruta.relative_to(destino).as_posix(),
                    "escenario": escenario.codigo,
                    "orden_en_escenario": orden,
                    "hoja": hoja,
                    "sha256": sha256(contenido).hexdigest(),
                    "filas_fisicas_sin_encabezado": len(lote.filas),
                    "filas_excel_con_datos": [
                        numero for numero, fila in enumerate(lote.filas, start=2)
                        if any(valor is not None for valor in fila)
                    ],
                })
    with (destino / "manifiesto.json").open("x", encoding="utf-8") as archivo:
        json.dump({
            "version_fixture": 1,
            "alcance": "Datos ficticios; expectativas de entrada, no resultados del lector.",
            "columnas": COLUMNAS,
            "aceptacion_ostel": "pendiente",
            "archivos": archivos,
        }, archivo, ensure_ascii=False, indent=2)
    return destino


if __name__ == "__main__":
    print(exportar_temporales())
