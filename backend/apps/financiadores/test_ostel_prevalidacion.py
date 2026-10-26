"""Pruebas puras con planillas ficticias; ejecutar con unittest, sin Django."""
from io import BytesIO
import builtins
import importlib
import json
import sys
from unittest import TestCase, main
from unittest.mock import patch
import warnings
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from . import ostel_prevalidacion as ostel


def archivo(filas=(), nombre="Hoja1", encabezados=ostel.COLUMNAS, preparar=None):
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    if preparar:
        preparar(libro)
    salida = BytesIO()
    libro.save(salida)
    libro.close()
    return salida.getvalue()


def cambiar(contenido, ruta, transformar):
    salida = BytesIO()
    with ZipFile(BytesIO(contenido)) as origen, ZipFile(salida, "w", ZIP_DEFLATED) as destino:
        for entrada in origen.infolist():
            datos = origen.read(entrada)
            destino.writestr(entrada.filename, transformar(datos) if entrada.filename == ruta else datos)
    return salida.getvalue()


class PrevalidacionOstelTests(TestCase):
    def codigos(self, resultado):
        return [a.codigo for a in resultado.anomalias]

    def test_conserva_valores_tipos_formatos_y_no_mutacion(self):
        def preparar(libro):
            libro.active["G2"].number_format = "000000"
        datos = ["?TIPO", "000123", " 001 ", None, "PERSONA FICTICIA", "SIN_DICCIONARIO",
                 123, "00123456789", None, "PLAN-X", "fecha sin homologar",
                 None, None, "FECVTO libre", "MOTIVO libre"]
        contenido = archivo([datos], preparar=preparar)
        copia = contenido[:]
        resultado = ostel.prevalidar_ostel(contenido)
        fila = resultado.filas["Hoja1"][0]
        self.assertEqual(fila.numero, 2)
        self.assertEqual(fila.celdas[1].valor, "000123")
        self.assertEqual(fila.celdas[2].valor, " 001 ")
        self.assertEqual(fila.celdas[6].valor, "123")
        self.assertEqual(fila.celdas[6].tipo, "n")
        self.assertEqual(fila.celdas[6].formato, "000000")
        self.assertEqual([c.valor for c in fila.celdas[13:15]], datos[13:15])
        self.assertEqual(resultado.anomalias, ())
        self.assertEqual(contenido, copia)
        self.assertNotIn("PERSONA FICTICIA", repr(resultado))
        self.assertNotIn("PERSONA FICTICIA", str(resultado.resumen))

    def test_vacias_ceros_espacios_y_duplicados_con_coordenadas(self):
        resultado = ostel.prevalidar_ostel(archivo([
            [], ["0001"], [None, ""], ["0001"], [0], [False], [" "],
        ]))
        self.assertEqual(resultado.resumen["filas_no_vacias"], {"Hoja1": 5})
        self.assertEqual(resultado.resumen["duplicados_exactos"], 1)
        self.assertEqual(resultado.anomalias, (ostel.Anomalia("fila_duplicada", "Hoja1", 5, fila_original=3),))
        self.assertEqual([f.numero for f in resultado.filas["Hoja1"]], [3, 5, 6, 7, 8])

    def test_ambas_hojas_y_volumenes_representativos(self):
        def preparar(libro):
            otra = libro.create_sheet("afimotook")
            otra.append(ostel.COLUMNAS)
            otra.append(["0001"])
        resultado = ostel.prevalidar_ostel(archivo([["0002"]], preparar=preparar))
        self.assertEqual(resultado.resumen["filas_no_vacias"], {"Hoja1": 1, "afimotook": 1})
        self.assertEqual(resultado.anomalias, ())
        for nombre, cantidad in (("Hoja1", 5754), ("afimotook", 7486)):
            with self.subTest(nombre=nombre):
                resultado = ostel.prevalidar_ostel(archivo(
                    ([f"FICTICIO-{i:06}"] * 39 for i in range(cantidad)), nombre=nombre,
                ))
                self.assertEqual(resultado.resumen["filas_no_vacias"], {nombre: cantidad})
                self.assertEqual(resultado.anomalias, ())

    def test_estructura_incorrecta_no_oculta_datos(self):
        encabezados = list(ostel.COLUMNAS)
        encabezados[0], encabezados[1] = encabezados[1], encabezados[0]
        resultado = ostel.prevalidar_ostel(archivo(
            [["valor"] + [None] * 38 + ["extra"]], encabezados=encabezados,
        ))
        self.assertEqual(resultado.resumen["filas_no_vacias"], {"Hoja1": 1})
        self.assertIn(ostel.Anomalia("encabezado_incorrecto", "Hoja1", 1, 1), resultado.anomalias)
        self.assertIn(ostel.Anomalia("encabezado_incorrecto", "Hoja1", 1, 2), resultado.anomalias)
        self.assertIn(ostel.Anomalia("columna_adicional", "Hoja1", 2, 40), resultado.anomalias)
        self.assertEqual(resultado.filas["Hoja1"][0].celdas[39].valor, "extra")

    def test_hoja_desconocida_no_filtra_nombre(self):
        resultado = ostel.prevalidar_ostel(archivo(nombre="NOMBRE_PRIVADO"))
        self.assertEqual(self.codigos(resultado), ["hoja_no_reconocida", "sin_hojas_reconocidas"])
        self.assertNotIn("NOMBRE_PRIVADO", repr(resultado))

    def test_formula_error_y_tipos_xml_sin_reglas_semanticas(self):
        resultado = ostel.prevalidar_ostel(archivo([["=1+1", "#DIV/0!"]]))
        self.assertEqual(self.codigos(resultado), ["formula", "error_excel"])
        self.assertEqual(resultado.filas["Hoja1"][0].celdas[0].formula, "1+1")
        self.assertEqual(resultado.resumen["filas_no_vacias"]["Hoja1"], 1)
        for anterior, nuevo, codigo in [
            (b't="n"><v>123</v>', b't="n"><v>NaN</v>', "numero_invalido"),
            (b't="n"><v>123</v>', b't="b"><v>3</v>', "booleano_invalido"),
            (b't="n"><v>123</v>', b't="inventado"><v>123</v>', "tipo_no_admitido"),
        ]:
            with self.subTest(codigo=codigo):
                contenido = cambiar(archivo([[123]]), "xl/worksheets/sheet1.xml",
                                    lambda d: d.replace(anterior, nuevo))
                self.assertIn(codigo, self.codigos(ostel.prevalidar_ostel(contenido)))

    def test_dimension_declarada_no_controla_conteo(self):
        contenido = cambiar(archivo([["00001"]]), "xl/worksheets/sheet1.xml",
                            lambda d: d.replace(b'ref="A1:AM2"', b'ref="A1"'))
        self.assertEqual(ostel.prevalidar_ostel(contenido).resumen["filas_no_vacias"]["Hoja1"], 1)

    def test_encabezado_ausente_y_faltante(self):
        contenido = cambiar(archivo(), "xl/worksheets/sheet1.xml",
                            lambda d: d.replace(b'<row r="1">', b'<row r="2">').replace(b'1"', b'2"'))
        self.assertIn("encabezado_ausente", self.codigos(ostel.prevalidar_ostel(contenido)))
        resultado = ostel.prevalidar_ostel(archivo(encabezados=ostel.COLUMNAS[:-1]))
        self.assertEqual(resultado.anomalias, (ostel.Anomalia("encabezado_incorrecto", "Hoja1", 1, 39),))

    def test_archivos_invalidos_y_limites(self):
        for contenido in (b"", b"no zip", "no bytes"):
            with self.subTest(contenido=contenido), self.assertRaises(ostel.ArchivoOstelInvalido):
                ostel.prevalidar_ostel(contenido)
        contenido = archivo([["ficticio"], ["otro"]])
        for nombre, limite in (("MAX_BYTES", 1), ("MAX_EXPANDIDO", 1), ("MAX_ENTRADAS", 1), ("MAX_FILAS", 1)):
            with self.subTest(limite=nombre), patch.object(ostel, nombre, limite):
                with self.assertRaises(ostel.ArchivoOstelInvalido):
                    ostel.prevalidar_ostel(contenido)

    def test_rechaza_dtd_xml_roto_y_relaciones_externas(self):
        for ruta, transformar in [
            ("xl/workbook.xml", lambda d: b'<!DOCTYPE workbook [<!ENTITY x "secreto">]>' + d),
            ("xl/workbook.xml", lambda d: b"<roto>"),
            ("xl/_rels/workbook.xml.rels", lambda d: d.replace(b' Target=', b' TargetMode="External" Target=')),
            ("[Content_Types].xml", lambda d: d.replace(b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")),
        ]:
            with self.subTest(ruta=ruta), self.assertRaises(ostel.ArchivoOstelInvalido):
                ostel.prevalidar_ostel(cambiar(archivo(), ruta, transformar))

    def test_rechaza_macros_objetos_y_entradas_duplicadas(self):
        for nombre in ("xl/vbaProject.bin", "xl/embeddings/obj.bin", "xl/externalLinks/link.xml", "xl/workbook.xml"):
            salida = BytesIO(archivo())
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                with ZipFile(salida, "a") as paquete:
                    paquete.writestr(nombre, b"ficticio")
            with self.subTest(nombre=nombre), self.assertRaises(ostel.ArchivoOstelInvalido):
                ostel.prevalidar_ostel(salida.getvalue())

    def test_rechaza_compresion_no_admitida(self):
        salida = BytesIO()
        with ZipFile(salida, "w", ZIP_BZIP2) as paquete:
            paquete.writestr("ficticio.xml", b"<ficticio/>")
        with self.assertRaises(ostel.ArchivoOstelInvalido):
            ostel.prevalidar_ostel(salida.getvalue())

    def test_shared_strings_con_ceros_y_rich_text(self):
        contenido = cambiar(archivo([["ficticio"]]), "xl/worksheets/sheet1.xml",
                            lambda d: d.replace(b'<c r="A2" t="inlineStr"><is><t>ficticio</t></is></c>',
                                                b'<c r="A2" t="s"><v>0</v></c>'))
        contenido = cambiar(contenido, "xl/_rels/workbook.xml.rels",
                            lambda d: d.replace(b"</Relationships>",
                                                b'<Relationship Id="strings" Type="' + ostel.REL[1:-1].encode()
                                                + b'/sharedStrings" Target="sharedStrings.xml"/></Relationships>'))
        salida = BytesIO(contenido)
        with ZipFile(salida, "a") as paquete:
            paquete.writestr("xl/sharedStrings.xml",
                             f'<sst xmlns="{ostel.NS[1:-1]}"><si><r><t>000</t></r>'
                             '<r><t xml:space="preserve">12 </t></r></si></sst>')
        resultado = ostel.prevalidar_ostel(salida.getvalue())
        self.assertEqual(resultado.filas["Hoja1"][0].celdas[0].valor, "00012 ")

    def test_coordenadas_duplicadas_fuera_de_limite_y_indices_invalidos(self):
        for transformar in (
            lambda d: d.replace(b'r="A2"', b'r="XFD2"'),
            lambda d: d.replace(b'r="A2"', b'r="A3"'),
            lambda d: d.replace(b'<row r="2">', b'<row r="10002">'),
            lambda d: d.replace(b'<c r="A2"', b'<c s="-1" r="A2"'),
            lambda d: d.replace(b'</row>', b'<c r="A1" t="n"><v>0</v></c></row>', 1),
        ):
            with self.subTest(transformar=transformar), self.assertRaises(ostel.ArchivoOstelInvalido):
                ostel.prevalidar_ostel(cambiar(archivo([[1]]), "xl/worksheets/sheet1.xml", transformar))

    def test_no_carga_django_ni_importador(self):
        importar = builtins.__import__

        def sin_django(nombre, *args, **kwargs):
            self.assertFalse(nombre == "django" or nombre.startswith("django."))
            self.assertNotIn("importaciones", nombre)
            return importar(nombre, *args, **kwargs)

        with patch("builtins.__import__", side_effect=sin_django):
            importlib.reload(ostel)

    def test_literales_xml_invalidos_conservan_valor_y_reportan_coordenada(self):
        for tipo, valor, codigo in (
            ("n", "1_000", "numero_invalido"),
            ("n", "１２", "numero_invalido"),
            ("d", "no-fecha", "fecha_xml_invalida"),
            ("d", "2026-02-30", "fecha_xml_invalida"),
        ):
            with self.subTest(tipo=tipo, valor=valor):
                contenido = cambiar(archivo([[123]]), "xl/worksheets/sheet1.xml",
                                    lambda d: d.replace(b't="n"><v>123</v>',
                                                        f't="{tipo}"><v>{valor}</v>'.encode()))
                resultado = ostel.prevalidar_ostel(contenido)
                self.assertIn(ostel.Anomalia(codigo, "Hoja1", 2, 1), resultado.anomalias)
                self.assertEqual(resultado.filas["Hoja1"][0].celdas[0].valor, valor)

    def test_fecha_xml_valida_y_serial_se_conservan_sin_conversion(self):
        for valor in ("2026-10-20", "2026-10-20T12:30:01.123Z", "12:30:01"):
            with self.subTest(valor=valor):
                contenido = cambiar(archivo([[123]]), "xl/worksheets/sheet1.xml",
                                    lambda d: d.replace(b't="n"><v>123</v>',
                                                        f't="d"><v>{valor}</v>'.encode()))
                resultado = ostel.prevalidar_ostel(contenido)
                self.assertEqual(resultado.anomalias, ())
                self.assertEqual(resultado.filas["Hoja1"][0].celdas[0].valor, valor)

    def test_indices_invalidos_informan_coordenada_sin_valor(self):
        for anterior, nuevo in (
            (b't="n"><v>123</v>', b't="s"><v>987654</v>'),
            (b'<c r="A2"', b'<c s="987654" r="A2"'),
        ):
            with self.subTest(nuevo=nuevo):
                contenido = cambiar(archivo([[123]]), "xl/worksheets/sheet1.xml",
                                    lambda d: d.replace(anterior, nuevo))
                with self.assertRaises(ostel.ArchivoOstelInvalido) as error:
                    ostel.prevalidar_ostel(contenido)
                self.assertIn("Hoja1, fila 2, columna 1", str(error.exception))
                self.assertNotIn("987654", str(error.exception))

    def test_formulas_compartidas_distintas_no_son_duplicados_exactos(self):
        contenido = cambiar(archivo([["=1"], ["=1"]]), "xl/worksheets/sheet1.xml",
                            lambda d: d.replace(b"<f>1</f>", b'<f t="shared" si="0"/>', 1)
                            .replace(b"<f>1</f>", b'<f t="shared" si="1"/>', 1))
        resultado = ostel.prevalidar_ostel(contenido)
        self.assertEqual(resultado.resumen["duplicados_exactos"], 0)
        self.assertEqual(self.codigos(resultado), ["formula", "formula"])
        mismo_grupo = cambiar(contenido, "xl/worksheets/sheet1.xml",
                              lambda d: d.replace(b'si="1"', b'si="0"'))
        self.assertEqual(ostel.prevalidar_ostel(mismo_grupo).resumen["duplicados_exactos"], 0)

    def test_partes_compartidas_resueltas_por_relacion(self):
        def preparar(libro):
            libro.active["A2"].number_format = "000000"
        contenido = archivo([[123]], preparar=preparar)
        salida = BytesIO()
        with ZipFile(BytesIO(contenido)) as origen, ZipFile(salida, "w", ZIP_DEFLATED) as destino:
            for entrada in origen.infolist():
                datos = origen.read(entrada)
                if entrada.filename == "xl/_rels/workbook.xml.rels":
                    datos = datos.replace(b"styles.xml", b"custom/estilos.xml").replace(
                        b"</Relationships>",
                        b'<Relationship Id="strings" Type="' + ostel.REL[1:-1].encode()
                        + b'/sharedStrings" Target="custom/textos.xml"/></Relationships>')
                elif entrada.filename == "[Content_Types].xml":
                    datos = datos.replace(b"/xl/styles.xml", b"/xl/custom/estilos.xml").replace(
                        b"</Types>", b'<Override PartName="/xl/custom/textos.xml" ContentType="'
                        b'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
                elif entrada.filename == "xl/worksheets/sheet1.xml":
                    datos = datos.replace(b't="n"><v>123</v>', b't="s"><v>0</v>')
                destino.writestr(
                    "xl/custom/estilos.xml" if entrada.filename == "xl/styles.xml" else entrada.filename, datos,
                )
            destino.writestr("xl/custom/textos.xml",
                             f'<sst xmlns="{ostel.NS[1:-1]}"><si><t>000123</t></si></sst>')
        resultado = ostel.prevalidar_ostel(salida.getvalue())
        celda = resultado.filas["Hoja1"][0].celdas[0]
        self.assertEqual((celda.valor, celda.formato), ("000123", "000000"))


def caso_demo(opcion):
    """Entradas exclusivamente ficticias, generadas en memoria bajo demanda."""
    fila = [None] * len(ostel.COLUMNAS)
    for campo, valor in {
        "TIPO": "SIN_HOMOLOGAR", "NROAFI": "000123", "APENOM": "PERSONA FICTICIA",
        "TIPODOC": "CODIGO_FICTICIO", "NRODOC": "00000123", "NROREF": "000009",
        "FECVTO": "SIN_INTERPRETAR", "MOTIVO": "SIN_DICCIONARIO",
    }.items():
        fila[ostel.COLUMNAS.index(campo)] = valor
    if opcion == "1":
        return archivo([fila])
    if opcion == "2":
        return archivo([fila, [], fila])
    if opcion == "3":
        return archivo([fila], encabezados=("CABECERA_INCORRECTA",) + ostel.COLUMNAS[1:])
    if opcion == "4":
        fila[0:2] = ["=1+1", "#DIV/0!"]
        return archivo([fila])
    if opcion == "5":
        return b"archivo ficticio que no es XLSX"
    if opcion in ("6", "7"):
        nombre, cantidad = ("Hoja1", 5754) if opcion == "6" else ("afimotook", 7486)
        return archivo(([f"FICTICIO-{i:06}"] * 39 for i in range(cantidad)), nombre=nombre)
    raise ValueError("Opción de demostración no admitida.")


def demo():
    """Consola de prueba local; no abre archivos reales ni registra datos."""
    from dataclasses import asdict

    print("OSTEL - prueba técnica local. Datos ficticios, sin DB ni API.")
    print("Sin anomalías NO significa apto para importar. 0 = salir.")
    opciones = (
        "1 Ceros y códigos originales | 2 Vacíos y duplicado | 3 Encabezado incorrecto\n"
        "4 Fórmula y error Excel | 5 Archivo inválido | 6 Padrón 5754 | 7 Histórico 7486"
    )
    while True:
        print(opciones)
        try:
            opcion = input("Caso [0-7]: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nFin de la prueba local.")
            return
        if opcion == "0":
            return
        if opcion not in tuple("1234567"):
            print("Elegí un número de 0 a 7.")
            continue
        try:
            resultado = ostel.prevalidar_ostel(caso_demo(opcion))
        except ostel.ArchivoOstelInvalido as error:
            print(f"Rechazado: {error}")
            continue
        print(json.dumps(resultado.resumen, ensure_ascii=False, indent=2))
        for anomalia in resultado.anomalias:
            print(json.dumps(asdict(anomalia), ensure_ascii=False))
        if opcion == "1":
            fila = resultado.filas["Hoja1"][0]
            print("Muestra ficticia de preservación:")
            for campo in ("NROAFI", "NRODOC", "NROREF", "FECVTO", "MOTIVO"):
                print(f"  {campo}: {fila.celdas[ostel.COLUMNAS.index(campo)].valor!r}")


class CuiltitOstelTests(TestCase):
    """Entrega OSTEL del 26/10/2026: CUILTIT como columna 40, con datos ficticios."""

    def fila(self, cuiltit="20990185011"):
        return [f"FICTICIO-{i:02}" for i in range(39)] + [cuiltit]

    def test_columna_40_cuiltit_se_acepta_y_conserva_texto(self):
        resultado = ostel.prevalidar_ostel(archivo(
            [self.fila("00990185011"), self.fila()], encabezados=ostel.COLUMNAS_CON_TITULAR,
        ))
        self.assertEqual(resultado.anomalias, ())
        self.assertEqual(resultado.columnas, {"Hoja1": ostel.COLUMNAS_CON_TITULAR})
        self.assertEqual([f.celdas[39].valor for f in resultado.filas["Hoja1"]],
                         ["00990185011", "20990185011"])
        self.assertEqual(ostel.COLUMNAS_CON_TITULAR[:39], ostel.COLUMNAS)

    def test_formato_de_39_columnas_sigue_vigente(self):
        resultado = ostel.prevalidar_ostel(archivo([self.fila()[:39]]))
        self.assertEqual(resultado.anomalias, ())
        self.assertEqual(resultado.columnas, {"Hoja1": ostel.COLUMNAS})

    def test_columna_40_sin_encabezado_exacto_es_adicional(self):
        for encabezado in ("CUILTIT ", "cuiltit", "OTRA", None):
            with self.subTest(encabezado=encabezado):
                resultado = ostel.prevalidar_ostel(archivo(
                    [self.fila()], encabezados=ostel.COLUMNAS + (encabezado,),
                ))
                self.assertEqual(resultado.columnas, {"Hoja1": ostel.COLUMNAS})
                self.assertIn(ostel.Anomalia("columna_adicional", "Hoja1", 2, 40), resultado.anomalias)

    def test_encabezado_cuiltit_formula_o_numero_no_habilita_columna(self):
        # openpyxl escribe el encabezado como cadena en línea (posición 40 = AN).
        original = b'<c r="AN1" t="inlineStr"><is><t>CUILTIT</t></is></c>'
        for nuevo in (b'<c r="AN1" t="str"><f>"CUILTIT"</f><v>CUILTIT</v></c>',
                      b'<c r="AN1"><v>1</v></c>'):
            with self.subTest(nuevo=nuevo):
                base = archivo([self.fila()], encabezados=ostel.COLUMNAS_CON_TITULAR)
                def reemplazar(d):
                    self.assertIn(original, d)
                    return d.replace(original, nuevo)
                contenido = cambiar(base, "xl/worksheets/sheet1.xml", reemplazar)
                resultado = ostel.prevalidar_ostel(contenido)
                self.assertEqual(resultado.columnas, {"Hoja1": ostel.COLUMNAS})
                self.assertIn(ostel.Anomalia("columna_adicional", "Hoja1", 2, 40), resultado.anomalias)

    def test_columna_41_con_cuiltit_sigue_siendo_adicional(self):
        resultado = ostel.prevalidar_ostel(archivo(
            [self.fila() + ["extra"]], encabezados=ostel.COLUMNAS_CON_TITULAR,
        ))
        self.assertEqual(resultado.anomalias, (ostel.Anomalia("columna_adicional", "Hoja1", 2, 41),))

    def test_esquema_se_decide_por_hoja(self):
        def preparar(libro):
            otra = libro.create_sheet("afimotook")
            otra.append(ostel.COLUMNAS_CON_TITULAR)
            otra.append(self.fila())
        resultado = ostel.prevalidar_ostel(archivo([self.fila()[:39]], preparar=preparar))
        self.assertEqual(resultado.anomalias, ())
        self.assertEqual(resultado.columnas, {"Hoja1": ostel.COLUMNAS, "afimotook": ostel.COLUMNAS_CON_TITULAR})


if __name__ == "__main__":
    if sys.argv[1:] == ["--demo"]:
        demo()
    else:
        main()
