"""Copia del archivo importado con comentarios de Excel en las celdas con error.

No se guarda nada: quien importa vuelve a enviar el mismo archivo y recibe una
copia anotada. Así no se conservan planillas con datos personales y la copia
mantiene sus hojas (incluida `_contexto` de la plantilla del sistema), por lo
que se puede corregir y volver a subir. Con un lote, la huella del archivo debe
coincidir con la del lote: no se anota un archivo con errores de otro.
"""
from hashlib import sha256
from io import BytesIO
from time import monotonic

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from .permisos import requerir_financiador

AUTOR = "Revisión de importación"
RESALTADO = PatternFill("solid", fgColor="FDE2E1")
AVISO = PatternFill("solid", fgColor="FEF3C7")
# Mismos textos que la revisión técnica de la UI (RevisionTecnicaExcel.jsx).
ANOMALIAS = {
    "encabezado_incorrecto": "Encabezado ausente, cambiado o fuera de orden. Se esperaba «{esperado}».",
    "encabezado_ausente": "Falta la fila de encabezados.",
    "columna_adicional": "Columna adicional con contenido: el formato tiene 40 columnas.",
    "formula": "Fórmula detectada; no se ejecutó. Reemplazala por su valor.",
    "error_excel": "La celda contiene un error de Excel.",
    "tipo_no_admitido": "Tipo de celda no admitido.",
    "booleano_invalido": "Valor booleano XLSX inválido.",
    "numero_invalido": "Representación numérica XLSX inválida.",
    "fecha_xml_invalida": "Representación de fecha XLSX no admitida.",
    "hoja_no_reconocida": "Hoja no reconocida: sólo se leen Hoja1 y afimotook.",
    "sin_hojas_reconocidas": "No hay hojas reconocidas: el archivo necesita una hoja Hoja1 o afimotook.",
}
# Filas idénticas no bloquean el lote: no se marcan.
NO_BLOQUEANTES = {"fila_duplicada"}


def _huella_ostel(lote, contenido):
    opcion = lote.resumen.get("cobertura_habilitada")
    opcion = b"" if opcion is None else str(opcion).encode()
    return sha256(lote.resumen.get("fuente", "base").encode() + opcion + contenido).hexdigest()


def _huella_sistema(lote, contenido):
    opcion = lote.resumen.get("cobertura_habilitada")
    return sha256((b"" if opcion is None else str(opcion).encode()) + contenido).hexdigest()


def _agregar(marcas, hoja, fila, columna, texto, aviso=False):
    actual = marcas.setdefault((hoja, fila, columna), {"textos": [], "aviso": True})
    if texto not in actual["textos"]:
        actual["textos"].append(texto)
    actual["aviso"] = actual["aviso"] and aviso


def _escribir(contenido, marcas):
    if not marcas:
        raise ValidationError("El archivo no tiene celdas con errores para marcar.")
    libro = load_workbook(BytesIO(contenido), keep_links=False)
    try:
        for (hoja, fila, columna), marca in marcas.items():
            for nombre in ([hoja] if hoja else libro.sheetnames[:1]):
                if nombre not in libro.sheetnames:
                    continue
                celda = libro[nombre].cell(fila, columna)
                texto = "\n".join(marca["textos"])
                celda.comment = Comment(texto, AUTOR, width=360, height=40 + 18 * (texto.count("\n") + len(texto) // 50))
                celda.fill = AVISO if marca["aviso"] else RESALTADO
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()
    finally:
        libro.close()


def anotar_ostel(*, financiador, usuario, contenido, lote=None):
    """Sin lote, marca las anomalías técnicas que rechazaron el archivo; con
    lote, los errores de las filas a revisar y las advertencias de ficha."""
    from .ostel_prevalidacion import COLUMNAS, ArchivoOstelInvalido, HOJAS, prevalidar_ostel

    requerir_financiador(usuario, financiador.pk, escritura=True)
    try:
        # También valida que el paquete sea seguro antes de abrirlo con openpyxl.
        resultado = prevalidar_ostel(contenido)
    except ArchivoOstelInvalido as exc:
        raise ValidationError(f"No se puede marcar este archivo: {exc}") from None
    marcas = {}
    if lote is None:
        for a in resultado.anomalias:
            if a.codigo in NO_BLOQUEANTES:
                continue
            texto = ANOMALIAS.get(a.codigo, "Anomalía técnica.")
            if a.codigo == "encabezado_incorrecto" and a.columna <= len(COLUMNAS):
                texto = texto.format(esperado=COLUMNAS[a.columna - 1])
            elif a.codigo == "encabezado_incorrecto":
                texto = "Encabezado fuera de las 40 columnas del formato."
            if a.codigo == "hoja_no_reconocida":
                # El lector no conserva nombres arbitrarios; acá se marca cada
                # hoja ajena del propio archivo de quien lo subió.
                libro = load_workbook(BytesIO(contenido), read_only=True, keep_links=False)
                try:
                    for nombre in libro.sheetnames:
                        if nombre not in HOJAS:
                            _agregar(marcas, nombre, 1, 1, texto)
                finally:
                    libro.close()
                continue
            _agregar(marcas, a.hoja, a.fila or 1, a.columna or 1, texto)
        return _escribir(contenido, marcas)
    if lote.version != 3 or lote.financiador_id != financiador.pk:
        raise ValidationError("El lote no corresponde a una importación OSTEL de este financiador.")
    if _huella_ostel(lote, contenido) != lote.huella:
        raise ValidationError("El archivo no es el mismo que se revisó en este lote.")
    from .ostel_importacion import FICHA
    columna_ficha = {**dict(FICHA), "sexo": "SEXO", "fecha_nacimiento": "FEC_NAC", "provincia": "CODPROV"}
    for fila in lote.filas:
        if fila["estado"] == "revision" and fila["errores"]:
            texto = " ".join(fila["errores"])
            for columna in fila.get("columnas_error") or ["NRODOC"]:
                _agregar(marcas, fila["hoja"], fila["fila_excel"], COLUMNAS.index(columna) + 1, texto)
        if fila["estado"] in ("aplicada", "omitida"):
            continue
        for campo, texto in zip(fila.get("advertidos") or [], fila.get("advertencias") or []):
            _agregar(marcas, fila["hoja"], fila["fila_excel"], COLUMNAS.index(columna_ficha[campo]) + 1,
                     f"Advertencia: {texto}", aviso=True)
    return _escribir(contenido, marcas)


def anotar_sistema(*, financiador, usuario, contenido, lote):
    """Marca las filas rechazadas o a revisar de un lote de la plantilla del sistema."""
    from .importaciones import _verificar_zip, encabezados, limites_importacion

    requerir_financiador(usuario, financiador.pk, escritura=True)
    if lote.version == 3 or lote.financiador_id != financiador.pk:
        raise ValidationError("El lote no corresponde a una importación de la plantilla del sistema.")
    limite = limites_importacion()
    if len(contenido) > limite["bytes"]:
        raise ValidationError("El archivo supera el tamaño máximo permitido.")
    if _huella_sistema(lote, contenido) != lote.huella:
        raise ValidationError("El archivo no es el mismo que se revisó en este lote.")
    _verificar_zip(contenido, limite, monotonic())
    columnas = encabezados(lote.tipo, lote.version)
    marcas = {}
    for fila in lote.filas:
        if fila["estado"] not in ("rechazada", "revision", "error_tecnico") or not fila["errores"]:
            continue
        texto = " ".join(fila["errores"])
        for columna in fila.get("columnas_error") or ["Documento"]:
            if not isinstance(columna, int) and columna not in columnas:
                columna = "Documento"
            posicion = columna if isinstance(columna, int) else columnas.index(columna) + 1
            _agregar(marcas, "Carga", fila["fila"], posicion, texto)
    return _escribir(contenido, marcas)
