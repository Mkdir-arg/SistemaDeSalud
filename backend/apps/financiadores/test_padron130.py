"""Padrón ampliado: persistencia API, compatibilidad e invariantes de confidencialidad."""
import json
from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from types import SimpleNamespace

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import Usuario
from apps.simulacion import contexto
from . import models as m
from . import padron as p
from .importaciones import aplicar_importacion, generar_plantilla, previsualizar_importacion
from .services import registrar_afiliado


def cuit(dni, prefijo="20"):
    base = prefijo + dni.zfill(8)
    digito = (11 - sum(int(n) * w for n, w in zip(base, (5, 4, 3, 2, 7, 6, 5, 4, 3, 2))) % 11) % 11
    if digito == 10:
        return cuit(dni, "23")
    return base + str(digito)


class PadronValidadoresTests(SimpleTestCase):
    def test_cbu_verifica_ambos_bloques_con_ceros(self):
        valor = "2850590940090418135201"
        p.validar_cbu(valor)
        for pos in (7, 21):
            invalido = valor[:pos] + str((int(valor[pos]) + 1) % 10) + valor[pos + 1:]
            with self.assertRaises(ValidationError):
                p.validar_cbu(invalido)
        for invalido in ("123", "a" * 22, "２" * 22, "0" * 22):
            with self.assertRaises(ValidationError):
                p.validar_cbu(invalido)

    def test_cuit_longitud_y_digito(self):
        p.validar_cuit("20329642330")
        for invalido in ("123", "20329642331", "A" * 11, "２" * 11, "0" * 11):
            with self.assertRaises(ValidationError):
                p.validar_cuit(invalido)


class Padron130ApiTests(APITestCase):
    def setUp(self):
        self.hoy = timezone.localdate()
        self.org = m.Financiador.objects.create(nombre="Demo A", tipo="obra_social")
        self.otra = m.Financiador.objects.create(nombre="Demo B", tipo="obra_social")
        self.user = Usuario.objects.create_user("padron130@test.local", "clave")
        self.member = m.MembresiaFinanciador.objects.create(financiador=self.org, usuario=self.user, rol="admin")
        self.a = m.Afiliado.objects.create(financiador=self.org, numero="001", documento="32964233", nombre="Afiliado legado", desde=self.hoy)
        self.b = m.Afiliado.objects.create(financiador=self.org, numero="002", documento="11111111", nombre="Familiar", desde=self.hoy)
        self.ajeno = m.Afiliado.objects.create(financiador=self.otra, numero="001", documento="22222222", nombre="Ajeno", desde=self.hoy)
        self.base = f"/api/financiadores/{self.org.pk}/"
        self.url = self.base + f"padron/{self.a.pk}/"
        self.client.force_authenticate(self.user)

    def patch(self, data, url=None):
        return self.client.patch(url or self.url, data, format="json")

    def permisos(self, salud=True, cobranza=True):
        self.member.acceso_padron_salud = salud
        self.member.acceso_padron_cobranza = cobranza
        self.member.save()

    def libro(self, filas, version=2, encabezado=None):
        libro = load_workbook(BytesIO(generar_plantilla(financiador=self.org, usuario=self.user, tipo="padron", version=version)))
        hoja = libro["Carga"]
        headers = [c.value for c in hoja[1]]
        if encabezado:
            hoja.cell(1, len(headers) + 1, encabezado)
        for pos, fila in enumerate(filas, 2):
            for col, campo in enumerate(headers, 1):
                if campo in fila:
                    hoja.cell(pos, col, fila[campo])
        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def fila(self, **extras):
        return {"N.º de afiliado": "001", "Documento": self.a.documento, "Nombre y apellido": self.a.nombre,
                "Vigente desde": self.hoy.isoformat(), **extras}

    def preview(self, archivo):
        return previsualizar_importacion(financiador=self.org, usuario=self.user, tipo="padron", archivo=archivo, clave=uuid4())

    def test_alta_exclusiva_rechaza_duplicado_sin_modificar(self):
        plan = m.Plan.objects.create(financiador=self.org, codigo="BASE", nombre="Base")
        self.a.plan = plan
        self.a.save(update_fields=["plan"])
        antes = (self.a.numero, self.a.nombre, self.a.plan_id, self.a.desde, self.a.revision)
        response = self.client.post(self.base + "padron/", {
            "numero": "REEMPLAZO", "documento": self.a.documento,
            "nombre": "No reemplazar", "desde": self.hoy.isoformat(),
            "motivo": "Alta manual", "solo_crear": True,
        }, format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("ya existe", str(response.data))
        self.a.refresh_from_db()
        self.assertEqual((self.a.numero, self.a.nombre, self.a.plan_id, self.a.desde, self.a.revision), antes)
        self.assertFalse(self.a.historial.exists())
        self.assertFalse(m.EventoCobertura.objects.filter(accion="actualizar_padron").exists())

    def test_alta_basica_preserva_motivo_y_rechaza_repeticion(self):
        payload = {"numero": "009", "documento": "12345678", "nombre": "Nueva persona",
                   "desde": self.hoy.isoformat(), "motivo": "Alta por oficina demo", "solo_crear": True}
        response = self.client.post(self.base + "padron/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        creada = m.Afiliado.objects.get(pk=response.data["id"])
        historial = self.client.get(self.base + f"padron/{creada.pk}/historial/")
        self.assertEqual(historial.status_code, 200, historial.data)
        self.assertEqual(historial.data["results"][0]["motivo"], payload["motivo"])
        self.assertEqual(creada.historial.get().registrado_por, self.user)
        payload["nombre"] = "No reemplazar"
        repetida = self.client.post(self.base + "padron/", payload, format="json")
        self.assertEqual(repetida.status_code, 400, repetida.data)
        creada.refresh_from_db()
        self.assertEqual(creada.nombre, "Nueva persona")
        self.assertEqual(creada.historial.count(), 1)

    def test_alta_basica_sin_campos_ampliados_conserva_motivo(self):
        payload = {"numero": "010", "documento": "12345679", "nombre": "Alta básica",
                   "desde": self.hoy.isoformat(), "motivo": "Documento revisado en oficina"}
        response = self.client.post(self.base + "padron/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        creada = m.Afiliado.objects.get(pk=response.data["id"])
        self.assertEqual(creada.historial.get().motivo, payload["motivo"])

    def test_post_legacy_sigue_permitiendo_actualizacion_incremental(self):
        response = self.client.post(self.base + "padron/", {
            "numero": self.a.numero, "documento": self.a.documento,
            "nombre": "Actualización legada", "desde": self.hoy.isoformat(),
        }, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.a.refresh_from_db()
        self.assertEqual(self.a.nombre, "Actualización legada")
        self.assertEqual(self.a.revision, 2)

    def test_plan_inactivo_no_modifica_ficha(self):
        plan = m.Plan.objects.create(financiador=self.org, codigo="INACTIVO", nombre="Inactivo", activo=False)
        response = self.patch({"revision": 1, "motivo": "Cambio de plan", "plan": plan.pk})
        self.assertEqual(response.status_code, 400, response.data)
        self.a.refresh_from_db()
        self.assertIsNone(self.a.plan_id)
        self.assertEqual(self.a.revision, 1)
        self.assertFalse(self.a.historial.exists())

    def test_plan_finalizado_requiere_reactivacion_explicita(self):
        plan = m.Plan.objects.create(financiador=self.org, codigo="NUEVO", nombre="Nuevo")
        finalizado = timezone.now()
        self.a.finalizado_en = finalizado
        self.a.save(update_fields=["finalizado_en"])
        response = self.patch({"revision": 1, "motivo": "Cambiar plan", "plan": plan.pk})
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("Reactivala", str(response.data))
        self.a.refresh_from_db()
        self.assertIsNone(self.a.plan_id)
        self.assertEqual(self.a.revision, 1)
        self.assertEqual(self.a.finalizado_en, finalizado)
        self.assertFalse(self.a.historial.exists())

    def test_correccion_plan_conserva_afiliacion_e_instantanea_anterior(self):
        anterior = m.Plan.objects.create(financiador=self.org, codigo="BASE", nombre="Base")
        nuevo = m.Plan.objects.create(financiador=self.org, codigo="SUPER", nombre="Superador")
        desde = self.hoy - timedelta(days=90)
        self.a.plan, self.a.desde = anterior, desde
        self.a.save(update_fields=["plan", "desde"])
        instantanea = m.HistorialAfiliacion.objects.create(afiliado=self.a, numero=self.a.numero,
            documento=self.a.documento, plan=anterior, desde=desde, registrado_por=self.user)
        response = self.patch({"revision": 1, "motivo": "Corrección administrativa", "plan": nuevo.pk})
        self.assertEqual(response.status_code, 200, response.data)
        self.a.refresh_from_db()
        instantanea.refresh_from_db()
        self.assertEqual(self.a.desde, desde)
        self.assertEqual(self.a.plan_id, nuevo.pk)
        self.assertEqual(instantanea.plan_id, anterior.pk)
        self.assertEqual(instantanea.desde, desde)
        historia = self.client.get(self.url + "historial/").data["results"][0]
        self.assertEqual(historia["cambios"]["plan_id"], {"anterior": anterior.pk, "nuevo": nuevo.pk})
        self.assertNotIn("desde", historia["cambios"])

    def test_observar_legajo_completo_requiere_volver_a_revision(self):
        response = self.client.post(self.url + "legajo/", {"revision": 1,
            "motivo": "Original visto", "tipo": "dni", "estado": "verificado",
            "presentado_en": self.hoy.isoformat(), "verificado_en": self.hoy.isoformat()}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        documento_id = response.data["documento"]["id"]
        response = self.patch({"revision": 2, "motivo": "Completar trámite",
            "credencial_tipo": "Externa", "credencial_fecha": self.hoy.isoformat(), "estado_tramite": "completo"})
        self.assertEqual(response.status_code, 200, response.data)
        payload = {"revision": 3, "motivo": "Observar original", "id": documento_id, "estado": "observado"}
        response = self.client.patch(self.url + "legajo/", payload, format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.a.refresh_from_db()
        self.assertEqual(self.a.revision, 3)
        self.assertEqual(m.DocumentoLegajo.objects.get(pk=documento_id).estado, "verificado")
        self.assertEqual(self.patch({"revision": 3, "motivo": "Revisar trámite", "estado_tramite": "en_revision"}).status_code, 200)
        payload["revision"] = 4
        response = self.client.patch(self.url + "legajo/", payload, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(m.DocumentoLegajo.objects.get(pk=documento_id).estado, "observado")

    def test_legado_contacto_persistencia_y_conflicto(self):
        detalle = self.client.get(self.url)
        self.assertEqual(detalle.status_code, 200, detalle.data)
        self.assertEqual(detalle.data["revision"], 1)
        self.assertEqual(detalle.data["cuil"], "")
        response = self.patch({"revision": 1, "motivo": "Corrección de contacto", "email": "persona@test.local", "calle": "Avenida Demo", "regimen": "dependencia"})
        self.assertEqual(response.status_code, 200, response.data)
        self.a.refresh_from_db()
        self.assertEqual(self.a.email, "persona@test.local")
        self.assertEqual(self.a.desde, self.hoy)
        self.assertIsNone(self.a.finalizado_en)
        self.assertEqual(self.client.get(self.url).data["calle"], "Avenida Demo")
        historia = self.client.get(self.url + "historial/").data
        self.assertEqual(historia["results"][0]["cambios"]["email"]["nuevo"], "persona@test.local")
        self.assertEqual(self.patch({"revision": 1, "motivo": "Otra edición", "telefono": "123"}).status_code, 409)
        self.assertEqual(self.patch({"revision": 2, "motivo": "Sin corte", "desde": self.hoy.isoformat()}).status_code, 400)

    def test_cuil_coherencia_y_unicidad(self):
        response = self.patch({"revision": 1, "motivo": "Identificación", "tipo_documento": "DNI", "cuil": "20329642330"})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.patch({"revision": 2, "motivo": "Cambio DNI", "documento": "11111111"}).status_code, 400)
        otro_url = self.base + f"padron/{self.b.pk}/"
        self.assertEqual(self.patch({"revision": 1, "motivo": "Duplicado", "cuil": "20329642330"}, otro_url).status_code, 400)
        self.assertEqual(self.patch({"revision": 2, "motivo": "Inválido", "cuil": "20329642331"}).status_code, 400)

    def test_familia_mismo_financiador_un_nivel_y_no_baja_cascada(self):
        url_b = self.base + f"padron/{self.b.pk}/"
        response = self.patch({"revision": 1, "motivo": "Familia informada", "titular": self.a.pk, "parentesco": "hijo"}, url_b)
        self.assertEqual(response.status_code, 200, response.data)
        familia = self.client.get(url_b + "familia/").data
        self.assertEqual({x["id"] for x in familia["integrantes"]}, {self.a.pk, self.b.pk})
        for titular in (self.a.pk, self.b.pk, self.ajeno.pk):
            self.assertEqual(self.patch({"revision": 1, "motivo": "Inválido", "titular": titular, "parentesco": "hijo"}).status_code, 400)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Baja titular", "baja_efectiva": self.hoy.isoformat(), "motivo_baja": "Renuncia"}).status_code, 200)
        self.b.refresh_from_db()
        self.assertIsNone(self.b.baja_efectiva)
        self.assertIsNone(self.b.finalizado_en)

    def test_fechas_administrativas_y_filtros_no_cortan_derechos(self):
        futura = self.hoy + timedelta(days=10)
        response = self.patch({"revision": 1, "motivo": "Programada", "alta_administrativa": self.hoy.isoformat(), "baja_efectiva": futura.isoformat(), "fin_prestacional_informado": futura.isoformat(), "motivo_baja": "Renuncia"})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["vigente"])
        self.assertEqual(response.data["estado_administrativo"], "baja_programada")
        listado = self.client.get(self.base + "padron/", {"estado_administrativo": "baja_programada", "alta_desde": self.hoy.isoformat(), "derecho": "vigentes"})
        self.assertEqual([x["id"] for x in listado.data["results"]], [self.a.pk])
        self.assertEqual(self.client.get(self.base + "padron/", {"alta_desde": futura.isoformat(), "alta_hasta": self.hoy.isoformat()}).status_code, 400)

    def test_legajo_metadata_y_tramite_explicito(self):
        self.assertEqual(self.patch({"revision": 1, "motivo": "Completar", "estado_tramite": "completo"}).status_code, 400)
        response = self.client.post(self.url + "legajo/", {"revision": 1, "motivo": "Original visto", "tipo": "dni", "estado": "verificado", "presentado_en": self.hoy.isoformat(), "verificado_en": self.hoy.isoformat(), "referencia": "Oficina 1"}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertFalse(response.data["documento"]["adjunto"])
        response = self.patch({"revision": 2, "motivo": "Completar", "credencial_tipo": "Externa", "credencial_fecha": self.hoy.isoformat(), "estado_tramite": "completo"})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["estado_administrativo"], "activo")
        self.assertEqual(self.client.post(self.url + "legajo/", {"revision": 3, "motivo": "CUD", "tipo": "cud"}, format="json").status_code, 400)

    def test_default_deny_autorizaciones_no_concede_y_plataforma_no_accede(self):
        self.member.resuelve_autorizaciones = True
        self.member.save()
        for bloque in ("salud", "cobranza"):
            self.assertEqual(self.client.get(self.url + bloque + "/").status_code, 403)
        admin = Usuario.objects.create_superuser("plataforma@test.local", "clave")
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.get(self.url + "salud/").status_code, 403)
        response = self.client.post(self.base + "usuarios/", {"email": admin.email, "nombre": "Plataforma", "rol": "admin", "acceso_padron_salud": True}, format="json")
        self.assertEqual(response.status_code, 403)
        self.assertFalse(m.MembresiaFinanciador.objects.filter(usuario=admin).exists())

    def test_designacion_real_independiente_revocacion_auditada(self):
        datos = {"email": self.user.email, "nombre": "Administrador", "rol": "admin", "acceso_padron_salud": True, "acceso_padron_cobranza": False}
        response = self.client.post(self.base + "usuarios/", datos, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(self.client.get(self.url + "salud/").status_code, 200)
        self.assertEqual(self.client.get(self.url + "cobranza/").status_code, 403)
        datos["acceso_padron_salud"] = False
        self.assertEqual(self.client.post(self.base + "usuarios/", datos, format="json").status_code, 201)
        self.assertEqual(self.client.get(self.url + "salud/").status_code, 403)
        self.assertEqual(m.EventoCobertura.objects.filter(accion="membresia_financiador").count(), 2)

    def test_sensibles_no_filtran_serializer_evento_historial_general(self):
        self.permisos()
        response = self.patch({"revision": 1, "motivo": "Corrección protegida", "cud": True, "cronicas": "Dato reservado", "parentesco_sensible": "discapacidad"}, self.url + "salud/")
        self.assertEqual(response.status_code, 200, response.data)
        protegido = self.client.get(self.url + "salud/historial/").data
        self.assertEqual(protegido["results"][0]["cambios"]["cronicas"]["nuevo"], "Dato reservado")
        for url in (self.url, self.base + "padron/", self.url + "historial/", self.url + "familia/"):
            salida = json.dumps(self.client.get(url).data, default=str)
            self.assertNotIn("Dato reservado", salida)
            self.assertNotIn("discapacidad", salida)
        evento = m.EventoCobertura.objects.get(accion="editar_padron_protegido")
        self.assertEqual(evento.motivo, "")
        self.a.refresh_from_db()
        self.assertEqual(self.a.revision, 1)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Conflicto", "cud": False}, self.url + "salud/").status_code, 409)

    def test_cobranza_validada_fuente_fecha_y_tarjeta_solo_emisor(self):
        self.permisos()
        payload = {"revision": 1, "motivo": "Dato declarado", "situacion": "deuda"}
        self.assertEqual(self.patch(payload, self.url + "cobranza/").status_code, 400)
        payload.update(fuente="Equipo de cobranza", informado_en=self.hoy.isoformat(), modalidad="debito_cbu", cbu="2850590940090418135201", cuit_titular="20329642330", titular_cuenta="Titular informado", dia_vencimiento=31)
        response = self.patch(payload, self.url + "cobranza/")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(response.data["estado_informado"])
        self.assertEqual(self.patch({"revision": 2, "motivo": "PAN", "pan": "123456"}, self.url + "cobranza/").status_code, 400)
        self.assertEqual(self.patch({"revision": 2, "motivo": "Día", "dia_vencimiento": 32}, self.url + "cobranza/").status_code, 400)
        self.assertEqual(self.patch({"revision": 2, "motivo": "CBU", "cbu": "2850590940090418135202"}, self.url + "cobranza/").status_code, 400)
        self.assertNotIn("cbu", self.client.get(self.url).data)

    def test_auditor_designado_lee_no_escribe_y_membresia_inactiva_denegada(self):
        self.permisos()
        self.member.rol = "auditor"
        self.member.save()
        for bloque in ("salud", "cobranza"):
            self.assertEqual(self.client.get(self.url + bloque + "/").status_code, 200)
            self.assertEqual(self.patch({"revision": 1, "motivo": "Intento"}, self.url + bloque + "/").status_code, 403)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Intento", "email": "nuevo@test.local"}).status_code, 403)
        self.member.activo = False
        self.member.save()
        self.assertEqual(self.client.get(self.url + "salud/").status_code, 404)

    def test_aislamiento_otro_financiador_todos_recursos(self):
        self.permisos()
        for sufijo in ("", "familia/", "historial/", "legajo/", "salud/", "cobranza/", "salud/historial/", "cobranza/historial/"):
            url = self.base + f"padron/{self.ajeno.pk}/" + sufijo
            self.assertEqual(self.client.get(url).status_code, 404, url)

    def test_importacion_v1_sigue_valida_nombre_auditado(self):
        archivo = self.libro([self.fila(**{"Nombre y apellido": "Nombre corregido"})], version=1)
        lote = self.preview(archivo)
        self.assertEqual(lote.version, 1)
        aplicar_importacion(importacion=lote, usuario=self.user)
        self.a.refresh_from_db()
        self.assertEqual(self.a.nombre, "Nombre corregido")
        h = self.a.historial.get()
        self.assertEqual(h.cambios["nombre"]["nuevo"], "Nombre corregido")
        self.assertEqual(h.lote, lote)

    def test_v2_preview_sin_efectos_aplicacion_blancos_e_idempotencia(self):
        self.a.email = "conservar@test.local"
        self.a.save()
        lote = self.preview(self.libro([self.fila(cuil="20329642330", tipo_documento="DNI", regimen="pasivo")]))
        self.assertEqual(lote.resumen["valida"], 1, lote.filas)
        self.a.refresh_from_db()
        self.assertEqual(self.a.cuil, "")
        aplicar_importacion(importacion=lote, usuario=self.user)
        self.assertEqual(lote.filas[0]["estado"], "aplicada", lote.filas)
        self.a.refresh_from_db()
        self.assertEqual(self.a.cuil, "20329642330")
        self.assertEqual(self.a.email, "conservar@test.local")
        count = self.a.historial.count()
        aplicar_importacion(importacion=lote, usuario=self.user)
        self.assertEqual(self.a.historial.count(), count)

    def test_v2_revalida_conflicto_sin_override_identidad(self):
        lote = self.preview(self.libro([self.fila(telefono="123")]))
        self.assertEqual(self.patch({"revision": 1, "motivo": "Edición concurrente", "telefono": "456"}).status_code, 200)
        aplicar_importacion(importacion=lote, usuario=self.user, revisiones_duplicados={"2": "Forzar"})
        self.assertEqual(lote.filas[0]["estado"], "rechazada", lote.filas)
        self.a.refresh_from_db()
        self.assertEqual(self.a.telefono, "456")

    def test_v2_titular_en_fila_anterior_y_no_inferencias(self):
        titular = self.fila(**{"N.º de afiliado": "NUEVO", "Documento": "12345678", "Nombre y apellido": "Titular nuevo"})
        familiar = self.fila(**{"N.º de afiliado": "FAMILIAR", "Documento": "12345679", "Nombre y apellido": "Familiar nuevo", "titular_documento": "12345678", "parentesco": "hijo"})
        lote = self.preview(self.libro([titular, familiar]))
        self.assertEqual(lote.resumen["valida"], 2, lote.filas)
        aplicar_importacion(importacion=lote, usuario=self.user)
        self.assertEqual(lote.resumen["aplicada"], 2, lote.filas)
        familiar_obj = m.Afiliado.objects.get(documento="12345679")
        self.assertEqual(familiar_obj.titular.documento, "12345678")

    def test_v2_rechaza_cuil_finalizado_y_familia_sin_titular(self):
        self.a.finalizado_en = timezone.now()
        self.a.save()
        casos = [self.fila(), self.fila(**{"Documento": "12345678", "cuil": "20329642331"}), self.fila(**{"Documento": "12345679", "titular_documento": "99999999", "parentesco": "hijo"})]
        lote = self.preview(self.libro(casos))
        self.assertEqual(lote.resumen["rechazada"], 3, lote.filas)
        self.assertEqual(m.Afiliado.objects.count(), 3)

    def test_v2_no_almacena_encabezados_sensibles_ni_parentesco_rechazado(self):
        with self.assertRaises(ValidationError):
            self.preview(self.libro([self.fila()], encabezado="cbu"))
        self.assertFalse(m.Importacion.objects.exists())
        lote = self.preview(self.libro([self.fila(parentesco="discapacidad")]))
        self.assertEqual(lote.filas[0]["estado"], "rechazada")
        self.assertNotIn("discapacidad", json.dumps(lote.filas))
        self.assertIn("omitido", json.dumps(lote.filas))

    def test_edicion_v1_invalida_revision_v2(self):
        registrar_afiliado(financiador=self.org, usuario=self.user, numero=self.a.numero, documento=self.a.documento, nombre="Nombre v1", plan=None, desde=self.hoy)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Vieja", "telefono": "123"}).status_code, 409)

    def test_alta_post_ampliada_atomica_y_legacy_sin_nuevos_obligatorios(self):
        payload = {"numero": "009", "documento": "12345678", "nombre": "Nueva persona", "desde": self.hoy.isoformat(), "tipo_documento": "DNI", "cuil": cuit("12345678"), "email": "nueva@test.local", "regimen": "monotributo", "titular": self.a.pk, "parentesco": "hija"}
        response = self.client.post(self.base + "padron/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        creada = m.Afiliado.objects.get(documento="12345678")
        self.assertEqual(creada.cuil, payload["cuil"])
        self.assertEqual(creada.titular_id, self.a.pk)
        payload.update(documento="12345679", cuil="20329642331")
        self.assertEqual(self.client.post(self.base + "padron/", payload, format="json").status_code, 400)
        self.assertFalse(m.Afiliado.objects.filter(documento="12345679").exists())
        legacy = self.client.post(self.base + "padron/", {"numero": "010", "documento": "12345680", "nombre": "Sin nuevos datos", "desde": self.hoy.isoformat()}, format="json")
        self.assertEqual(legacy.status_code, 201, legacy.data)

    def test_historial_simulado_atribuye_autor_real_sin_conceder_permiso_real(self):
        root = Usuario.objects.create_superuser("root130@test.local", "clave")
        sesion = SimpleNamespace(cuenta=self.user)
        contexto.activar(sesion, root)
        try:
            self.permisos()
            p.editar_ficha(afiliado=self.a, usuario=self.user, entrada={"revision": 1, "motivo": "Simulación", "telefono": "123"})
            p.editar_protegido(afiliado=self.a, usuario=self.user, bloque="salud", entrada={"revision": 1, "motivo": "Protegido", "cud": True})
        finally:
            contexto.limpiar()
        self.assertEqual(self.a.historial.get().registrado_por, root)
        self.assertEqual(m.HistorialPadronProtegido.objects.get().registrado_por, root)
        self.assertEqual(m.EventoCobertura.objects.get(accion="editar_padron_protegido").usuario, root)

    def test_v2_no_afecta_omitidos_y_revalida_titular_al_confirmar(self):
        lote = self.preview(self.libro([self.fila(**{"Documento": "12345678", "N.º de afiliado": "009", "titular_documento": self.a.documento, "parentesco": "hijo"})]))
        self.assertEqual(lote.resumen["valida"], 1, lote.filas)
        # Tras preview el titular pasa a ser familiar: la confirmación debe rechazar.
        self.assertEqual(self.patch({"revision": 1, "motivo": "Nuevo vínculo", "titular": self.b.pk, "parentesco": "hijo"}).status_code, 200)
        aplicar_importacion(importacion=lote, usuario=self.user)
        self.assertEqual(lote.resumen["rechazada"], 1, lote.filas)
        self.assertFalse(m.Afiliado.objects.filter(documento="12345678").exists())
        self.b.refresh_from_db()
        self.assertIsNone(self.b.finalizado_en)

    def test_v1_finalizacion_y_reactivacion_invalidan_revision(self):
        response = self.client.post(self.base + "finalizar-afiliacion/", {"afiliado": self.a.pk, "motivo": "Cierre explícito"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Vieja", "telefono": "123"}).status_code, 409)
        response = self.client.post(self.base + "reactivar-afiliacion/", {"afiliado": self.a.pk, "plan": None, "motivo": "Reactivación explícita"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.patch({"revision": 2, "motivo": "Vieja", "telefono": "123"}).status_code, 409)

    def test_plan_tipo_y_busqueda_cuil_paginacion(self):
        response = self.client.post(self.base + "planes/", {"codigo": "SUPER", "nombre": "Superador", "tipo": "superador"}, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(self.patch({"revision": 1, "motivo": "Plan", "plan": response.data["id"], "cuil": "20329642330"}).status_code, 200)
        listado = self.client.get(self.base + "padron/", {"search": "20329642330", "page_size": 1})
        self.assertEqual(listado.data["count"], 1)
        self.assertEqual(listado.data["results"][0]["plan_tipo"], "superador")

    def test_datos_ilegales_en_general_y_numericos_no_se_silencian(self):
        self.assertEqual(self.patch({"revision": 1, "motivo": "Banco", "cbu": "2850590940090418135201"}).status_code, 400)
        self.assertEqual(self.patch({"revision": 1, "motivo": "CUIL", "cuil": 20329642330}).status_code, 400)
        self.assertFalse(self.a.historial.exists())

    def test_capability_gestion_no_depende_rol_sintetico_plataforma(self):
        response = self.client.get('/api/financiadores/')
        self.assertTrue(response.data['results'][0]['gestiona_permisos_padron'])
        self.assertTrue(self.client.get(self.base).data['gestiona_permisos_padron'])
        root = Usuario.objects.create_superuser('capability130@test.local', 'clave')
        self.client.force_authenticate(root)
        response = self.client.get('/api/financiadores/')
        self.assertTrue(all(not org['gestiona_permisos_padron'] for org in response.data['results']))
        self.assertFalse(self.client.get(self.base).data['gestiona_permisos_padron'])
        self.member.rol = 'operador'
        self.member.save()
        self.client.force_authenticate(self.user)
        self.assertFalse(self.client.get('/api/financiadores/').data['results'][0]['gestiona_permisos_padron'])

    def test_emisor_tarjeta_rechaza_pan_y_cvv_camuflados(self):
        self.permisos()
        for emisor in ('4111111111111111', 'Visa 4111 1111 1111 1111', 'CVV 123', '123'):
            response = self.patch({'revision': 1, 'motivo': 'Registro tarjeta', 'emisor_tarjeta': emisor}, self.url + 'cobranza/')
            self.assertEqual(response.status_code, 400)
        self.assertFalse(m.CobranzaPadron.objects.exists())

    def test_servicio_no_acepta_campos_prestacionales_ni_ownership(self):
        with self.assertRaises(ValidationError):
            p.editar_ficha(afiliado=self.a, usuario=self.user, entrada={'revision': 1, 'motivo': 'Intento', 'finalizado_en': timezone.now()})
        self.permisos()
        with self.assertRaises(ValidationError):
            p.editar_protegido(afiliado=self.a, usuario=self.user, bloque='salud', entrada={'revision': 1, 'motivo': 'Intento', 'afiliado_id': self.ajeno.pk})
