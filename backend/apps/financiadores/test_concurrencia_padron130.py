"""Pruebas que requieren PostgreSQL dedicado; SQLite no acredita locks reales."""
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event
from queue import Queue
from time import monotonic, sleep
from unittest import skipUnless
from unittest.mock import patch
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from django.core.exceptions import ValidationError
from django.db import DatabaseError, close_old_connections, connection, transaction
from django.test import TransactionTestCase
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Usuario
from .models import Afiliado, Financiador, MembresiaFinanciador, Plan
from . import padron
from .services import registrar_afiliado
from .padron import ConflictoPadron, editar_ficha
from .importaciones import aplicar_importacion, generar_plantilla, previsualizar_importacion


@skipUnless(connection.vendor == "postgresql", "Requiere locks reales en PostgreSQL local dedicado; no ejecutado en SQLite.")
class ConcurrenciaPadron130Tests(TransactionTestCase):
    def setUp(self):
        self.org = Financiador.objects.create(nombre="A", tipo="obra_social")
        self.usuario = Usuario.objects.create_user("concurrencia130@test.local", "clave")
        MembresiaFinanciador.objects.create(financiador=self.org, usuario=self.usuario, rol="operador")
        self.a = Afiliado.objects.create(financiador=self.org, numero="A", documento="10000001", nombre="A", desde=timezone.localdate())
        self.b = Afiliado.objects.create(financiador=self.org, numero="B", documento="10000002", nombre="B", desde=timezone.localdate())

    def test_vinculos_inversos_simultaneos_no_crean_ciclo(self):
        barrera = Barrier(2)
        def vincular(afiliado, titular):
            close_old_connections()
            try:
                barrera.wait(timeout=10)
                try:
                    editar_ficha(afiliado=afiliado, usuario=self.usuario, entrada={"revision": 1, "motivo": "Familia", "titular": titular.pk, "parentesco": "hijo"})
                    return "aceptado"
                except (ValidationError, ConflictoPadron):
                    return "rechazado"
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as pool:
            primero = pool.submit(vincular, self.a, self.b)
            segundo = pool.submit(vincular, self.b, self.a)
            resultados = [primero.result(timeout=20), segundo.result(timeout=20)]
        self.assertCountEqual(resultados, ["aceptado", "rechazado"])
        self.a.refresh_from_db()
        self.b.refresh_from_db()
        self.assertFalse(self.a.titular_id and self.b.titular_id)

    def test_plan_desactivado_mientras_se_asigna_se_revalida_bajo_lock(self):
        plan = Plan.objects.create(financiador=self.org, codigo="BASE", nombre="Base")
        validando, continuar = Event(), Event()
        pid_editor = Queue(maxsize=1)
        validar_original = padron.validar_generales

        def validar_pausado(obj):
            validando.set()
            if not continuar.wait(timeout=15):
                raise RuntimeError("No se liberó la edición de la prueba.")
            return validar_original(obj)

        def asignar():
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '15s'")
                    cursor.execute("SELECT pg_backend_pid()")
                    pid_editor.put(cursor.fetchone()[0])
                try:
                    editar_ficha(afiliado=self.a, usuario=self.usuario, entrada={
                        "revision": 1, "motivo": "Asignación concurrente", "plan": plan.pk,
                    })
                    return "aceptado"
                except ValidationError:
                    return "rechazado"
            finally:
                connection.close()

        with ThreadPoolExecutor(max_workers=1) as pool, patch.object(padron, "validar_generales", side_effect=validar_pausado):
            try:
                with transaction.atomic():
                    bloqueado = Plan.objects.select_for_update().get(pk=plan.pk)
                    bloqueado.activo = False
                    bloqueado.save(update_fields=["activo"])
                    with connection.cursor() as cursor:
                        cursor.execute("SELECT pg_backend_pid()")
                        pid_desactivador = cursor.fetchone()[0]
                    futuro = pool.submit(asignar)
                    pid = pid_editor.get(timeout=10)
                    limite = monotonic() + 10
                    while True:
                        with connection.cursor() as cursor:
                            cursor.execute("SELECT pg_blocking_pids(%s)", [pid])
                            bloqueadores = cursor.fetchone()[0]
                        # Con el fix espera el lock; sin él llega a validar con
                        # activo=True antes del commit. Ambos órdenes se prueban.
                        if pid_desactivador in bloqueadores or validando.is_set():
                            break
                        if monotonic() >= limite:
                            self.fail("La asignación no alcanzó el punto de sincronización.")
                        sleep(0.02)
                # Confirmar la desactivación antes de permitir guardar al editor.
            finally:
                continuar.set()
            self.assertEqual(futuro.result(timeout=20), "rechazado")
        self.a.refresh_from_db()
        self.assertIsNone(self.a.plan_id)
        self.assertEqual(self.a.revision, 1)
        self.assertFalse(self.a.historial.exists())

    def test_altas_exclusivas_simultaneas_no_sobrescriben_ganador(self):
        barrera = Barrier(2)
        crear_original = Afiliado.objects.get_or_create

        def crear_sincronizado(*args, **kwargs):
            barrera.wait(timeout=10)
            return crear_original(*args, **kwargs)

        def registrar(nombre):
            close_old_connections()
            try:
                with connection.cursor() as cursor:
                    cursor.execute("SET lock_timeout = '15s'")
                try:
                    registrar_afiliado(financiador=self.org, usuario=self.usuario,
                        numero=nombre, documento="10000003", nombre=nombre,
                        plan=None, desde=timezone.localdate(), motivo="Alta exclusiva",
                        solo_crear=True)
                    return nombre
                except ValidationError:
                    return None
            finally:
                connection.close()

        with ThreadPoolExecutor(max_workers=2) as pool, patch.object(Afiliado.objects, "get_or_create", side_effect=crear_sincronizado):
            futuros = [pool.submit(registrar, nombre) for nombre in ("Uno", "Dos")]
            resultados = [futuro.result(timeout=20) for futuro in futuros]
        ganadores = [nombre for nombre in resultados if nombre is not None]
        self.assertEqual(len(ganadores), 1)
        afiliado = Afiliado.objects.get(financiador=self.org, documento="10000003")
        self.assertEqual((afiliado.nombre, afiliado.numero), (ganadores[0], ganadores[0]))
        self.assertEqual(afiliado.revision, 1)
        self.assertEqual(afiliado.historial.count(), 1)
        self.assertEqual(afiliado.historial.get().motivo, "Alta exclusiva")

    def test_contacto_y_familia_no_bloquean_el_commit_de_auditoria(self):
        registrado, continuar = Event(), Event()
        original = padron.registrar_historial

        def pausar(obj, usuario, delta, motivo, *args, **kwargs):
            resultado = original(obj, usuario, delta, motivo, *args, **kwargs)
            if motivo == "Contacto concurrente":
                registrado.set()
                if not continuar.wait(timeout=10):
                    raise RuntimeError("No se liberó el commit del contacto.")
            return resultado

        def contacto():
            close_old_connections()
            try:
                editar_ficha(afiliado=self.a, usuario=self.usuario, entrada={
                    "revision": 1, "motivo": "Contacto concurrente", "telefono": "111",
                })
            finally:
                connection.close()

        def liberar(execute, sql, params, many, context):
            if '"financiadores_afiliado"' in sql and "FOR UPDATE" in sql:
                continuar.set()
            return execute(sql, params, many, context)

        with ThreadPoolExecutor(max_workers=1) as pool, patch.object(padron, "registrar_historial", side_effect=pausar):
            futuro = pool.submit(contacto)
            try:
                self.assertTrue(registrado.wait(timeout=10))
                with connection.execute_wrapper(liberar):
                    editar_ficha(afiliado=self.b, usuario=self.usuario, entrada={
                        "revision": 1, "motivo": "Familia concurrente", "titular": self.a.pk, "parentesco": "hijo",
                    })
                futuro.result(timeout=15)
            finally:
                continuar.set()
        self.a.refresh_from_db()
        self.b.refresh_from_db()
        self.assertEqual(self.a.telefono, "111")
        self.assertEqual(self.b.titular_id, self.a.pk)

    def test_revocacion_y_escritura_protegida_no_bloquean_auditoria(self):
        miembro = MembresiaFinanciador.objects.get(financiador=self.org, usuario=self.usuario)
        miembro.rol, miembro.acceso_padron_salud = "admin", True
        miembro.save(update_fields=["rol", "acceso_padron_salud"])
        registrado, continuar = Event(), Event()
        original = padron.auditar

        def pausar(*args, **kwargs):
            resultado = original(*args, **kwargs)
            registrado.set()
            if not continuar.wait(timeout=10):
                raise RuntimeError("No se liberó el commit protegido.")
            return resultado

        def escribir():
            close_old_connections()
            try:
                padron.editar_protegido(afiliado=self.a, usuario=self.usuario, bloque="salud",
                    entrada={"revision": 1, "motivo": "Escritura concurrente", "dbt": True})
            finally:
                connection.close()

        def liberar(execute, sql, params, many, context):
            if sql.startswith("UPDATE") and '"financiadores_membresiafinanciador"' in sql:
                continuar.set()
            return execute(sql, params, many, context)

        cliente = APIClient()
        cliente.force_authenticate(self.usuario)
        with ThreadPoolExecutor(max_workers=1) as pool, patch.object(padron, "auditar", side_effect=pausar):
            futuro = pool.submit(escribir)
            try:
                self.assertTrue(registrado.wait(timeout=10))
                with connection.execute_wrapper(liberar):
                    response = cliente.post(f"/api/financiadores/{self.org.pk}/usuarios/", {
                        "email": self.usuario.email, "nombre": "Concurrente", "rol": "admin",
                        "acceso_padron_salud": False,
                    }, format="json")
                self.assertEqual(response.status_code, 201, response.data)
                futuro.result(timeout=15)
            finally:
                continuar.set()
        miembro.refresh_from_db()
        self.assertFalse(miembro.acceso_padron_salud)
        self.assertTrue(padron.m.SaludPadron.objects.get(afiliado=self.a).dbt)
        self.assertEqual(cliente.get(f"/api/financiadores/{self.org.pk}/padron/{self.a.pk}/salud/").status_code, 403)

    def _probar_importacion_y_contacto(self, version):
        errores_sql = []
        libro = load_workbook(BytesIO(generar_plantilla(
            financiador=self.org, usuario=self.usuario, tipo="padron", version=version,
        )))
        hoja = libro["Carga"]
        fila = {"N.º de afiliado": self.a.numero, "Documento": self.a.documento,
                "Nombre y apellido": "A importado", "Vigente desde": self.a.desde.isoformat()}
        if version == 2:
            fila.update(titular_documento=self.b.documento, parentesco="hijo")
        for columna, celda in enumerate(hoja[1], 1):
            if celda.value in fila:
                hoja.cell(2, columna, fila[celda.value])
        salida = BytesIO()
        libro.save(salida)
        lote = previsualizar_importacion(financiador=self.org, usuario=self.usuario,
            tipo="padron", archivo=salida.getvalue(), clave=uuid4())
        self.assertEqual(lote.resumen["valida"], 1)

        registrado, continuar = Event(), Event()
        original = padron.registrar_historial

        def pausar(obj, usuario, delta, motivo, *args, **kwargs):
            resultado = original(obj, usuario, delta, motivo, *args, **kwargs)
            if motivo == "Contacto concurrente con importación":
                registrado.set()
                if not continuar.wait(timeout=10):
                    raise RuntimeError("No se liberó el commit del contacto.")
            return resultado

        def contacto():
            close_old_connections()
            try:
                editar_ficha(afiliado=self.a, usuario=self.usuario, entrada={
                    "revision": 1, "motivo": "Contacto concurrente con importación", "telefono": "111",
                })
            finally:
                connection.close()

        def liberar(execute, sql, params, many, context):
            if '"financiadores_afiliado"' in sql and "FOR UPDATE" in sql:
                continuar.set()
            try:
                return execute(sql, params, many, context)
            except DatabaseError as exc:
                errores_sql.append(type(exc.__cause__ or exc).__name__)
                raise

        with ThreadPoolExecutor(max_workers=1) as pool, patch.object(padron, "registrar_historial", side_effect=pausar):
            futuro = pool.submit(contacto)
            try:
                self.assertTrue(registrado.wait(timeout=10))
                with connection.execute_wrapper(liberar):
                    aplicar_importacion(importacion=lote, usuario=self.usuario)
                futuro.result(timeout=15)
            finally:
                continuar.set()
        lote.refresh_from_db()
        self.a.refresh_from_db()
        self.assertEqual(errores_sql, [])
        self.assertEqual(self.a.telefono, "111")
        if version == 2:
            self.assertEqual(lote.filas[0]["estado"], "rechazada")
            self.assertEqual(lote.filas[0]["errores"], [
                "El afiliado cambió desde la vista previa; generá una nueva importación para revisar el conflicto.",
            ])
            self.assertEqual((self.a.nombre, self.a.titular_id, self.a.revision), ("A", None, 2))
            self.assertEqual(self.a.historial.count(), 1)
            nuevo = previsualizar_importacion(financiador=self.org, usuario=self.usuario,
                tipo="padron", archivo=salida.getvalue(), clave=uuid4())
            self.assertEqual(nuevo.resumen["valida"], 1)
            aplicar_importacion(importacion=nuevo, usuario=self.usuario)
            nuevo.refresh_from_db()
            self.a.refresh_from_db()
            self.assertEqual(nuevo.filas[0]["estado"], "aplicada", nuevo.filas[0]["errores"])
            self.assertEqual((self.a.titular_id, self.a.revision), (self.b.pk, 4))
        else:
            self.assertEqual(lote.filas[0]["estado"], "aplicada", lote.filas[0]["errores"])
            self.assertEqual((self.a.titular_id, self.a.revision), (None, 3))
        self.assertEqual(self.a.nombre, "A importado")
        self.assertEqual(self.a.telefono, "111")

    def test_importacion_v1_y_contacto_no_bloquean_auditoria(self):
        self._probar_importacion_y_contacto(1)

    def test_importacion_v2_y_contacto_no_bloquean_auditoria(self):
        self._probar_importacion_y_contacto(2)

    def test_ediciones_misma_revision_solo_una_confirma(self):
        barrera = Barrier(2)
        def editar(telefono):
            close_old_connections()
            try:
                barrera.wait(timeout=10)
                try:
                    editar_ficha(afiliado=self.a, usuario=self.usuario, entrada={"revision": 1, "motivo": "Contacto", "telefono": telefono})
                    return "aceptado"
                except ConflictoPadron:
                    return "conflicto"
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as pool:
            primero = pool.submit(editar, "111")
            segundo = pool.submit(editar, "222")
            resultados = [primero.result(timeout=20), segundo.result(timeout=20)]
        self.assertCountEqual(resultados, ["aceptado", "conflicto"])
        self.a.refresh_from_db()
        self.assertEqual(self.a.revision, 2)

    def test_importacion_existente_y_vinculo_inverso_no_invierten_locks(self):
        libro = load_workbook(BytesIO(generar_plantilla(financiador=self.org, usuario=self.usuario, tipo="padron", version=2)))
        hoja = libro["Carga"]
        fila = {"N.º de afiliado": self.a.numero, "Documento": self.a.documento,
                "Nombre y apellido": self.a.nombre, "Vigente desde": self.a.desde.isoformat(),
                "titular_documento": self.b.documento, "parentesco": "hijo"}
        for columna, celda in enumerate(hoja[1], 1):
            if celda.value in fila:
                hoja.cell(2, columna, fila[celda.value])
        salida = BytesIO()
        libro.save(salida)
        lote = previsualizar_importacion(financiador=self.org, usuario=self.usuario, tipo="padron", archivo=salida.getvalue(), clave=uuid4())
        self.assertEqual(lote.resumen["valida"], 1)
        barrera = Barrier(2)
        def importar():
            close_old_connections()
            try:
                barrera.wait(timeout=10)
                aplicar_importacion(importacion=lote, usuario=self.usuario)
                return lote.filas[0]["estado"]
            finally:
                close_old_connections()
        def vincular():
            close_old_connections()
            try:
                barrera.wait(timeout=10)
                try:
                    editar_ficha(afiliado=self.b, usuario=self.usuario, entrada={"revision": 1, "motivo": "Vínculo inverso", "titular": self.a.pk, "parentesco": "hijo"})
                    return "aplicada"
                except (ValidationError, ConflictoPadron):
                    return "rechazada"
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as pool:
            primero = pool.submit(importar)
            segundo = pool.submit(vincular)
            resultados = [primero.result(timeout=20), segundo.result(timeout=20)]
        self.assertCountEqual(resultados, ["aplicada", "rechazada"])
        self.a.refresh_from_db()
        self.b.refresh_from_db()
        self.assertFalse(self.a.titular_id and self.b.titular_id)
