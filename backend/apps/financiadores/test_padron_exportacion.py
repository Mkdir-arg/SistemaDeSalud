from datetime import date
from io import BytesIO
from unittest.mock import patch

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import Usuario
from . import models as m
from .padron_consulta import COLUMNAS_PADRON


class ExportacionPadronTests(APITestCase):
    def setUp(self):
        self.org = m.Financiador.objects.create(nombre="Organización ficticia", tipo="mutual")
        self.otra = m.Financiador.objects.create(nombre="Otra organización ficticia", tipo="mutual")
        self.user = Usuario.objects.create_user("exportacion@example.test", "clave-test")
        self.member = m.MembresiaFinanciador.objects.create(financiador=self.org, usuario=self.user, rol="admin")
        self.a = m.Afiliado.objects.create(financiador=self.org, numero="00025", documento="00123456", nombre="=SUM(1,2)", desde=timezone.localdate(), cuil="00000000001", localidad="Árbol\x01 Unicode", telefono="12345678901234567890", fecha_nacimiento=date(1990, 1, 2))
        m.Afiliado.objects.create(financiador=self.otra, numero="secreto-ajeno", documento="99999999", nombre="Ajeno", desde=timezone.localdate())
        m.SaludPadron.objects.create(afiliado=self.a, cronicas="SALUD-NO-EXPORTABLE")
        m.CobranzaPadron.objects.create(afiliado=self.a, banco="BANCO-NO-EXPORTABLE")
        self.url = f"/api/financiadores/{self.org.pk}/padron/exportar/"
        self.client.force_authenticate(self.user)

    def libro(self, parametros=None):
        respuesta = self.client.get(self.url, parametros or {})
        self.assertEqual(respuesta.status_code, 200, getattr(respuesta, "data", ""))
        self.assertIn("no-store", [v.strip() for v in respuesta["Cache-Control"].split(",")])
        self.assertEqual(respuesta["X-Content-Type-Options"], "nosniff")
        self.assertIn("application/vnd.openxmlformats", respuesta["Content-Type"])
        self.assertIn(".xlsx", respuesta["Content-Disposition"])
        return load_workbook(BytesIO(respuesta.content), data_only=False)

    def test_textos_identificadores_fechas_null_sin_formulas_ni_enlaces(self):
        columnas = ("nombre", "numero", "documento", "cuil", "telefono", "localidad", "fecha_nacimiento", "alta_administrativa")
        libro = self.libro({"columnas": ",".join(columnas)})
        hoja = libro["Padrón"]
        self.assertEqual([c.value for c in hoja[1]], [COLUMNAS_PADRON[k]["label"] for k in columnas])
        self.assertEqual([c.value for c in hoja[2]][:6], ["=SUM(1,2)", "00025", "00123456", "00000000001", "12345678901234567890", "Árbol Unicode"])
        self.assertTrue(all(c.data_type == "s" and c.number_format == "@" for c in list(hoja[2])[:6]))
        self.assertEqual(hoja.cell(2, 7).value.date(), date(1990, 1, 2))
        self.assertIsNone(hoja.cell(2, 8).value)
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(hoja.auto_filter.ref, "A1:H2")
        self.assertFalse(libro._external_links)
        self.assertTrue(all(c.data_type != "f" and c.hyperlink is None for row in hoja for c in row))

    def test_exporta_todo_mas_que_pagina_mismo_orden_y_filtro(self):
        for i in range(12):
            m.Afiliado.objects.create(financiador=self.org, numero=f"000{i}", documento=f"800000{i:02}", nombre=f"Ficticio {i:02}", desde=timezone.localdate())
        params = {"search": "Ficticio", "ordering": "-nombre", "columnas": "nombre,numero,documento", "page": 2, "page_size": 3}
        listado = self.client.get(f"/api/financiadores/{self.org.pk}/padron/", params)
        self.assertEqual(listado.status_code, 200)
        self.assertEqual(listado.data["count"], 12)
        self.assertEqual(len(listado.data["results"]), 3)
        filas = list(self.libro(params)["Padrón"].values)[1:]
        self.assertEqual(len(filas), 12)
        self.assertEqual([r[0] for r in filas], [f"Ficticio {i:02}" for i in reversed(range(12))])
        self.assertEqual([r[2] for r in filas[3:6]], [r["documento"] for r in listado.data["results"]])
        evento = m.EventoCobertura.objects.get(accion="exportar_padron")
        self.assertEqual(evento.usuario, self.user)
        self.assertEqual(evento.financiador, self.org)
        self.assertIn("registros=12", evento.motivo)
        self.assertIn("columnas=nombre,numero,documento", evento.motivo)
        self.assertNotIn("Ficticio", evento.motivo)

    def test_roles_consulta_scope_y_revocacion(self):
        for rol in ("admin", "operador", "auditor"):
            self.member.rol = rol; self.member.save()
            filas = list(self.libro()["Padrón"].values)
            self.assertEqual(len(filas), 2)
            self.assertNotIn("secreto-ajeno", str(filas))
        # get_queryset oculta organizaciones sin membresía antes de autorizar.
        self.assertEqual(self.client.get(f"/api/financiadores/{self.otra.pk}/padron/exportar/").status_code, 404)
        self.member.activo = False; self.member.save()
        self.assertEqual(self.client.get(self.url).status_code, 404)
        self.client.force_authenticate(None)
        self.assertIn(self.client.get(self.url).status_code, (401, 403))

    def test_plantilla_vigente_nombre_y_legacy_compatible(self):
        url = f"/api/financiadores/{self.org.pk}/plantilla/"
        for version, nombre in ((2, "plantilla-padron.xlsx"), (1, "plantilla.xlsx")):
            respuesta = self.client.get(url, {"tipo": "padron", "version": version})
            self.assertEqual(respuesta.status_code, 200)
            self.assertIn(nombre, respuesta["Content-Disposition"])
            self.assertIn("Carga", load_workbook(BytesIO(respuesta.content)).sheetnames)

    def test_allowlist_parametros_y_formato_invalidos_no_auditan_descarga(self):
        casos = [{"formato": "csv"}, {"columnas": "nombre,numero,cbu"}, {"columnas": "nombre,numero,cronicas"}, {"columnas": "nombre,numero,parentesco_sensible"}, {"columnas": "nombre,numero,acciones"}, {"ordering": "salud__cronicas"}, {"flt_cbu": "algo"}, {"alta_desde": "2026-09-10", "alta_hasta": "2026-09-01"}, {"financiador": self.otra.pk}, {"columnas": "nombre,numero,numero"}]
        for parametros in casos:
            with self.subTest(parametros=parametros):
                self.assertEqual(self.client.get(self.url, parametros).status_code, 400)
        self.assertFalse(m.EventoCobertura.objects.filter(accion="exportar_padron").exists())

    def test_todas_columnas_solo_administracion_y_motivo_acotado(self):
        columnas = ",".join(COLUMNAS_PADRON)
        hoja = self.libro({"columnas": columnas})["Padrón"]
        self.assertEqual(len(list(hoja.values)[0]), len(COLUMNAS_PADRON))
        self.assertNotIn("SALUD-NO-EXPORTABLE", str(list(hoja.values)))
        self.assertNotIn("BANCO-NO-EXPORTABLE", str(list(hoja.values)))
        self.assertLessEqual(len(m.EventoCobertura.objects.get(accion="exportar_padron").motivo), 255)

    def test_cero_resultados_contiene_encabezados_y_audita(self):
        hoja = self.libro({"search": "sin-resultados"})["Padrón"]
        self.assertEqual(len(list(hoja.values)), 1)
        self.assertIn("registros=0", m.EventoCobertura.objects.get(accion="exportar_padron").motivo)

    def test_limite_5001_no_trunca_ni_audita(self):
        m.Afiliado.objects.bulk_create([m.Afiliado(financiador=self.org, numero=str(i), documento=f"700{i:05}", nombre="Ficticio", desde=timezone.localdate()) for i in range(5000)])
        respuesta = self.client.get(self.url)
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn("5000", str(respuesta.data))
        self.assertFalse(m.EventoCobertura.objects.filter(accion="exportar_padron").exists())

    def test_fallo_auditoria_no_entrega_archivo(self):
        with patch("apps.financiadores.padron_exportacion.auditar", side_effect=RuntimeError("auditoría no disponible")):
            with self.assertRaisesRegex(RuntimeError, "auditoría no disponible"):
                self.client.get(self.url)
        self.assertFalse(m.EventoCobertura.objects.filter(accion="exportar_padron").exists())

    def test_texto_excesivo_no_trunca(self):
        with patch("apps.financiadores.padron_exportacion.valor_columna_padron", return_value="a" * 32768):
            self.assertEqual(self.client.get(self.url).status_code, 400)
        self.assertFalse(m.EventoCobertura.objects.filter(accion="exportar_padron").exists())

    def test_limite_exacto_y_prefijos_peligrosos_son_texto(self):
        for valor in ("=1+1", "+123", "-123", "@SUM(A1)", " \t=1+1", "https://example.test"):
            with self.subTest(valor=valor):
                m.Afiliado.objects.filter(pk=self.a.pk).update(nombre=valor)
                celda = self.libro()["Padrón"].cell(2, 1)
                self.assertEqual(celda.value, valor)
                self.assertEqual(celda.data_type, "s")
                self.assertIsNone(celda.hyperlink)
        m.Afiliado.objects.bulk_create([m.Afiliado(financiador=self.org, numero=str(i), documento=f"600{i:05}", nombre="Ficticio", desde=timezone.localdate()) for i in range(4999)])
        self.assertEqual(len(list(self.libro()["Padrón"].values)), 5001)


    def test_titular_mayor_que_precision_excel_es_texto_y_consulta_numerica(self):
        id_a, id_b = 2 ** 53 + 17, 2 ** 53 + 117
        for numero, pk in enumerate((id_a, id_b), 1):
            m.Afiliado.objects.create(pk=pk, financiador=self.org, numero=f"T{numero}", documento=f"9000000{numero}", nombre=f"Titular {numero}", desde=timezone.localdate())
        m.Afiliado.objects.filter(pk=self.a.pk).update(nombre="Hijo A", titular_id=id_a, parentesco="hijo")
        m.Afiliado.objects.create(financiador=self.org, numero="H2", documento="90000003", nombre="Hijo B", desde=timezone.localdate(), titular_id=id_b, parentesco="hijo")
        parametros = {"search": "Hijo", "ordering": "-titular", "columnas": "nombre,numero,titular"}
        hoja = self.libro(parametros)["Padrón"]
        self.assertEqual([hoja.cell(fila, 3).value for fila in (2, 3)], [str(id_b), str(id_a)])
        self.assertTrue(all(hoja.cell(fila, 3).data_type == "s" and hoja.cell(fila, 3).number_format == "@" for fila in (2, 3)))
        parametros["flt_titular"] = str(id_a)
        listado = self.client.get(f"/api/financiadores/{self.org.pk}/padron/", parametros)
        self.assertEqual(listado.status_code, 200)
        self.assertEqual(listado.data["count"], 1)
        self.assertEqual(listado.data["results"][0]["titular"], id_a)
        filas = list(self.libro(parametros)["Padrón"].values)
        self.assertEqual(filas[1], ("Hijo A", "00025", str(id_a)))
