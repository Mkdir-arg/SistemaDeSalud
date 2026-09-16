"""Recorrido HTTP continuo; fixtures ficticios, sin sustituir servicios de negocio."""
import csv
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import uuid4

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import Membresia, Usuario
from apps.casos.models import Caso
from apps.finanzas.models import (
    AccesoFinanciero, ConcesionFinanciera, MovimientoDinero, ObligacionFinanciera,
)
from apps.flujos.models import Conexion, Nodo
from apps.registros.models import Ciudadano

from . import models as m
from .test_cobertura import CoberturaSetup


class PilotoIntegralTests(CoberturaSetup, APITestCase):
    def setUp(self):
        super().setUp()
        self.caso_b, self.prestacion_b = self.otro_hospital()
        self.hospital_b = self.caso_b.institucion
        self.usuario_b = Usuario.objects.create_user("piloto-b@hospital.local", "x")
        for usuario, caso, prestacion in (
            (self.usuario, self.caso, self.prestacion),
            (self.usuario_b, self.caso_b, self.prestacion_b),
        ):
            miembro = Membresia.objects.create(
                usuario=usuario, institucion=caso.institucion, rol="medico",
            )
            miembro.areas.add(caso.area_actual)
            # El piloto designa usuarios concretos; ser médico no otorga Finanzas.
            for accion in ("registrar_aceptacion", "ver_dinero", "registrar_dinero", "aprobar_dinero"):
                permiso = ConcesionFinanciera.objects.create(membresia=miembro, accion=accion)
                permiso.areas.add(caso.area_actual)
            fin = Nodo.objects.create(version=caso.version, tipo="fin", titulo="Fin")
            Conexion.objects.create(version=caso.version, origen=prestacion.nodo, destino=fin)
        self.membresia.rol = "admin"
        self.membresia.save(update_fields=["rol"])
        self.mutual = m.Financiador.objects.create(nombre="Mutual piloto", tipo="mutual")
        self.operador_mutual = Usuario.objects.create_user("piloto@mutual.local", "x")
        m.MembresiaFinanciador.objects.create(
            financiador=self.mutual, usuario=self.operador_mutual, rol="admin",
        )
        m.Convenio.objects.create(
            financiador=self.mutual, institucion=self.institucion, estado="activo",
            propuesto_por="plataforma", creado_por=self.admin,
        )

    def post(self, url, datos, estado=201):
        respuesta = self.client.post(url, datos, format="json")
        self.assertEqual(respuesta.status_code, estado, respuesta.data)
        return respuesta.data

    def get(self, url, datos=None):
        respuesta = self.client.get(url, datos or {})
        self.assertEqual(respuesta.status_code, 200, respuesta.content)
        return respuesta

    def plantilla(self, financiador, tipo):
        respuesta = self.get(f"/api/financiadores/{financiador.pk}/plantilla/", {"tipo": tipo})
        return load_workbook(BytesIO(respuesta.content))

    def importar(self, financiador, tipo, filas, validas, rechazadas):
        base = f"/api/financiadores/{financiador.pk}/"
        libro = self.plantilla(financiador, tipo)
        # La plantilla formatea filas vacías: escribir desde la 2, sin append.
        for numero, valores in enumerate(filas, 2):
            for columna, valor in enumerate(valores, 1):
                libro["Carga"].cell(numero, columna, valor)
        salida = BytesIO()
        libro.save(salida)
        libro.close()
        antes = (m.Afiliado.objects.count(), m.ConsumoExterno.objects.count())
        respuesta = self.client.post(base + "importaciones/", {
            "tipo": tipo, "clave": str(uuid4()),
            "archivo": SimpleUploadedFile("piloto.xlsx", salida.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        }, format="multipart")
        self.assertEqual(respuesta.status_code, 201, respuesta.data)
        lote = respuesta.data
        self.assertEqual(lote["resumen"]["valida"], validas)
        self.assertEqual(lote["resumen"]["rechazada"], rechazadas)
        self.assertEqual((m.Afiliado.objects.count(), m.ConsumoExterno.objects.count()), antes)
        aplicado = self.post(base + "confirmar-importacion/", {"importacion": lote["id"]}, 200)
        self.assertEqual(aplicado["resumen"]["aplicada"], validas)
        despues = (m.Afiliado.objects.count(), m.ConsumoExterno.objects.count())
        repetido = self.post(base + "confirmar-importacion/", {"importacion": lote["id"]}, 200)
        self.assertEqual(repetido["resumen"], aplicado["resumen"])
        self.assertEqual((m.Afiliado.objects.count(), m.ConsumoExterno.objects.count()), despues)
        if rechazadas:
            errores = self.get(base + "rechazos/", {"importacion": lote["id"]})
            detalle = load_workbook(BytesIO(errores.content), read_only=True)
            valores = [str(valor) for hoja in detalle for fila in hoja.values for valor in fila if valor is not None]
            self.assertIn("DESCONOCIDO", " ".join(valores))
            detalle.close()
        return lote["id"]

    def caso_para(self, origen, prestacion, afiliado, usuario):
        ciudadano, _ = Ciudadano.objects.get_or_create(
            institucion=origen.institucion, documento=afiliado.documento,
            defaults={"nombre": afiliado.nombre},
        )
        caso = Caso.objects.create(
            institucion=origen.institucion, version=origen.version, ciudadano=ciudadano,
            area_actual=origen.area_actual, nodo_actual=prestacion.nodo,
            estado=Caso.Estado.EN_EVALUACION,
        )
        self.client.force_authenticate(usuario)
        contexto = self.get(f"/api/casos/{caso.pk}/cobertura/").data["contexto"]
        self.post(f"/api/casos/{caso.pk}/cobertura-afiliacion/", {
            "contexto": contexto, "afiliado": afiliado.pk,
            "motivo": "Padrón verificado al ingresar en el piloto",
        }, 201)
        return caso

    def cotizar(self, caso, prestacion, usuario):
        self.client.force_authenticate(usuario)
        base = f"/api/casos/{caso.pk}/"
        contexto = self.get(base + "cobertura/").data["contexto"]
        cotizacion = self.post(base + "cobertura-evaluar/", {
            "contexto": contexto, "prestacion": prestacion.pk,
        }, 200)
        return contexto, cotizacion

    def confirmar(self, caso, prestacion, usuario, importe, disponibles, acepta=True):
        contexto, cotizacion = self.cotizar(caso, prestacion, usuario)
        self.assertEqual(cotizacion["importe_paciente"], importe)
        self.assertEqual(cotizacion["disponibles"], disponibles)
        payload = {
            "contexto": contexto, "prestacion": prestacion.pk, "firma": cotizacion["firma"],
            "clave": str(uuid4()), "acepta": acepta,
        }
        antes = ObligacionFinanciera.objects.count()
        reserva = self.post(f"/api/casos/{caso.pk}/cobertura-confirmar/", payload)
        repetida = self.post(f"/api/casos/{caso.pk}/cobertura-confirmar/", payload)
        self.assertEqual(reserva["id"], repetida["id"])
        self.assertEqual(ObligacionFinanciera.objects.count(), antes)
        if acepta:
            self.assertEqual(reserva["aceptacion"]["importe"], importe)
            self.assertEqual(reserva["aceptacion"]["prestacion"], prestacion.pk)
        self.post(f"/api/casos/{caso.pk}/avanzar/", {
            "titulo": "Atención del piloto", "contenido": "Prestación realizada", "firmada": False,
        }, 200)
        caso.refresh_from_db()
        self.assertEqual(caso.estado, Caso.Estado.CERRADO)
        registro = m.ReservaCobertura.objects.get(pk=reserva["id"])
        self.assertEqual(registro.estado, "realizada")
        self.assertIsNotNone(registro.hecho_id)
        cargos = ObligacionFinanciera.objects.count()
        self.post(f"/api/casos/{caso.pk}/avanzar/", {}, 400)
        self.assertEqual(ObligacionFinanciera.objects.count(), cargos)
        return m.DistribucionCobro.objects.get(reserva=registro)

    def test_importar_atender_cobrar_y_conciliar_en_organizaciones_aisladas(self):
        # La parametría del segundo pagador nace desde el portal, con otro catálogo.
        self.client.force_authenticate(self.operador_mutual)
        base_mutual = f"/api/financiadores/{self.mutual.pk}/"
        plan = self.post(base_mutual + "planes/", {"codigo": "M70", "nombre": "Mutual 70 %"})
        exclusiva = m.PrestacionComun.objects.create(codigo="RX", nombre="Radiografía", categoria="imagenes")
        for prestacion in (self.comun, exclusiva):
            self.post(base_mutual + "reglas/", {
                "plan": plan["id"], "prestacion": prestacion.pk, "porcentaje": "70.00",
                "cupo": 6, "periodo": "mes", "vigente_desde": str(self.hoy),
            })
        for financiador, usuario, codigos in (
            (self.financiador, self.operador, {"CONS"}),
            (self.mutual, self.operador_mutual, {"CONS", "RX"}),
        ):
            self.client.force_authenticate(usuario)
            libro = self.plantilla(financiador, "consumos")
            self.assertEqual({fila[0] for fila in list(libro["Prestaciones"].values)[1:]}, codigos)
            libro.close()

        documento = "00444555"
        self.client.force_authenticate(self.operador)
        lote = self.importar(self.financiador, "padron", [
            ["00002", documento, "Persona piloto", "BASE", self.hoy],
            ["00003", "00555666", "Otra persona piloto", "BASE", self.hoy],
            ["00004", "00666777", "Fila rechazada", "DESCONOCIDO", self.hoy],
        ], 2, 1)
        self.assertTrue(m.Afiliado.objects.filter(pk=self.afiliado.pk, finalizado_en=None).exists())
        afiliado = m.Afiliado.objects.get(financiador=self.financiador, documento=documento)
        self.importar(self.financiador, "consumos", [
            ["00002", documento, "CONS", self.hoy, 4, "PILOTO-EXT-1"],
            ["00003", "00555666", "CONS", self.hoy, 1, "PILOTO-EXT-2"],
            ["00002", documento, "DESCONOCIDO", self.hoy, 1, "PILOTO-ERROR"],
        ], 2, 1)
        self.client.force_authenticate(self.operador_mutual)
        self.importar(self.mutual, "padron", [
            ["M0001", documento, "Persona piloto", "M70", self.hoy],
        ], 1, 0)
        afiliado_mutual = m.Afiliado.objects.get(financiador=self.mutual, documento=documento)
        # Un usuario de la mutual no puede recuperar un lote ni actividad de la OS.
        base_os = f"/api/financiadores/{self.financiador.pk}/"
        self.assertEqual(self.client.get(base_os + "rechazos/", {"importacion": lote}).status_code, 404)
        self.assertEqual(self.client.get(base_os + "actividad/").status_code, 404)

        caso_a = self.caso_para(self.caso, self.prestacion, afiliado, self.usuario)
        caso_b = self.caso_para(self.caso_b, self.prestacion_b, afiliado, self.usuario_b)
        cuenta_a = self.confirmar(caso_a, self.prestacion, self.usuario, "20.00", 2)
        cuenta_b = self.confirmar(caso_b, self.prestacion_b, self.usuario_b, "40.00", 1)
        self.assertEqual(cuenta_a.importe_financiador, Decimal("80"))
        self.assertEqual(cuenta_b.importe_financiador, Decimal("160"))
        siguiente = self.caso_para(self.caso, self.prestacion, afiliado, self.usuario)
        _, cotizacion = self.cotizar(siguiente, self.prestacion, self.usuario)
        self.assertEqual(cotizacion["estado"], "no_cubierta")
        particular = self.confirmar(siguiente, self.prestacion, self.usuario, "100.00", 0)
        self.assertIsNone(particular.obligacion_financiador_id)
        self.assertEqual(particular.obligacion_paciente.importe_original, Decimal("100"))

        # La misma persona en otra organización tiene un cupo independiente.
        caso_mutual = self.caso_para(self.caso, self.prestacion, afiliado_mutual, self.usuario)
        sin_aceptacion = self.confirmar(caso_mutual, self.prestacion, self.usuario, "30.00", 6, acepta=False)
        self.assertEqual(sin_aceptacion.estado, "pendiente")
        self.assertIsNone(sin_aceptacion.obligacion_paciente_id)

        pago = {"importe": "30.00", "fecha": str(self.hoy), "clave": str(uuid4()), "aprobado": True}
        url_cobro = f"/api/obligaciones-financieras/{cuenta_a.obligacion_financiador_id}/movimientos/"
        self.post(url_cobro, pago)
        self.post(url_cobro, pago)
        self.assertEqual(MovimientoDinero.objects.count(), 1)
        cuentas = self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk}).data
        self.assertEqual(cuentas["count"], 4)
        self.assertEqual(cuentas["resumen"]["importe_original"], "270.00")
        self.assertEqual(cuentas["resumen"]["registrado_neto"], "30.00")
        self.assertEqual(cuentas["resumen"]["pendiente"], "240.00")
        pendientes = self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk, "vista": "pendientes"}).data
        self.assertEqual(pendientes["count"], 1)
        self.assertEqual(pendientes["results"][0]["caso"], caso_mutual.pk)
        self.assertEqual(pendientes["results"][0]["importe_pendiente"], "30.00")
        self.assertEqual(self.client.get("/api/seguimiento-cobros/", {"institucion": self.hospital_b.pk}).status_code, 403)
        self.assertEqual(self.client.get(f"/api/casos/{caso_b.pk}/cobertura/").status_code, 404)
        self.client.force_authenticate(self.usuario_b)
        cuentas_b = self.get("/api/seguimiento-cobros/", {"institucion": self.hospital_b.pk}).data
        self.assertEqual(cuentas_b["resumen"]["pendiente"], "200.00")
        self.assertEqual(self.client.post(url_cobro, pago, format="json").status_code, 404)

        self.client.force_authenticate(self.operador)
        actividad = self.get(base_os + "actividad/").data
        self.assertEqual(actividad["count"], 3)
        self.assertEqual(actividad["resumen"]["importe_asignado"], "240.00")
        self.assertEqual(self.client.get(f"/api/casos/{caso_a.pk}/cobertura/").status_code, 404)
        self.assertEqual(self.client.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk}).status_code, 403)
        self.client.force_authenticate(self.operador_mutual)
        mutual = self.get(base_mutual + "actividad/").data
        self.assertEqual(mutual["count"], 1)
        self.assertEqual(mutual["resumen"]["importe_asignado"], "70.00")

        # El CSV hospitalario concilia el saldo y mantiene separada la deuda pendiente.
        self.client.force_authenticate(self.usuario)
        archivo = self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk, "formato": "csv"})
        filas = list(csv.DictReader(StringIO(archivo.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(len(filas), 4)
        self.assertEqual(sum(Decimal(fila["Saldo pendiente de cobro (ARS; coma decimal)"].replace(",", ".")) for fila in filas), Decimal("240"))
        self.assertEqual(AccesoFinanciero.objects.get(recurso="seguimiento-cobros", accion="exportar_cuentas").resultados, 4)

        # Un consumo externo tardío informa la discrepancia sin reescribir lo acordado.
        comprometido = m.ReservaCobertura.objects.get(pk=cuenta_b.reserva_id)
        evaluacion, aceptacion = comprometido.evaluacion, comprometido.aceptacion
        self.client.force_authenticate(self.operador)
        self.importar(self.financiador, "consumos", [
            ["00002", documento, "CONS", self.hoy, 1, "PILOTO-EXT-TARDIO"],
        ], 1, 0)
        comprometido.refresh_from_db()
        self.assertTrue(comprometido.discrepancia)
        self.assertEqual(comprometido.evaluacion, evaluacion)
        self.assertEqual(comprometido.aceptacion, aceptacion)
        self.assertEqual(comprometido.cubiertas, 1)
        self.assertEqual(self.get(base_os + "actividad/").data["resumen"]["importe_asignado"], "240.00")
        nueva_atencion = self.caso_para(self.caso_b, self.prestacion_b, afiliado, self.usuario_b)
        _, reevaluacion = self.cotizar(nueva_atencion, self.prestacion_b, self.usuario_b)
        self.assertEqual(reevaluacion["disponibles"], 0)
        self.assertEqual(reevaluacion["importe_paciente"], "200.00")
        self.client.force_authenticate(self.usuario)
        resumen_final = self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk}).data["resumen"]
        self.assertEqual(resumen_final, cuentas["resumen"])
        self.assertEqual(MovimientoDinero.objects.count(), 1)

        # La designación expresa habilita una resolución documentada; cobrar no basta.
        url_resolucion = f"/api/coberturas/{sin_aceptacion.reserva_id}/resolver/"
        rechazo = {"decision": "rechazar", "importe": "30.00", "motivo": "El hospital no asume el copago", "clave": str(uuid4())}
        self.post(url_resolucion, rechazo, 403)
        self.conceder("resolver_cobertura", area=self.area)
        self.post(url_resolucion, rechazo, 200)
        sin_aceptacion.refresh_from_db()
        self.assertEqual(sin_aceptacion.estado, "pendiente")
        acuerdo = self.post(url_resolucion, {
            "decision": "paciente", "importe": "30.00", "motivo": "Aceptación posterior documentada",
            "evidencia": "Aceptación de esta consulta por $30 en acta ficticia PILOTO-1", "clave": str(uuid4()),
        }, 200)
        resolucion = m.ResolucionSaldo.objects.get(pk=acuerdo["id"])
        self.assertEqual(resolucion.registrado_por_id, self.usuario.pk)
        self.assertEqual(resolucion.motivo, "Aceptación posterior documentada")
        self.assertEqual(resolucion.obligacion.importe_original, Decimal("30"))
        sin_aceptacion.refresh_from_db()
        self.assertEqual(sin_aceptacion.estado, "resuelta")
        self.assertEqual(self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk, "vista": "pendientes"}).data["count"], 0)
        conciliado = self.get("/api/seguimiento-cobros/", {"institucion": self.institucion.pk}).data
        self.assertEqual(conciliado["count"], 5)
        self.assertEqual(conciliado["resumen"]["importe_original"], "300.00")
        self.assertEqual(conciliado["resumen"]["pendiente"], "270.00")
        self.assertEqual(MovimientoDinero.objects.count(), 1)
