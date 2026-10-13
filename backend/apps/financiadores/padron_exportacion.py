"""Exportación administrativa acotada; nunca consulta recursos protegidos."""
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import re
from pathlib import Path

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from rest_framework import serializers
from drf_spectacular.utils import OpenApiParameter, OpenApiTypes

from .padron_consulta import COLUMNAS_PADRON, consultar_padron, validar_columnas_padron, valor_columna_padron
from .services import auditar

LIMITE_EXPORTACION_PADRON = 5000
XML_INVALIDO = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def parametros_padron():
    parametros = [OpenApiParameter(name, OpenApiTypes.STR, description=descripcion) for name, descripcion in (
        ("search", "Número, documento, nombre, apellido o CUIL."), ("documento", "Documento exacto."),
        ("plan", "ID de plan de esta organización."), ("estado", "vigentes, futuras, finalizadas o todos."),
        ("derecho", "Mismos valores que estado; prevalece sobre estado."),
        ("estado_administrativo", "activo, en_tramite, baja_programada o baja_administrativa."),
        ("regimen", "Régimen administrativo."), ("columnas", "Claves separadas por coma; nombre y numero obligatorias. Solo administración."),
    )]
    parametros.append(OpenApiParameter("ordering", OpenApiTypes.STR, enum=[key for key in COLUMNAS_PADRON] + [f"-{key}" for key in COLUMNAS_PADRON], description="Una columna; desempate pk ascendente y nulos últimos. Ausente: pk."))
    for prefijo in ("alta", "baja"):
        for extremo in ("desde", "hasta"):
            parametros.append(OpenApiParameter(f"{prefijo}_{extremo}", OpenApiTypes.DATE))
    for key, columna in COLUMNAS_PADRON.items():
        if columna["type"] in ("date", "datetime"):
            for extremo in ("desde", "hasta"):
                parametros.append(OpenApiParameter(f"flt_{key}_{extremo}", OpenApiTypes.DATE))
        else:
            parametros.append(OpenApiParameter(f"flt_{key}", OpenApiTypes.STR, enum=list(columna["options"]) if columna.get("options") else None))
        if columna.get("nullable"):
            parametros.append(OpenApiParameter(f"flt_{key}_vacio", OpenApiTypes.BOOL))
    return parametros


def celda_padron(hoja, valor):
    celda = WriteOnlyCell(hoja)
    if isinstance(valor, datetime):
        celda.value = timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
        celda.number_format = "dd/mm/yyyy hh:mm"
    elif isinstance(valor, date):
        celda.value = valor
        celda.number_format = "dd/mm/yyyy"
    elif valor is not None:
        # Excel limita cada celda. No permitir el truncamiento silencioso de openpyxl.
        texto = XML_INVALIDO.sub("", str(valor))
        if len(texto) > 32767:
            raise serializers.ValidationError("Un texto supera el límite de XLSX. Acotá los filtros para exportar otros registros.")
        celda.value = texto
        # Asignar el tipo DESPUÉS del valor: openpyxl infiere fórmula al ver '='.
        celda.data_type = "s"
        celda.number_format = "@"
    return celda


def exportar_padron(request, org):
    if request.query_params.get("formato", "xlsx") != "xlsx":
        raise serializers.ValidationError("El formato admitido es XLSX.")
    columnas = validar_columnas_padron(request.query_params.get("columnas"))
    qs = consultar_padron(financiador=org, parametros=request.query_params)
    # Una sola lectura acotada: página y tamaño no limitan el archivo.
    filas = list(qs[:LIMITE_EXPORTACION_PADRON + 1].iterator(chunk_size=250))
    if len(filas) > LIMITE_EXPORTACION_PADRON:
        raise serializers.ValidationError(f"La exportación admite hasta {LIMITE_EXPORTACION_PADRON} registros. Acotá los filtros para descargar el resultado completo.")
    libro = Workbook(write_only=True)
    hoja = libro.create_sheet("Padrón")
    hoja.freeze_panes = "A2"
    hoja.append([celda_padron(hoja, COLUMNAS_PADRON[key]["label"]) for key in columnas])
    try:
        for afiliado in filas:
            hoja.append([celda_padron(hoja, valor_columna_padron(afiliado, key)) for key in columnas])
        from openpyxl.utils import get_column_letter
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
        contenido = BytesIO()
        libro.save(contenido)
    finally:
        # Cierra también el generador/archivo temporal si una celda fue rechazada.
        if not hoja.closed:
            hoja.close()
        if hoja._writer and Path(hoja._writer.out).exists():
            hoja._writer.cleanup()
        libro.close()
    claves = ",".join(columnas)
    motivo = f"XLSX; registros={len(filas)}; columnas={claves}"
    if len(motivo) > 255:
        motivo = f"XLSX; registros={len(filas)}; columnas={len(columnas)}; columnas_sha256={sha256(claves.encode()).hexdigest()}"
    # Ningún byte sale si falla la auditoría. No registrar búsquedas ni personas.
    with transaction.atomic():
        auditar(request.user, "exportar_padron", org.pk, financiador=org, motivo=motivo)
    response = HttpResponse(contenido.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="padron-{org.pk}-{timezone.localdate().isoformat()}.xlsx"'
    response["Cache-Control"] = "no-store"
    response["X-Content-Type-Options"] = "nosniff"
    return response
