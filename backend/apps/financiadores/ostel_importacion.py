"""Carga OSTEL con habilitación elegida por lote para personas nuevas.

Las filas con evidencia de baja no son órdenes de alta ni de finalización.
Se conservan códigos sin homologar. No se deducen familias, bajas o reingresos.
CUILTIT (entrega del 26/10/2026) se conserva como evidencia; vincular grupos
familiares queda para #181.
"""
from collections import Counter
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from math import isfinite
from uuid import UUID

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction, IntegrityError
from django.utils import timezone
from openpyxl import load_workbook

from .models import Afiliado, Financiador, Importacion
from .ostel_prevalidacion import prevalidar_ostel, ArchivoOstelInvalido, COLUMNAS, MAX_BYTES
from .permisos import requerir_financiador
from . import padron
from .services import auditar, registrar_afiliado
from .importacion_cobertura import revisar_cobertura, validar_opcion

VERSION = 3
CAMPOS = ("numero_interno", "numero", "documento", "cuil", "nombre")
# Columna OSTEL de cada campo de identidad: el archivo con comentarios marca ahí el error.
COLUMNA = dict(zip(CAMPOS, ("NUMERO", "NROAFI", "NRODOC", "CUIL", "APENOM")))
# `original` conserva las 40 columnas del archivo, sin homologar: es la
# evidencia completa de cada fila. Las familias NO se arman aquí: #181 las
# construirá a partir de CUILTIT y CODPAR. NROREF y NROAFI siguen sin
# significado confirmado por OSTEL; CODLOCA no se homologa.
ORIGINAL = COLUMNAS

# Ficha que OSTEL actualiza al aplicar (nuevos y existentes). Un valor vacío
# del archivo no borra el dato existente; un valor inválido no se escribe y
# queda como advertencia de la fila, sin bloquearla.
FICHA = (("calle", "A_CALLE"), ("altura", "A_NRO"), ("piso", "A_PISO"),
         ("departamento", "A_DPTO"), ("codigo_postal", "CODPOS"), ("localidad", "DESLOCA"),
         ("telefono", "A_TEL1"), ("celular", "CELU1"), ("email", "EMAIL1"))
# SUPUESTO pendiente de confirmación de OSTEL: CODPROV usa la letra de provincia
# de ISO 3166-2:AR (la del DNI). Nombres oficiales de Georef, como texto libre
# en `provincia`. Una letra desconocida no se traduce. CODLOCA no se homologa.
PROVINCIAS = {
    "A": "Salta", "B": "Buenos Aires", "C": "Ciudad Autónoma de Buenos Aires",
    "D": "San Luis", "E": "Entre Ríos", "F": "La Rioja", "G": "Santiago del Estero",
    "H": "Chaco", "J": "San Juan", "K": "Catamarca", "L": "La Pampa", "M": "Mendoza",
    "N": "Misiones", "P": "Formosa", "Q": "Neuquén", "R": "Río Negro", "S": "Santa Fe",
    "T": "Tucumán", "U": "Chubut", "V": "Tierra del Fuego, Antártida e Islas del Atlántico Sur",
    "W": "Corrientes", "X": "Córdoba", "Y": "Jujuy", "Z": "Santa Cruz",
}
# Campos de ficha que una persona puede corregir al revisar la fila.
EDITABLES = tuple(campo for campo, _ in FICHA) + ("sexo", "fecha_nacimiento", "provincia")


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, bool) or not isinstance(valor, (str, int, float)):
        raise ValidationError("Identificador o texto con tipo no admitido.")
    if isinstance(valor, float):
        if not valor.is_integer():
            raise ValidationError("Los identificadores no admiten decimales.")
        valor = int(valor)
    return str(valor).strip()


def _admitido(valor):
    try:
        _texto(valor)
        return True
    except ValidationError:
        return False


def _json(valor):
    if isinstance(valor, float) and not isfinite(valor):
        raise ValidationError("El archivo contiene un número fuera de rango. Corregí la fuente antes de continuar.")
    return valor.isoformat() if isinstance(valor, (date, datetime)) else valor


def _ficha(raw):
    """Devuelve (ficha, advertencias, campos advertidos); nunca rechaza la fila por la ficha.

    `campos` va alineado con las advertencias: al corregir un campo en la
    revisión, su advertencia deja de aplicar.
    """
    ficha, avisos, campos = {}, [], []

    def avisar(campo, texto):
        avisos.append(texto)
        campos.append(campo)

    for campo, columna in FICHA:
        try:
            valor = _texto(raw.get(columna))
        except ValidationError:
            avisar(campo, f"{columna} tiene un tipo no admitido; no se actualiza {campo}.")
            continue
        if not valor:
            continue
        if len(valor) > Afiliado._meta.get_field(campo).max_length:
            avisar(campo, f"{columna} supera el largo admitido; no se actualiza {campo}.")
            continue
        if campo == "email":
            try:
                validate_email(valor)
            except ValidationError:
                avisar(campo, "EMAIL1 no es un correo válido; no se actualiza el email.")
                continue
        ficha[campo] = valor
    sexo = raw.get("SEXO")
    if sexo in ("F", "M"):
        ficha["sexo"] = sexo
    elif sexo not in (None, ""):
        avisar("sexo", "SEXO sin equivalencia; no se actualiza el sexo.")
    nacimiento = raw.get("FEC_NAC")
    if isinstance(nacimiento, datetime):
        nacimiento = nacimiento.date()
    if isinstance(nacimiento, date) and nacimiento <= timezone.localdate():
        ficha["fecha_nacimiento"] = nacimiento.isoformat()
    elif nacimiento not in (None, ""):
        avisar("fecha_nacimiento", "FEC_NAC no es una fecha válida; no se actualiza la fecha de nacimiento.")
    provincia = raw.get("CODPROV")
    if isinstance(provincia, str) and provincia.strip() in PROVINCIAS:
        ficha["provincia"] = PROVINCIAS[provincia.strip()]
    elif provincia not in (None, ""):
        avisar("provincia", "CODPROV sin equivalencia de provincia; no se actualiza la provincia.")
    return ficha, avisos, campos


def _ficha_corregida(valores):
    """Valida la ficha corregida por una persona. A diferencia del archivo, un
    valor inválido rechaza la corrección: quien revisa lo puede arreglar. Vacío
    significa «no actualizar», igual que en el archivo."""
    if not isinstance(valores, dict) or set(valores) - set(EDITABLES):
        raise ValidationError("La ficha corregida tiene campos no admitidos.")
    ficha = {}
    for campo, valor in valores.items():
        if not isinstance(valor, str):
            raise ValidationError("La ficha corregida sólo admite texto.")
        valor = valor.strip()
        if not valor:
            continue
        if campo == "sexo":
            if valor not in ("F", "M"):
                raise ValidationError("Elegí el sexo F o M, o dejalo vacío.")
        elif campo == "fecha_nacimiento":
            try:
                fecha = date.fromisoformat(valor)
            except ValueError:
                raise ValidationError("La fecha de nacimiento debe tener el formato AAAA-MM-DD.") from None
            if fecha > timezone.localdate():
                raise ValidationError("La fecha de nacimiento no puede ser futura.")
        elif campo == "provincia":
            if valor not in PROVINCIAS.values():
                raise ValidationError("Elegí una provincia de la lista.")
        elif len(valor) > Afiliado._meta.get_field(campo).max_length:
            raise ValidationError(f"El campo {campo} supera el largo admitido.")
        elif campo == "email":
            try:
                validate_email(valor)
            except ValidationError:
                raise ValidationError("El email corregido no es un correo válido.") from None
        ficha[campo] = valor
    return ficha


def _guardar(lote):
    opcion = lote.resumen.get("cobertura_habilitada")
    lote.resumen = {"total": len(lote.filas), **{
        estado: sum(f["estado"] == estado for f in lote.filas)
        for estado in ("valida", "revision", "rechazada", "aplicada", "error_tecnico", "omitida")}}
    # Filas cuya ficha no se actualiza completa: la UI lo avisa antes de aplicar.
    lote.resumen["advertencias"] = sum(bool(f.get("advertencias")) for f in lote.filas)
    lote.resumen["fuente"] = lote.filas[0]["fuente"] if lote.filas else "base"
    if opcion is not None:
        lote.resumen["cobertura_habilitada"] = opcion
    lote.estado = "preview" if any(f["estado"] not in ("aplicada", "omitida") for f in lote.filas) else "aplicada"
    lote.save(update_fields=["filas", "resumen", "estado"])


def _error(texto, *campos):
    """Error de fila con las columnas OSTEL donde se origina (puede no haber)."""
    return ValidationError(texto, params={"columnas": tuple(COLUMNA[c] for c in campos)})


def _columnas(exc):
    columnas = {c for e in getattr(exc, "error_list", ()) for c in (e.params or {}).get("columnas", ())}
    return sorted(columnas, key=COLUMNAS.index)


def _marcar(fila, exc):
    # Un error sin columna propia (cobertura, padrón cambiado) es de la persona:
    # se marca en su documento, nunca en toda una columna ajena.
    fila.update(estado="revision", errores=list(exc.messages), columnas_error=_columnas(exc) or [COLUMNA["documento"]])


def _validar(fila, lote, repetidos, referencias=None):
    d = fila["datos"]
    if fila.get("bloqueo_fuente"):
        raise ValidationError(fila["bloqueo_fuente"], params={"columnas": tuple(fila.get("columnas_bloqueo", ()))})
    if set(d) != set(CAMPOS):
        raise ValidationError("Campos administrativos incompletos.")
    for campo in CAMPOS:
        valor = d[campo]
        if not isinstance(valor, str) or not valor.strip() or len(valor) > (160 if campo == "nombre" else 80):
            raise _error(f"Completá {campo} con un valor válido.", campo)
    if not d["documento"].isascii() or not d["documento"].isdigit():
        raise _error("Revisá el documento; se requiere su representación numérica sin separadores.", "documento")
    try:
        padron.validar_cuit(d["cuil"])
    except ValidationError as exc:
        raise _error(" ".join(exc.messages), "cuil") from None
    if len(d["documento"]) > 8 or d["cuil"][2:10] != d["documento"].zfill(8):
        raise _error("Documento y CUIL no concuerdan; requiere revisión humana.", "documento", "cuil")
    for campo in ("numero_interno", "documento", "cuil"):
        if repetidos[campo][d[campo]] > 1:
            raise _error(f"{campo} repetido en el lote. Revisá todas las filas relacionadas.", campo)
    qs = Afiliado.objects.filter(financiador=lote.financiador)
    obj = referencias["documento"].get(d["documento"]) if referencias is not None else qs.filter(documento=d["documento"]).first()
    for campo in ("numero_interno", "cuil"):
        otro = referencias[campo].get(d[campo]) if referencias is not None else qs.filter(**{campo: d[campo]}).first()
        if otro and (obj is None or otro.pk != obj.pk):
            raise _error(f"{campo} pertenece a otra identidad del padrón; no se fusionan personas.", campo)
        if obj and getattr(obj, campo) and getattr(obj, campo) != d[campo]:
            raise _error(f"{campo} contradice el padrón; revisá la identidad antes de aplicar.", campo)
    if obj:
        if obj.finalizado_en:
            raise _error("Afiliación finalizada: un posible reingreso requiere gestión expresa.", "documento")
        if fila.get("revision_padron", 0) != obj.revision:
            raise _error("El padrón cambió desde la revisión. Revisá y guardá la fila nuevamente.", "documento")
        distintos = [k for k in CAMPOS if getattr(obj, k) != d[k]]
        if distintos and not fila.get("cambios_revisados"):
            raise _error("Hay diferencias con el padrón actual. Compará los valores y confirmá la corrección con un motivo.", *distintos)
    elif fila.get("revision_padron", 0):
        raise ValidationError("La identidad cambió desde la revisión; volvé a revisar la fila.")
    revisar_cobertura(fila, obj, lote.resumen.get("cobertura_habilitada"))
    return obj


def _repetidos(lote):
    return {k: Counter(f["datos"].get(k) for f in lote.filas if f["estado"] != "omitida")
            for k in ("numero_interno", "documento", "cuil")}


def _revisar(lote):
    repetidos = _repetidos(lote)
    referencias = {k: {} for k in ("documento", "numero_interno", "cuil")}
    for obj in Afiliado.objects.filter(financiador=lote.financiador):
        for campo in referencias:
            if getattr(obj, campo):
                referencias[campo][getattr(obj, campo)] = obj
    for fila in lote.filas:
        if fila["estado"] in ("aplicada", "omitida"):
            continue
        try:
            _validar(fila, lote, repetidos, referencias)
            fila.update(estado="valida", errores=[], columnas_error=[])
        except ValidationError as exc:
            _marcar(fila, exc)
    _guardar(lote)


@transaction.atomic
def previsualizar(*, financiador, usuario, archivo, clave, fuente="base", cobertura_habilitada=None):
    requerir_financiador(usuario, financiador.pk, escritura=True)
    validar_opcion(cobertura_habilitada)
    if fuente not in ("base", "bajas"):
        raise ValidationError("Elegí padrón base o evidencia de bajas.")
    try:
        clave = UUID(str(clave))
    except (ValueError, TypeError):
        raise ValidationError("La clave de importación no es válida.") from None
    contenido = archivo.read(MAX_BYTES + 1)
    opcion = b"" if cobertura_habilitada is None else str(cobertura_habilitada).encode()
    huella = sha256(fuente.encode() + opcion + contenido).hexdigest()
    previo = Importacion.objects.filter(financiador=financiador, clave=clave).first()
    if previo:
        if previo.version != VERSION or previo.huella != huella:
            raise ValidationError("La clave pertenece a otro archivo o fuente.")
        return previo
    try:
        lectura = prevalidar_ostel(contenido)
    except ArchivoOstelInvalido as exc:
        raise ValidationError(str(exc)) from None
    # No asignar nombres de campos a posiciones cuyo encabezado no se verificó.
    if any(a.codigo != "fila_duplicada" for a in lectura.anomalias):
        raise ValidationError("Corregí primero las anomalías técnicas del archivo; no se creó el lote.")
    libro = load_workbook(BytesIO(contenido), read_only=True, data_only=False, keep_links=False)
    filas = []
    existentes = {o.documento: o for o in Afiliado.objects.filter(financiador=financiador)}
    try:
        for hoja in libro:
            posiciones = {f.numero for f in lectura.filas[hoja.title]}
            # Las dimensiones declaradas del XLSX no son confiables. El lector
            # ya verificó las coordenadas reales y el ancho exacto de 40 columnas.
            it = hoja.iter_rows(min_row=1, max_row=max(posiciones, default=1), max_col=len(COLUMNAS), values_only=True)
            nombres = next(it)
            for numero, valores in enumerate(it, 2):
                if numero not in posiciones:
                    continue
                raw = dict(zip(nombres, valores))
                # Evidencia cruda de las 40 columnas, sin asignar titular,
                # parentesco ni derechos. Ver ORIGINAL.
                original = {k: _json(raw.get(k)) for k in ORIGINAL}
                try:
                    d = {k: _texto(raw[COLUMNA[k]]) for k in CAMPOS}
                    error, bloqueo = "", []
                except ValidationError:
                    d = dict.fromkeys(CAMPOS, "")
                    error = "Tipos de identificadores no admitidos. Corregí el archivo y volvé a cargarlo."
                    bloqueo = [COLUMNA[k] for k in CAMPOS if not _admitido(raw[COLUMNA[k]])]
                ficha, advertencias, advertidos = _ficha(raw)
                if fuente == "bajas" or raw.get("ESTADO") == "B" or raw.get("FECBAJ") is not None:
                    error = "Evidencia de baja: no se aplica automáticamente. Requiere resolver sus movimientos."
                    bloqueo = [c for c in ("FECBAJ", "ESTADO") if (raw.get("FECBAJ") is not None if c == "FECBAJ"
                                                                     else raw.get("ESTADO") == "B")] or ["FECBAJ"]
                if error:
                    # Una fila bloqueada nunca se aplica: su ficha no se usa y
                    # advertir sobre ella sólo agregaría ruido.
                    advertencias, advertidos = [], []
                obj = existentes.get(d["documento"])
                filas.append({"fila": len(filas) + 1, "hoja": hoja.title, "fila_excel": numero,
                              "datos": d, "original": original, "estado": "revision", "errores": [],
                              "ficha": ficha, "advertencias": advertencias, "advertidos": advertidos,
                              "revision_padron": obj.revision if obj else 0, "revision": 1,
                              "padron_actual": {k: getattr(obj, k) for k in CAMPOS} if obj else None,
                              "bloqueo_fuente": error, "columnas_bloqueo": bloqueo, "correcciones": [], "fuente": fuente})
    finally:
        libro.close()
    Financiador.objects.select_for_update(no_key=True).get(pk=financiador.pk)
    previo = Importacion.objects.filter(financiador=financiador, clave=clave).first()
    if previo:
        if previo.version != VERSION or previo.huella != huella:
            raise ValidationError("La clave pertenece a otro archivo o fuente.")
        return previo
    lote = Importacion.objects.create(financiador=financiador, tipo="padron", version=VERSION,
        clave=clave, huella=huella, filas=filas, creado_por=usuario,
        resumen={} if cobertura_habilitada is None else {"cobertura_habilitada": cobertura_habilitada})
    _revisar(lote)
    auditar(usuario, "previsualizar_ostel", lote.pk, financiador=financiador)
    return lote


@transaction.atomic
def subsanar(*, importacion, usuario, fila, revision, datos, motivo, accion="corregir", ficha=None):
    requerir_financiador(usuario, importacion.financiador_id, escritura=True)
    lote = Importacion.objects.select_for_update().get(pk=importacion.pk)
    if lote.version != VERSION or lote.tipo != "padron":
        raise ValidationError("El lote no admite esta corrección.")
    actual = next((f for f in lote.filas if f["fila"] == fila), None)
    if not actual or actual["estado"] in ("aplicada", "omitida") or actual["revision"] != revision:
        raise ValidationError("La fila cambió o ya fue aplicada. Recargá el lote.")
    if accion not in ("corregir", "omitir"):
        raise ValidationError("Acción no admitida.")
    if actual["bloqueo_fuente"] and accion != "omitir":
        raise ValidationError(actual["bloqueo_fuente"])
    if not motivo.strip() or len(motivo) > 255 or set(datos) != set(CAMPOS):
        raise ValidationError("Completá los campos de identidad y el motivo de corrección.")
    if accion == "corregir" and len(actual["correcciones"]) >= 50:
        raise ValidationError("Se alcanzó el límite de revisiones de esta fila; excluila y prepará un nuevo lote corregido.")
    corregidos = {k: _texto(v) for k, v in datos.items()}
    ficha_nueva = _ficha_corregida(ficha) if accion == "corregir" and ficha is not None else None
    obj = Afiliado.objects.filter(financiador=lote.financiador, documento=corregidos["documento"]).first()
    destino_cambio = obj and (actual.get("padron_actual") or {}).get("documento") != obj.documento
    if accion == "corregir" and (destino_cambio or (obj.revision if obj else 0) != actual["revision_padron"]):
        actual.update(padron_actual={k: getattr(obj, k) for k in CAMPOS} if obj else None,
                      revision_padron=obj.revision if obj else 0, cambios_revisados=False,
                      revision=actual["revision"] + 1, estado="revision",
                      errores=["El padrón cambió. Compará los valores actuales y revisá nuevamente; no se guardó la corrección."])
        _guardar(lote)
        return lote
    actual["correcciones"].append({"anterior": actual["datos"], "nuevo": corregidos, "accion": accion,
        "motivo": motivo, "usuario": usuario.pk, "fecha": timezone.now().isoformat(),
        **({"ficha_anterior": actual.get("ficha") or {}, "ficha_nueva": ficha_nueva} if ficha_nueva is not None else {})})
    actual["datos"] = corregidos
    if ficha_nueva is not None:
        actual["ficha"] = ficha_nueva
    if ficha_nueva is not None and "advertidos" in actual:
        # Una persona revisó la ficha: lo que corrigió deja de estar advertido y
        # lo que dejó vacío sigue sin actualizarse. Los lotes anteriores a
        # `advertidos` conservan sus advertencias tal cual.
        pares = [(c, a) for c, a in zip(actual.get("advertidos") or [], actual.get("advertencias") or [])
                 if c not in ficha_nueva]
        actual.update(advertidos=[c for c, _ in pares], advertencias=[a for _, a in pares])
    actual["revision"] += 1
    if accion == "omitir":
        actual["estado"] = "omitida"
    actual["revision_padron"] = obj.revision if obj else 0
    actual["padron_actual"] = {k: getattr(obj, k) for k in CAMPOS} if obj else None
    actual["cambios_revisados"] = accion == "corregir"
    if accion == "corregir":
        try:
            revisar_cobertura(actual, obj, lote.resumen.get("cobertura_habilitada"), motivo)
        except ValidationError:
            pass  # _revisar mantiene la diferencia nueva pendiente de otra revisión.
    _revisar(lote)
    auditar(usuario, "subsanar_ostel", lote.pk, financiador=lote.financiador)
    return lote


@transaction.atomic
def aplicar(*, importacion, usuario):
    requerir_financiador(usuario, importacion.financiador_id, escritura=True)
    lote = Importacion.objects.select_for_update().get(pk=importacion.pk)
    if lote.version != VERSION or lote.tipo != "padron":
        raise ValidationError("El lote no corresponde al formato administrativo.")
    Financiador.objects.select_for_update(no_key=True).get(pk=lote.financiador_id)
    repetidos = _repetidos(lote)
    for fila in [f for f in lote.filas if f["estado"] == "valida"][:100]:
        try:
            with transaction.atomic():
                d = fila["datos"]
                # Mismo orden de locks que escritores de padrón: financiador, afiliado.
                list(Afiliado.objects.select_for_update().filter(financiador=lote.financiador,
                                                               documento=d["documento"]))
                obj = _validar(fila, lote, repetidos)
                if obj is None:
                    obj = registrar_afiliado(financiador=lote.financiador, usuario=usuario,
                        numero=d["numero"], documento=d["documento"], nombre=d["nombre"],
                        plan=None, desde=timezone.localdate(), cobertura_habilitada=lote.resumen.get("cobertura_habilitada", False),
                        solo_crear=True, origen="importacion", lote=lote,
                        motivo="Alta OSTEL con habilitación seleccionada al revisar el lote")
                entrada = {k: d[k] for k in ("numero", "nombre", "cuil")}
                # OSTEL pisa la ficha administrativa; editar_ficha registra el
                # delta anterior/nuevo en el historial y sube la revisión.
                ficha = dict(fila.get("ficha") or {})
                if "fecha_nacimiento" in ficha:
                    ficha["fecha_nacimiento"] = date.fromisoformat(ficha["fecha_nacimiento"])
                entrada.update(ficha)
                # NUMERO es la clave externa individual; NROAFI se guarda en numero
                # para mostrarlo, sin sentido de familia (pendiente de OSTEL).
                if obj.numero_interno != d["numero_interno"]:
                    anterior = obj.numero_interno
                    obj.numero_interno = d["numero_interno"]
                    obj.revision += 1
                    obj.full_clean()
                    obj.save(update_fields=["numero_interno", "revision"])
                    padron.registrar_historial(obj, usuario,
                        {"numero_interno": {"anterior": anterior, "nuevo": obj.numero_interno}},
                        "Identificación interna OSTEL", "importacion", lote)
                padron.editar_ficha(afiliado=obj, usuario=usuario,
                    entrada={**entrada, "revision": obj.revision, "motivo": "Carga administrativa OSTEL"},
                    origen="importacion", lote=lote)
                fila.update(estado="aplicada", errores=[], columnas_error=[], afiliado=obj.pk)
        except ValidationError as exc:
            _marcar(fila, exc)
            actual = Afiliado.objects.filter(financiador=lote.financiador, documento=fila["datos"]["documento"]).first()
            fila.update(padron_actual={k: getattr(actual, k) for k in CAMPOS} if actual else None,
                        revision_padron=actual.revision if actual else 0, cambios_revisados=False,
                        revision=fila["revision"] + 1)
        except IntegrityError:
            fila.update(estado="revision", errores=["Conflicto concurrente de identidad. Revisá la fila nuevamente."])
    _guardar(lote)
    auditar(usuario, "aplicar_ostel", lote.pk, financiador=lote.financiador)
    return lote
