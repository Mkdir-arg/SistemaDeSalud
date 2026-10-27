"""Lectura y carga idempotente de diccionarios de códigos de un financiador.

El diccionario OSTEL no contiene datos personales. La carga sólo crea o
actualiza códigos: no borra los ausentes ni asigna parentescos, familias,
planes o derechos (eso queda para #181 y siguientes).
"""
from datetime import date, datetime
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import CodigoFinanciador

MAX_BYTES = 5 * 1024 * 1024
# Tablas de Catalogo_completo confirmadas por las hojas con nombre del
# diccionario OSTEL. Las demás conservan su número: "ostel_05", "ostel_10"...
TABLAS_OSTEL = {"01": "planes", "02": "motivos_baja", "03": "tipos_documento", "04": "tipos_afiliado"}
PARENTESCOS = ("CODPAR", "DESCRI", "ACTIVO")
CATALOGO = ("TABLA", "CODIGO", "DESCRIP", "ACTIVO")


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _atributo(valor):
    return valor.isoformat() if isinstance(valor, (date, datetime)) else valor


def _hoja(libro, nombre, requeridas):
    if nombre not in libro.sheetnames:
        raise ValidationError(f"Falta la hoja {nombre} del diccionario.")
    filas = libro[nombre].iter_rows(values_only=True)
    encabezado = [_texto(v) for v in next(filas, ())]
    faltantes = [c for c in requeridas if c not in encabezado]
    if faltantes:
        raise ValidationError(f"La hoja {nombre} no tiene las columnas {', '.join(faltantes)}.")
    for valores in filas:
        if any(v not in (None, "") for v in valores):
            yield dict(zip(encabezado, valores))


def leer_diccionario_ostel(contenido):
    """Devuelve ({(tabla, codigo): {"descripcion", "activo", "atributos"}}, omitidas).

    Las filas sin código (OSTEL informa un «DESCONOCIDO» sin código en la tabla
    12) no se cargan: no hay clave con la cual buscarlas. Se cuentan como omitidas.
    """
    from openpyxl import load_workbook

    if not isinstance(contenido, bytes) or len(contenido) > MAX_BYTES:
        raise ValidationError("El diccionario debe ser un XLSX de hasta 5 MiB.")
    try:
        libro = load_workbook(BytesIO(contenido), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ValidationError("El diccionario no es una planilla XLSX legible.") from None
    codigos, omitidas = {}, 0
    try:
        fuentes = [("parentescos", fila, "CODPAR", "DESCRI") for fila in _hoja(libro, "Parentescos", PARENTESCOS)]
        for fila in _hoja(libro, "Catalogo_completo", CATALOGO):
            tabla = _texto(fila["TABLA"])
            fuentes.append((TABLAS_OSTEL.get(tabla, f"ostel_{tabla}"), fila, "CODIGO", "DESCRIP"))
        for tabla, fila, columna_codigo, columna_descripcion in fuentes:
            codigo = _texto(fila[columna_codigo])
            if not codigo:
                omitidas += 1
                continue
            if len(codigo) > 40 or len(tabla) > 40:
                raise ValidationError(f"Código demasiado largo en la tabla {tabla}.")
            if (tabla, codigo) in codigos:
                raise ValidationError(f"Código repetido en la tabla {tabla}: {codigo}.")
            excluidas = {columna_codigo, columna_descripcion, "ACTIVO", "TABLA"}
            codigos[(tabla, codigo)] = {
                "descripcion": _texto(fila[columna_descripcion])[:160],
                "activo": fila["ACTIVO"] is not False,
                "atributos": {k: _atributo(v) for k, v in fila.items() if k and k not in excluidas},
            }
    finally:
        libro.close()
    return codigos, omitidas


@transaction.atomic
def cargar_codigos(financiador, codigos):
    """Crea o actualiza; repetir la misma carga no cambia nada."""
    resultado = {"creados": 0, "actualizados": 0, "sin_cambios": 0}
    existentes = {(c.tabla, c.codigo): c for c in
                  CodigoFinanciador.objects.select_for_update().filter(financiador=financiador)}
    for (tabla, codigo), datos in sorted(codigos.items()):
        actual = existentes.get((tabla, codigo))
        if actual is None:
            CodigoFinanciador.objects.create(financiador=financiador, tabla=tabla, codigo=codigo, **datos)
            resultado["creados"] += 1
        elif any(getattr(actual, k) != v for k, v in datos.items()):
            for k, v in datos.items():
                setattr(actual, k, v)
            actual.save()
            resultado["actualizados"] += 1
        else:
            resultado["sin_cambios"] += 1
    resultado["no_informados"] = len(existentes.keys() - codigos.keys())
    return resultado
