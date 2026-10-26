"""Pruebas del generador y su contrato con el lector #180, sin Django/DB."""

from dataclasses import FrozenInstanceError
from datetime import date, datetime
from io import BytesIO
from hashlib import sha256
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import load_workbook

from .ostel_synthetic import (
    COLUMNAS, HOJAS, Lote, exportar_temporales, generar_escenarios, generar_xlsx,
)


def campos(fila):
    return dict(zip(COLUMNAS, fila))


def checksum_valido(cuil):
    # Comprobación independiente: suma ponderada incluyendo el verificador.
    return sum(int(n) * p for n, p in zip(cuil, (5, 4, 3, 2, 7, 6, 5, 4, 3, 2, 1))) % 11 == 0


class EscenariosOSTELTests(unittest.TestCase):
    def setUp(self):
        self.catalogo = {e.codigo: e for e in generar_escenarios()}

    def filas(self, codigo, lote=0):
        return [campos(f) for f in self.catalogo[codigo].lotes[lote].filas]

    def test_esquema_exacto_y_cobertura(self):
        esperado = (
            "TIPO,NROAFI,CODEMP,CODPAR,APENOM,TIPODOC,NRODOC,CUIL,SEXO,COD_PLAN,"
            "FEC_NAC,FECING,FECBAJ,FECVTO,MOTIVO,NROBENE,NROPENS,CARNET,NIV,INCAPA,"
            "FACTURA,TABVAL,NUMERO,NROREF,A_CALLE,A_NRO,A_PISO,A_DPTO,A_TEL1,CELU1,"
            "EMAIL1,CODPOS,DESLOCA,CODPROV,CODZON,CODLOCA,CODEST,CODNAC,ESTADO,CUILTIT"
        )
        self.assertEqual(COLUMNAS, tuple(esperado.split(",")))
        self.assertEqual(len(COLUMNAS), 40)
        self.assertEqual(set(self.catalogo), {
            "base", "numero_compartido", "referencia_presente", "referencia_ausente",
            "referencia_no_aaa", "cuil_invalido", "documento_discordante",
            "filas_vacias", "tipos_mixtos", "baja_repetida", "posible_reingreso",
            "vencimiento_lejano", "baja_anterior_ingreso", "lote_parcial",
            "referencia_ambigua", "fila_repetida", "codigos_sin_homologar",
            "contactos_ausentes", "texto_literal", "fuentes_discordantes",
            "identidad_individual", "cuil_distinto_mismo_numero",
            "numero_distinto_mismo_cuil", "categorias_sin_cobertura",
            "parentescos_confirmados", "cambio_empresa", "cambio_categoria",
        })
        self.assertEqual(len(generar_escenarios()), len(self.catalogo))
        for escenario in self.catalogo.values():
            for lote in escenario.lotes:
                for fila in lote.filas:
                    self.assertEqual(len(fila), 40)

    def test_determinismo_e_inmutabilidad(self):
        self.assertEqual(generar_escenarios(), generar_escenarios())
        with self.assertRaises(FrozenInstanceError):
            self.catalogo["base"].codigo = "cambio"
        with self.assertRaises(TypeError):
            self.catalogo["base"].lotes[0].filas[0][0] = "cambio"

    def test_identidades_marcadas_y_sin_contactos_operativos(self):
        for escenario in self.catalogo.values():
            for lote in escenario.lotes:
                for fila in lote.filas:
                    if all(v is None for v in fila):
                        continue
                    datos = campos(fila)
                    self.assertTrue(datos["APENOM"].startswith("PERSONA FICTICIA OSTEL "))
                    if datos["EMAIL1"] is not None:
                        self.assertTrue(datos["EMAIL1"].endswith("@example.invalid"))
                    self.assertIsNone(datos["A_TEL1"])
                    self.assertIsNone(datos["CELU1"])

    def test_numero_compartido_con_personas_distintas(self):
        a, b = self.filas("numero_compartido")
        self.assertEqual(a["NROAFI"], b["NROAFI"])
        for columna in ("APENOM", "NRODOC", "CUIL"):
            self.assertNotEqual(a[columna], b[columna])

    def test_identidad_confirmada_y_conflictos_sin_resolucion_automatica(self):
        a, b = self.filas("identidad_individual")
        self.assertEqual(a["NROAFI"], b["NROAFI"])
        self.assertNotEqual(a["CUIL"], b["CUIL"])
        self.assertNotEqual(a["NUMERO"], b["NUMERO"])
        for caso, igual, distinto in (
            ("cuil_distinto_mismo_numero", "NUMERO", "CUIL"),
            ("numero_distinto_mismo_cuil", "CUIL", "NUMERO"),
        ):
            a, b = self.filas(caso)
            self.assertEqual(a[igual], b[igual])
            self.assertNotEqual(a[distinto], b[distinto])

    def test_cambios_administrativos_conservan_identidad_y_fechas(self):
        for caso, campo in (("cambio_empresa", "CODEMP"), ("cambio_categoria", "TIPO")):
            a, = self.filas(caso, 0)
            b, = self.filas(caso, 1)
            self.assertNotEqual(a[campo], b[campo])
            self.assertEqual({k: v for k, v in a.items() if k != campo},
                             {k: v for k, v in b.items() if k != campo})
            self.assertIsNone(b["FECBAJ"])

    def test_categoria_y_parentesco_no_prueban_vigencia_ni_familia(self):
        a, p = self.filas("categorias_sin_cobertura")
        self.assertEqual((a["TIPO"], p["TIPO"]), ("A", "P"))
        self.assertIsNotNone(a["FECBAJ"])
        self.assertIsNone(p["FECBAJ"])
        self.assertEqual((a["COD_PLAN"], p["COD_PLAN"]), ("CC", "CC"))
        parientes = self.filas("parentescos_confirmados")
        self.assertEqual(len(parientes), 56)
        self.assertEqual(len({r["CODPAR"] for r in parientes}), 56)
        self.assertTrue({"AAA", "H01", "H15", "XD1", "XD2"}.issubset(
            {r["CODPAR"] for r in parientes}))
        # CUILTIT ficticio coherente: el titular se apunta a sí mismo y los
        # familiares al CUIL del titular. No arma familias (eso es #181).
        titular, = (r for r in parientes if r["CODPAR"] == "AAA")
        self.assertEqual(titular["CUILTIT"], titular["CUIL"])
        self.assertTrue(all(r["CUILTIT"] == titular["CUIL"] for r in parientes))
        self.assertTrue(all(r["NROREF"] is None for r in parientes))

    def test_referencias(self):
        for codigo in ("referencia_presente", "referencia_no_aaa"):
            a, b = self.filas(codigo)
            self.assertEqual(b["NROREF"], a["NROAFI"])
        self.assertTrue(self.filas("referencia_presente")[0]["NROAFI"].endswith("AAA"))
        self.assertFalse(self.filas("referencia_no_aaa")[0]["NROAFI"].endswith("AAA"))
        ausentes = self.filas("referencia_ausente")
        self.assertNotIn(ausentes[0]["NROREF"], {f["NROAFI"] for f in ausentes})
        presente = self.filas("referencia_presente")
        self.assertEqual(presente[1]["CUILTIT"], presente[0]["CUIL"])
        self.assertNotIn(ausentes[0]["CUILTIT"], {f["CUIL"] for f in ausentes})

    def test_checksum_y_discrepancia_son_estimulos_distintos(self):
        base = self.filas("base")[0]
        self.assertTrue(checksum_valido(base["CUIL"]))
        self.assertEqual(base["CUIL"][2:10], base["NRODOC"])
        invalido = self.filas("cuil_invalido")[0]
        self.assertFalse(checksum_valido(invalido["CUIL"]))
        self.assertEqual(invalido["CUIL"][2:10], invalido["NRODOC"])
        discordante = self.filas("documento_discordante")[0]
        self.assertTrue(checksum_valido(discordante["CUIL"]))
        self.assertNotEqual(discordante["CUIL"][2:10], discordante["NRODOC"])

    def test_vacias_y_tipos_sin_normalizacion(self):
        filas = self.catalogo["filas_vacias"].lotes[0].filas
        self.assertEqual([i for i, f in enumerate(filas) if all(v is None for v in f)], [0, 2, 4])
        a, b, c = self.filas("tipos_mixtos")
        self.assertEqual(a["NROAFI"], "000185")
        self.assertEqual([type(f["NROAFI"]) for f in (a, b, c)], [str, int, float])
        self.assertIs(type(c["NRODOC"]), bool)
        self.assertIsInstance(a["FECING"], date)
        self.assertEqual(b["FECING"], "2026-01-01")
        self.assertEqual(c["FECING"], "no-es-fecha")

    def test_secuencias_entre_archivos(self):
        baja, repetida = self.catalogo["baja_repetida"].lotes
        self.assertNotEqual(baja.nombre, repetida.nombre)
        self.assertEqual(baja.filas, repetida.filas)
        antes, despues = (self.filas("posible_reingreso", i)[0] for i in (0, 1))
        for columna in ("NROAFI", "NRODOC", "CUIL"):
            self.assertEqual(antes[columna], despues[columna])
        self.assertLess(antes["FECBAJ"], despues["FECING"])
        self.assertIsNone(despues["FECBAJ"])

    def test_fechas_y_lote_parcial(self):
        self.assertEqual(self.filas("vencimiento_lejano")[0]["FECVTO"], date(2099, 12, 31))
        conflicto = self.filas("baja_anterior_ingreso")[0]
        self.assertLess(conflicto["FECBAJ"], conflicto["FECING"])
        a, b, c = self.filas("lote_parcial")
        self.assertTrue(checksum_valido(a["CUIL"]))
        self.assertFalse(checksum_valido(b["CUIL"]))
        self.assertEqual(b["CUIL"][2:10], b["NRODOC"])
        self.assertLess(c["FECBAJ"], c["FECING"])
        self.assertEqual(len({f["NROAFI"] for f in (a, b, c)}), 3)

    def test_referencia_ambigua_no_equivale_a_fila_duplicada(self):
        a, b, familiar = self.filas("referencia_ambigua")
        self.assertEqual(a["NROAFI"], familiar["NROREF"])
        self.assertEqual(b["NROAFI"], familiar["NROREF"])
        self.assertEqual(len({f["NRODOC"] for f in (a, b, familiar)}), 3)
        duplicadas = self.catalogo["fila_repetida"].lotes[0].filas
        self.assertEqual(len(duplicadas), 2)
        self.assertEqual(duplicadas[0], duplicadas[1])

    def test_codigos_contactos_y_texto_sin_inferencias(self):
        codigos = self.filas("codigos_sin_homologar")[0]
        for columna, valor in {
            "NUMERO": "000042", "TIPO": "TIPO-FICTICIO", "TIPODOC": "DOC-FICTICIO",
            "CODPAR": "PAR-FICTICIO", "MOTIVO": "MOT-FICTICIO", "INCAPA": "INC-FICTICIA",
        }.items():
            self.assertEqual(codigos[columna], valor)
        ausentes = self.filas("contactos_ausentes")[0]
        for columna in (
            "A_CALLE", "A_NRO", "A_PISO", "A_DPTO", "A_TEL1", "CELU1", "EMAIL1",
            "CODPOS", "DESLOCA", "CODPROV", "CODZON", "CODLOCA",
        ):
            self.assertIsNone(ausentes[columna])
        texto = self.filas("texto_literal")[0]
        self.assertEqual(texto["NROAFI"], " 000185-Ñ ")
        self.assertEqual(texto["NUMERO"], "000042")
        self.assertEqual(texto["NROREF"], "REF-FICTICIA/α")

    def test_fuentes_discordantes_no_sugieren_orden_temporal(self):
        a, b = (self.filas("fuentes_discordantes", i)[0] for i in (0, 1))
        for columna in ("NROAFI", "NRODOC", "CUIL", "FECING"):
            self.assertEqual(a[columna], b[columna])
        self.assertIsNotNone(a["FECBAJ"])
        self.assertIsNone(b["FECBAJ"])
        self.assertNotEqual(a["COD_PLAN"], b["COD_PLAN"])

    def test_xlsx_ambas_hojas_todos_los_casos(self):
        for escenario in self.catalogo.values():
            for lote in escenario.lotes:
                for hoja in HOJAS:
                    with self.subTest(caso=escenario.codigo, lote=lote.nombre, hoja=hoja):
                        libro = load_workbook(BytesIO(generar_xlsx(lote, hoja=hoja)))
                        try:
                            self.assertEqual(libro.sheetnames, [hoja])
                            filas = list(libro[hoja].values)
                            self.assertEqual(filas[0], COLUMNAS)
                            esperado = [
                                tuple(datetime.combine(v, datetime.min.time()) if type(v) is date else v
                                      for v in fila) for fila in lote.filas
                            ]
                            self.assertEqual(filas[1:], esperado)
                            for original, leida in zip(esperado, filas[1:]):
                                self.assertEqual(tuple(map(type, original)), tuple(map(type, leida)))
                        finally:
                            libro.close()

    def test_exportador_rechaza_hoja_y_ancho_incorrectos(self):
        with self.assertRaises(ValueError):
            generar_xlsx(self.catalogo["base"].lotes[0], hoja="Carga")
        # 39 ya no es válido: CUILTIT es obligatorio desde el 26/10/2026.
        for ancho in (0, 38, 39, 41):
            with self.subTest(ancho=ancho), self.assertRaises(ValueError):
                generar_xlsx(Lote("malformado", ((None,) * ancho,)))

    def test_xlsx_solo_filas_vacias_y_lote_sin_filas(self):
        for cantidad in (0, 1, 3):
            with self.subTest(cantidad=cantidad):
                filas = ((None,) * len(COLUMNAS),) * cantidad
                libro = load_workbook(BytesIO(generar_xlsx(Lote("vacias", filas))))
                try:
                    self.assertEqual(list(libro.active.values), [COLUMNAS, *filas])
                finally:
                    libro.close()

    def test_exportacion_temporal_completa_sin_sobrescritura(self):
        with TemporaryDirectory() as temporal:
            with patch("tempfile.mkdtemp", return_value=temporal):
                destino = exportar_temporales()
                self.assertEqual(destino, Path(temporal))
                rutas = {
                    Path(hoja) / escenario.codigo / f"{lote.nombre}.xlsx"
                    for hoja in HOJAS for escenario in generar_escenarios()
                    for lote in escenario.lotes
                }
                self.assertEqual({p.relative_to(destino) for p in destino.rglob("*.xlsx")}, rutas)
                manifiesto = json.loads((destino / "manifiesto.json").read_text(encoding="utf-8"))
                self.assertEqual(manifiesto["columnas"], list(COLUMNAS))
                self.assertEqual(manifiesto["aceptacion_ostel"], "pendiente")
                entradas = manifiesto["archivos"]
                self.assertEqual(len(entradas), len(rutas))
                self.assertEqual({e["archivo"] for e in entradas}, {p.as_posix() for p in rutas})
                indice = {e["archivo"]: e for e in entradas}
                for ruta in rutas:
                    libro = load_workbook(destino / ruta)
                    try:
                        self.assertEqual(libro.sheetnames, [ruta.parts[0]])
                        self.assertEqual(next(libro.active.values), COLUMNAS)
                        entrada = indice[ruta.as_posix()]
                        self.assertEqual(entrada["hoja"], libro.active.title)
                        filas = list(libro.active.values)[1:]
                        self.assertEqual(entrada["filas_fisicas_sin_encabezado"], len(filas))
                        self.assertEqual(entrada["filas_excel_con_datos"], [
                            i for i, fila in enumerate(filas, 2)
                            if any(v is not None for v in fila)
                        ])
                        self.assertEqual(entrada["sha256"], sha256((destino / ruta).read_bytes()).hexdigest())
                        escenario = self.catalogo[entrada["escenario"]]
                        esperado = escenario.lotes[entrada["orden_en_escenario"] - 1]
                        self.assertEqual(ruta.stem, esperado.nombre)
                    finally:
                        libro.close()
                original = {ruta: (destino / ruta).read_bytes() for ruta in rutas}
                manifiesto_original = (destino / "manifiesto.json").read_bytes()
                with self.assertRaises(FileExistsError):
                    exportar_temporales()
                self.assertEqual(original, {ruta: (destino / ruta).read_bytes() for ruta in rutas})
                self.assertEqual(manifiesto_original, (destino / "manifiesto.json").read_bytes())

    def test_texto_literal_no_formula(self):
        fila = list(self.catalogo["base"].lotes[0].filas[0])
        fila[COLUMNAS.index("APENOM")] = "=1+1"
        libro = load_workbook(BytesIO(generar_xlsx(Lote("literal", (tuple(fila),)))))
        try:
            self.assertEqual(libro.active["E2"].value, "=1+1")
            self.assertEqual(libro.active["E2"].data_type, "s")
        finally:
            libro.close()


class ContratoLectorOSTELTests(unittest.TestCase):
    """Contrato local publicado por #180; no prueba aplicación ni cobertura."""

    def test_manifiesto_y_celdas_de_todos_los_archivos_con_lector_real(self):
        from openpyxl.utils.datetime import to_excel
        from .ostel_prevalidacion import (
            COLUMNAS as COLUMNAS_LECTOR, HOJAS as HOJAS_LECTOR, prevalidar_ostel,
        )

        self.assertEqual(COLUMNAS_LECTOR, COLUMNAS)
        self.assertEqual(HOJAS_LECTOR, HOJAS)
        catalogo = {e.codigo: e for e in generar_escenarios()}
        with TemporaryDirectory() as temporal:
            with patch("tempfile.mkdtemp", return_value=temporal):
                destino = exportar_temporales()
            manifiesto = json.loads((destino / "manifiesto.json").read_text(encoding="utf-8"))
            for entrada in manifiesto["archivos"]:
                with self.subTest(archivo=entrada["archivo"]):
                    contenido = (destino / entrada["archivo"]).read_bytes()
                    resultado = prevalidar_ostel(contenido)
                    hoja = entrada["hoja"]
                    self.assertEqual(set(resultado.filas), {hoja})
                    filas = resultado.filas[hoja]
                    self.assertEqual([f.numero for f in filas], entrada["filas_excel_con_datos"])
                    duplicados = 1 if entrada["escenario"] == "fila_repetida" else 0
                    self.assertEqual(resultado.resumen, {
                        "filas_no_vacias": {hoja: len(entrada["filas_excel_con_datos"])},
                        "duplicados_exactos": duplicados, "anomalias": duplicados,
                    })
                    self.assertEqual([
                        (a.codigo, a.hoja, a.fila, a.columna, a.fila_original)
                        for a in resultado.anomalias
                    ], [("fila_duplicada", hoja, 3, None, 2)] if duplicados else [])
                    lote = catalogo[entrada["escenario"]].lotes[entrada["orden_en_escenario"] - 1]
                    for fila in filas:
                        self.assertEqual(len(fila.celdas), len(COLUMNAS))
                        for columna, (celda, original) in enumerate(
                            zip(fila.celdas, lote.filas[fila.numero - 2]), start=1
                        ):
                            with self.subTest(fila=fila.numero, columna=columna):
                                self.assertIsNone(celda.formula)
                                if type(original) is date:
                                    self.assertEqual(celda.tipo, "n")
                                    self.assertEqual(celda.valor, str(int(to_excel(original))))
                                    self.assertEqual(celda.formato, "yyyy-mm-dd")
                                elif type(original) is bool:
                                    self.assertEqual((celda.tipo, celda.valor),
                                                     ("b", "1" if original else "0"))
                                elif isinstance(original, (int, float)):
                                    self.assertEqual((celda.tipo, celda.valor), ("n", str(original)))
                                else:
                                    self.assertEqual(celda.valor, original)
                    self.assertEqual(sha256(contenido).hexdigest(), entrada["sha256"])
                    self.assertEqual((destino / entrada["archivo"]).read_bytes(), contenido)

    def test_lecturas_independientes_sin_consolidar_movimientos(self):
        from .ostel_prevalidacion import prevalidar_ostel

        for escenario in generar_escenarios():
            if len(escenario.lotes) < 2:
                continue
            for hoja in HOJAS:
                with self.subTest(escenario=escenario.codigo, hoja=hoja):
                    archivos = [generar_xlsx(lote, hoja=hoja) for lote in escenario.lotes]
                    resultados = [prevalidar_ostel(archivo) for archivo in archivos]
                    # Leer otra fuente no modifica el resultado de volver a leer la primera.
                    self.assertEqual(prevalidar_ostel(archivos[0]), resultados[0])
                    self.assertTrue(all(r.resumen["filas_no_vacias"] == {hoja: 1}
                                        for r in resultados))
                    self.assertTrue(all(r.anomalias == () for r in resultados))
                    if escenario.codigo == "baja_repetida":
                        self.assertEqual(resultados[0], resultados[1])
