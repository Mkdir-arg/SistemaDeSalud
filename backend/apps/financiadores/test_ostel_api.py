"""Integración HTTP OSTEL con base de pruebas aislada y fixtures del contrato #185."""
from unittest.mock import patch
from io import BytesIO

from openpyxl import load_workbook

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from django.test import TestCase
from rest_framework.test import APIClient

from apps.accounts.models import Usuario
from .models import Afiliado, EventoCobertura, Financiador, Importacion, MembresiaFinanciador
from .ostel_prevalidacion import COLUMNAS, COLUMNAS_CON_TITULAR, HOJAS, MAX_BYTES, prevalidar_ostel
from .ostel_synthetic import COLUMNAS as COLUMNAS_SINTETICAS, Lote, generar_escenarios, generar_xlsx


class OstelApiTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.org = Financiador.objects.create(nombre="OSTEL FICTICIA #180", tipo="obra_social")
        cls.usuario = Usuario.objects.create_user("ostel-api@example.invalid", nombre="Prueba ficticia")
        cls.miembro = MembresiaFinanciador.objects.create(
            financiador=cls.org, usuario=cls.usuario, rol="operador",
        )

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.usuario)
        self.url = f"/api/financiadores/{self.org.pk}/prevalidar-ostel/"

    def enviar(self, contenido=None, nombre="ficticio.xlsx"):
        if contenido is None:
            contenido = generar_xlsx(generar_escenarios()[0].lotes[0])
        return self.client.post(self.url, {"archivo": SimpleUploadedFile(nombre, contenido)}, format="multipart")

    def test_contrato_185_todos_los_archivos_por_api_sin_persistir_padron(self):
        self.assertEqual(COLUMNAS, COLUMNAS_SINTETICAS)
        enviados = 0
        for hoja in HOJAS:
            for escenario in generar_escenarios():
                for lote in escenario.lotes:
                    with self.subTest(hoja=hoja, escenario=escenario.codigo, lote=lote.nombre):
                        respuesta = self.enviar(generar_xlsx(lote, hoja=hoja))
                        self.assertEqual(respuesta.status_code, 200, respuesta.data)
                        esperadas = sum(any(v not in (None, "") for v in fila) for fila in lote.filas)
                        self.assertEqual(respuesta.data["resumen"]["filas_no_vacias"], {hoja: esperadas})
                        # Los estímulos funcionales no se convierten en rechazos técnicos.
                        self.assertEqual(respuesta.data["resumen"]["anomalias"],
                                         1 if escenario.codigo == "fila_repetida" else 0)
                        self.assertNotIn("PERSONA FICTICIA", str(respuesta.data))
                        self.assertEqual(respuesta["Cache-Control"], "private, no-store")
                        enviados += 1
        self.assertEqual(enviados, len(HOJAS) * sum(len(c.lotes) for c in generar_escenarios()))
        self.assertEqual(Afiliado.objects.count(), 0)
        self.assertEqual(Importacion.objects.count(), 0)
        eventos = EventoCobertura.objects.filter(accion="prevalidar_ostel")
        self.assertEqual(eventos.count(), enviados)
        self.assertFalse(eventos.exclude(motivo="", objeto=str(self.org.pk)).exists())

    def test_permisos_existentes_y_ambito_antes_de_leer(self):
        for cambios, estado in (({"rol": "auditor"}, 403), ({"activo": False}, 404)):
            with self.subTest(cambios=cambios):
                MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(**cambios)
                with patch("apps.financiadores.ostel_prevalidacion.prevalidar_ostel") as lector:
                    self.assertEqual(self.enviar().status_code, estado)
                    lector.assert_not_called()
                MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(rol="operador", activo=True)
        ajena = Financiador.objects.create(nombre="Ajena ficticia", tipo="mutual")
        self.url = f"/api/financiadores/{ajena.pk}/prevalidar-ostel/"
        self.assertEqual(self.enviar().status_code, 404)
        self.client.force_authenticate(None)
        self.assertEqual(self.enviar().status_code, 401)
        self.assertEqual(EventoCobertura.objects.count(), 0)

    def test_admin_puede_revisar_sin_designaciones_protegidas(self):
        MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(rol="admin")
        self.assertEqual(self.enviar().status_code, 200)

    def test_textos_del_contrato_185_se_conservan_sin_homologar(self):
        for escenario in generar_escenarios():
            for lote in escenario.lotes:
                with self.subTest(escenario=escenario.codigo, lote=lote.nombre):
                    resultado = prevalidar_ostel(generar_xlsx(lote))
                    por_numero = {fila.numero: fila for fila in resultado.filas["Hoja1"]}
                    for numero, original in enumerate(lote.filas, 2):
                        for columna, valor in enumerate(original):
                            if isinstance(valor, str) and valor != "":
                                self.assertEqual(por_numero[numero].celdas[columna].valor, valor)

    def test_archivo_faltante_invalido_extension_y_tamano(self):
        self.assertEqual(self.client.post(self.url, {}, format="multipart").status_code, 400)
        for contenido, nombre in ((b"roto", "ficticio.xlsx"), (b"roto", "ficticio.xls"),
                                  (b"x" * (MAX_BYTES + 1), "ficticio.xlsx")):
            with self.subTest(nombre=nombre, largo=len(contenido)):
                respuesta = self.enviar(contenido, nombre)
                self.assertEqual(respuesta.status_code, 400)
                self.assertIn("archivo", respuesta.data)
        self.assertFalse(EventoCobertura.objects.exists())

    def test_detalle_acotado_no_recorta_conteos(self):
        fila = generar_escenarios()[0].lotes[0].filas[0]
        respuesta = self.enviar(generar_xlsx(Lote("duplicados", (fila,) * 502)))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.data["resumen"]["filas_no_vacias"]["Hoja1"], 502)
        self.assertEqual(respuesta.data["anomalias_por_codigo"], {"fila_duplicada": 501})
        self.assertEqual(len(respuesta.data["anomalias"]), 500)
        self.assertTrue(respuesta.data["detalle_truncado"])
        self.assertEqual(respuesta.data["anomalias"][0]["fila_original"], 2)

    def test_auditoria_fallida_no_entrega_resultado(self):
        with patch("apps.financiadores.views.auditar", side_effect=IntegrityError):
            respuesta = self.enviar()
        self.assertEqual(respuesta.status_code, 400)
        self.assertNotIn("resumen", respuesta.data)
        self.assertFalse(Importacion.objects.exists())

    def test_plantillas_por_formato_preservan_sistema_y_ofrecen_ostel_vacia(self):
        url = f"/api/financiadores/{self.org.pk}/plantilla/"
        for parametros, hoja in (
            ({"tipo": "padron", "version": 2}, "Carga"),
            ({"tipo": "consumos"}, "Carga"),
            ({"tipo": "padron", "formato": "ostel"}, "Hoja1"),
        ):
            with self.subTest(parametros=parametros):
                respuesta = self.client.get(url, parametros)
                self.assertEqual(respuesta.status_code, 200)
                libro = load_workbook(BytesIO(respuesta.content))
                try:
                    self.assertIn(hoja, libro.sheetnames)
                    if hoja == "Hoja1":
                        self.assertEqual(tuple(c.value for c in libro[hoja][1]), COLUMNAS_CON_TITULAR)
                        self.assertEqual(prevalidar_ostel(respuesta.content).anomalias, ())
                        self.assertEqual(libro[hoja].max_row, 1)
                        self.assertEqual(libro[hoja].column_dimensions["B"].number_format, "@")
                        self.assertEqual(prevalidar_ostel(respuesta.content).resumen["filas_no_vacias"], {"Hoja1": 0})
                finally:
                    libro.close()
        for parametros in ({"tipo": "consumos", "formato": "ostel"}, {"tipo": "padron", "formato": "desconocido"}):
            self.assertEqual(self.client.get(url, parametros).status_code, 400)
        MembresiaFinanciador.objects.filter(pk=self.miembro.pk).update(rol="auditor")
        self.assertEqual(self.client.get(url, {"tipo": "padron", "formato": "ostel"}).status_code, 403)
