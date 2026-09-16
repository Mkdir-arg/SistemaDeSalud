"""Genera muestras de formato para revisión; no es el importador de Cauce.

Usa openpyxl ya instalado en el entorno de artefactos. No modifica dependencias
del backend. Desde raíz: python docs/plans/financiadores/crear_plantillas_demo.py
"""

from datetime import datetime
from pathlib import Path
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


SALIDA = Path(__file__).resolve().parent / "plantillas"
AZUL, VERDE, GRIS = "16324F", "147D73", "EAF1F5"
CATALOGO = {"PREST-DEMO-001": "Consulta", "PREST-DEMO-002": "Radiografía",
            "PREST-DEMO-003": "Ecografía"}
FILAS = 200  # Sólo el área preparada de la muestra; no define un límite productivo.


def titulo(hoja, texto, columnas):
    hoja.sheet_view.showGridLines = False
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columnas)
    hoja.cell(1, 1, texto).font = Font(name="Arial", bold=True, size=15, color="FFFFFF")
    hoja.cell(1, 1).fill = PatternFill("solid", fgColor=AZUL)
    hoja.row_dimensions[1].height = 28
    hoja.freeze_panes = "A6"


def encabezados(hoja, nombres):
    for i, nombre in enumerate(nombres, 1):
        celda = hoja.cell(5, i, nombre)
        celda.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[celda.column_letter].width = 25 if i != 3 else 46
    hoja.row_dimensions[5].height = 33


def validacion(hoja, tipo, rango, formula1, formula2=None, mensaje="Revisá el valor ingresado."):
    dv = DataValidation(type=tipo, formula1=formula1, formula2=formula2,
                        allow_blank=True, showErrorMessage=True, showInputMessage=True)
    dv.errorTitle, dv.error = "Dato a revisar", mensaje
    dv.promptTitle, dv.prompt = "Completar la plantilla", mensaje
    if tipo in ("date", "whole"):
        dv.operator = "between"
    if tipo == "list":
        dv.showDropDown = False
    hoja.add_data_validation(dv)
    dv.add(rango)


def crear(nombre_archivo, financiador, codigos, padron=False):
    libro = Workbook()
    instrucciones = libro.active
    instrucciones.title = "LEEME"
    titulo(instrucciones, f"Cauce · {financiador} · MUESTRA PARA REVISIÓN", 3)
    notas = [
        "PROTOTIPO de formato: el sistema actual todavía no dispone de este importador.",
        "Datos y organizaciones ficticios. Revisar antes de usar como contrato de producción.",
        "Completá la hoja Carga. Los ejemplos están separados y no son filas para importar.",
        "Número de afiliado y documento se completan como texto: conservá los ceros iniciales.",
        "La lista de selección y la hoja de referencia corresponden a este financiador de ejemplo.",
        "Las fechas se muestran como día/mes/año. No ingresar fórmulas ni macros.",
        "La muestra prepara 200 filas. El volumen y los límites reales siguen por definir (Q07).",
        "La validación de Excel ayuda a cargar; el servidor debe comprobar todos los datos.",
        "Las hojas de contexto o referencia no acreditan permisos ni pertenencia a una organización.",
    ]
    if padron:
        notas += ["Carga incremental: agrega/actualiza los incluidos. No realiza bajas.",
                  "Nombre, plan y vigente desde son columnas propuestas (Q03–Q05).",
                  "No se importa una fecha de baja ni se desactiva a personas omitidas."]
    else:
        notas += ["Prestación, fecha, cantidad, número de afiliado y documento son obligatorios.",
                  "Referencia externa es opcional. Ayuda a reconocer el mismo consumo al reenviarlo.",
                  "Un consumo externo actualiza cupo y no genera cargos de un hospital de Cauce.",
                  "Antes de aplicar: resumen de válidas, rechazadas y filas que necesitan revisión."]
    instrucciones.column_dimensions["A"].width = 112
    for fila, nota in enumerate(notas, 3):
        celda = instrucciones.cell(fila, 1, nota)
        celda.font = Font(name="Arial", size=11)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        instrucciones.row_dimensions[fila].height = 31

    carga = libro.create_sheet("Carga")
    nombres = (["N.º de afiliado", "Documento", "Nombre y apellido (propuesto)",
                "Plan (propuesto)", "Vigente desde (propuesto)"] if padron else
               ["N.º de afiliado", "Documento", "Prestación", "Fecha de prestación",
                "Cantidad", "Referencia externa (opcional)"])
    titulo(carga, "PADRÓN · BORRADOR" if padron else "CONSUMOS EXTERNOS · BORRADOR", len(nombres))
    carga.cell(3, 1, "Completá desde la fila 6. Consultá LEEME y la hoja de referencia.")
    encabezados(carga, nombres)
    for fila in carga.iter_rows(min_row=6, max_row=5 + FILAS, min_col=1, max_col=len(nombres)):
        for celda in fila:
            celda.font = Font(name="Arial", size=11, color="0000FF")
            celda.alignment = Alignment(vertical="center")
            celda.number_format = "@"
        fila[4 if padron else 3].number_format = "dd/mm/yyyy"
        if not padron:
            fila[4].number_format = "0"
    tabla = Table(displayName="DatosCarga", ref=f"A5:{chr(64 + len(nombres))}{5 + FILAS}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    carga.add_table(tabla)

    referencia = libro.create_sheet("Planes" if padron else "Prestaciones")
    titulo(referencia, "Referencia ficticia de " + financiador, 3)
    encabezados(referencia, ["Código estable", "Nombre", "Selección"])
    registros = ([("PLAN-DEMO-01", "Plan Básico"), ("PLAN-DEMO-02", "Plan Ampliado")]
                 if padron else [(codigo, CATALOGO[codigo]) for codigo in codigos])
    selecciones = []
    for fila, (codigo, etiqueta) in enumerate(registros, 6):
        seleccion = f"{codigo} | {etiqueta}"
        selecciones.append(seleccion)
        for columna, valor in enumerate((codigo, etiqueta, seleccion), 1):
            referencia.cell(fila, columna, valor).font = Font(name="Arial", size=11)
    referencia.column_dimensions["C"].width = 47
    nombre_lista = "OpcionesDeCarga"
    libro.defined_names.add(DefinedName(nombre_lista,
                            attr_text=f"'{referencia.title}'!$C$6:$C${5 + len(registros)}"))
    columna_lista = "D" if padron else "C"
    validacion(carga, "list", f"{columna_lista}6:{columna_lista}{5 + FILAS}", nombre_lista,
               mensaje="Elegí una opción de la hoja de referencia de este financiador.")
    columna_fecha = "E" if padron else "D"
    validacion(carga, "date", f"{columna_fecha}6:{columna_fecha}{5 + FILAS}", "2", "2958465",
               mensaje="Ingresá una fecha de Excel válida. La vigencia se verifica al importar.")
    if not padron:
        validacion(carga, "whole", f"E6:E{5 + FILAS}", "1", "2147483647",
                   mensaje="Ingresá una cantidad entera positiva; los límites reales se validan al importar.")

    ejemplos = libro.create_sheet("Ejemplos ficticios")
    titulo(ejemplos, "EJEMPLOS FICTICIOS · No se importan desde esta hoja", len(nombres))
    encabezados(ejemplos, nombres)
    for i in range(2):
        valores = ([f"0000123{i}", f"0000000{i + 1}", f"Persona de ejemplo {i + 1}",
                    selecciones[i], datetime(2026, 9, 1)] if padron else
                   [f"0000123{i}", f"0000000{i + 1}", selecciones[i], datetime(2026, 9, 10),
                    i + 1, "EXT-DEMO-001" if i == 0 else None])
        for columna, valor in enumerate(valores, 1):
            celda = ejemplos.cell(6 + i, columna, valor)
            celda.font = Font(name="Arial", size=11)
            celda.number_format = "dd/mm/yyyy" if isinstance(valor, datetime) else ("0" if isinstance(valor, int) else "@")

    contexto = libro.create_sheet("_contexto")
    contexto.append(["Muestra", "No acredita identidad ni permisos"])
    contexto.append(["financiador", financiador])
    contexto.append(["contrato", "borrador-1"])
    contexto.sheet_state = "hidden"
    libro.properties.creator = "Cauce · preparación de diseño"
    destino = SALIDA / nombre_archivo
    libro.save(destino)

    # Verificación del artefacto generado, sin probar un importador inexistente.
    revision = load_workbook(destino)
    if any(celda.data_type == "f" for hoja in revision for fila in hoja for celda in fila):
        raise RuntimeError("La muestra no debe contener celdas con fórmulas.")
    if revision["Ejemplos ficticios"]["A6"].value != "00001230":
        raise RuntimeError("Se perdieron los ceros del identificador de ejemplo.")
    if any(celda.value is not None for fila in revision["Carga"].iter_rows(min_row=6) for celda in fila):
        raise RuntimeError("La hoja de carga debe estar vacía.")
    print(f"OK: {destino.name}; carga vacía; identificadores textuales; catálogo personalizado; sin fórmulas")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    SALIDA.mkdir(parents=True, exist_ok=True)
    crear("consumos-obra-social-a-demo.xlsx", "Obra social A — ficticia", ["PREST-DEMO-001", "PREST-DEMO-002"])
    crear("consumos-mutual-b-demo.xlsx", "Mutual B — ficticia", ["PREST-DEMO-001", "PREST-DEMO-003"])
    crear("padron-obra-social-a-demo.xlsx", "Obra social A — ficticia", [], padron=True)
